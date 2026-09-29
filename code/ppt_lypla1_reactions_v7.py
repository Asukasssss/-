from pathlib import Path
import ast,shutil,json
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
ROOT=Path(__file__).resolve().parents[1];O=ROOT/'results/BRCA/07_INTEGRATION/20260929T_lypla1_reactions_v7';O.mkdir(exist_ok=True)
V=Path('D:/CodexData/visualizations/2026/09/29/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/CAMP_LYPLA1_reactions_v7');V.mkdir(exist_ok=True)
OLD=V.parent/'CAMP_case_art_v4/CAMP_完整组会汇报_生化插图精修_25页.pptx';R=Presentation(OLD)
A=O/'lipid_reactions.png';shutil.copy2('D:/CodexData/generated_images/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/exec-9dee593a-c4c3-44d6-8bc5-cef16d3c30bb.png',A)
mod=ast.parse((ROOT/'code/ppt_editable_main_v1.py').read_text(encoding='utf-8'));ns=dict(G='00553B',INK='183D34',M='6A7973',PALE='F0F5F1',LINE='DDE7E0',FONT='Microsoft YaHei')
exec(compile(ast.Module(body=[n for n in mod.body if isinstance(n,(ast.Import,ast.ImportFrom,ast.FunctionDef))],type_ignores=[]),'<helpers>','exec'),ns)
txt,box,line,el=[ns[n] for n in ['txt','box','line','el']];C=ns['PP_ALIGN'].CENTER
s=R.slides.add_slide(R.slide_layouts[6]);s.shapes.add_picture(str(ROOT/'results/BRCA/07_INTEGRATION/20260929T_illustrated_main_v2/assets/sysu_logo.png'),Inches(.65),Inches(.13),width=Inches(1.93),height=Inches(.75))
txt(s,10.02,.42,2.62,.28,'CAMP  /  组会汇报',11,'6A7973',align=ns['PP_ALIGN'].RIGHT);line(s,.65,.94,12.67,.94);line(s,.65,.94,1.4,.94,'00553B',2)
txt(s,.67,1.13,12,.52,'LYPLA1 如何连接 LPC、GPC 与油酸？',27,'00553B',True)
txt(s,.69,1.74,12,.32,'两条不同的磷脂水解反应：底物不同，释放的脂肪酸也不同',14,'6A7973')
def art(rect,x,y,w,h):
 iw,ih=Image.open(A).size;l,t,r,b=rect;f=min(w/(r-l),h/(b-t));dw,dh=(r-l)*f,(b-t)*f
 p=s.shapes.add_picture(str(A),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh));p.crop_left=l/iw;p.crop_right=1-r/iw;p.crop_top=t/ih;p.crop_bottom=1-b/ih
def flow(x,y,w):
 z=line(s,x,y,x+w,y,'00553B',1.8);z._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(el('a:tailEnd',type='triangle',w='med',len='med'))
for y,title,rects,left,right,role in [
 (2.19,'01  溶血磷脂水解',[(180,30,480,500),(550,45,970,460),(1060,30,1450,500)],'LPC(16:0)','GPC  +  棕榈酸','切下 LPC 的棕榈酰基，留下 GPC'),
 (4.40,'02  磷脂酰胆碱 sn-1 位水解',[(110,520,485,990),(550,525,970,955),(1050,520,1485,990)],'sn-1 含油酰基的 PC','2-酰基 LPC  +  油酸','切下 sn-1 油酰基，释放油酸')]:
 box(s,.68,y,11.97,2.05,'F2F6F2');txt(s,.9,y+.11,7,.28,title,14,'00553B',True)
 art(rects[0],1.0,y+.49,1.28,1.23);art(rects[1],5.64,y+.49,1.0,.86);art(rects[2],9.52,y+.45,2.23,1.18)
 txt(s,2.45,y+.84,2.9,.56,left,17,'00553B',True)
 txt(s,6.68,y+.65,1.12,.3,'LYPLA1',15,'00553B',True)
 flow(5.37,y+1.48,3.14);txt(s,5.17,y+1.70,3.6,.25,'催化酯键水解',11,'6A7973',align=C)
 txt(s,9.12,y+1.70,3.16,.3,right,16,'00553B',True,align=C)
 txt(s,2.47,y+1.51,2.65,.48,role,11,'6A7973')
txt(s,.69,6.70,12,.31,'油酸在第二条反应中是脂肪酸产物；不是 LPC(16:0) 水解的产物。',15,'00553B',True)
line(s,.68,7.12,12.65,7.12);txt(s,.69,7.20,11.4,.22,'来源：UniProt O75608 · Rhea 40435 / 41720｜概念示意，省略水与质子；不代表患者内净通量。',8.5,'6A7973')
ns['notes'](s,'插图由image_gen生成，脂质尾数与产物对应为概念图，不是精确化学结构或实际蛋白构象。LYPLA1溶血磷脂酶与磷脂酶A1作用分别展示。油酸是Rhea41720产物。来源https://www.rhea-db.org/rhea/41720及原冻结UniProt O75608映射。')
ids=R.slides._sldIdLst;item=ids[-1];ids.remove(item);ids.insert(9,item)
for n,sl in enumerate(R.slides,1):
 if n<10:continue
 for sh in list(sl.shapes):
  if sh.has_text_frame and sh.text.strip().isdigit() and sh.left>Inches(11.9) and sh.top>Inches(6.9):sl.shapes._spTree.remove(sh._element)
 txt(sl,12.12,7.03,.52,.27,str(n),11,'00553B',True,ns['PP_ALIGN'].RIGHT)
 for sh in sl.shapes:
  for e in sh._element.xpath('.//a:effectRef'):e.set('idx','0')
  for rp in sh._element.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr'):
   for tag,face in [('a:latin','Microsoft YaHei'),('a:ea','微软雅黑')]:
    old=rp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag.split(':')[1])
    if old is not None:rp.remove(old)
    rp.append(el(tag,typeface=face))
PPT='CAMP_LYPLA1反应图文精修_26页.pptx';R.save(O/PPT);shutil.copy2(O/PPT,V/PPT)
print(O/PPT)
