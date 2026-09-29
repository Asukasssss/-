from pathlib import Path
import ast,json,shutil
from pptx import Presentation
from pptx.util import Inches
ROOT=Path(__file__).resolve().parents[1];RUN='20260929T_complete_cases_v3';O=ROOT/'results/BRCA/07_INTEGRATION'/RUN;O.mkdir(exist_ok=True)
V=Path('D:/CodexData/visualizations/2026/09/29/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/CAMP_complete_cases_v3');V.mkdir(exist_ok=True)
SRC=Path('D:/CodexData/visualizations/2026/09/25/01a0d8cd-4545-75b0-ba43-922588deea1a/CAMP_representative_cases_v3')
OLD=ROOT/'results/BRCA/07_INTEGRATION/20260929T_illustrated_main_v2'
R=Presentation(OLD/'CAMP_泛癌与LYPLA1_插图精修可编辑版.pptx')
module=ast.parse((ROOT/'code/ppt_editable_main_v1.py').read_text(encoding='utf-8'))
ns=dict(G='00553B',INK='183D34',M='6A7973',PALE='F0F5F1',LINE='DDE7E0',FONT='Microsoft YaHei')
exec(compile(ast.Module(body=[n for n in module.body if isinstance(n,(ast.Import,ast.ImportFrom,ast.FunctionDef))],type_ignores=[]),'<native converter>','exec'),ns)
ns.update(R=R,SOURCES=[],AUDIT=[])
files=['01_BRCA_ASNS_CAMP','02_BRCA_ASNS_全细胞与患者','04_PDAC_PRAD_SLC6A6_CAMP','05_PDAC_SLC6A6_全细胞与患者','06_PRAD_SLC6A6_全细胞与患者','11_COAD_UCKL1_CAMP','09_COAD_UCKL1_全细胞与患者']
for n,name in enumerate(files,19):
 s=ns['convert'](SRC/(name+'.pdf'),n)
 for sh in list(s.shapes):
  if sh.top+sh.height<Inches(1.065) or (sh.has_text_frame and sh.text.strip().isdigit() and sh.left>Inches(11.9) and sh.top>Inches(6.9)):s.shapes._spTree.remove(sh._element)
 s.shapes.add_picture(str(OLD/'assets/sysu_logo.png'),Inches(.65),Inches(.13),width=Inches(1.93),height=Inches(.75))
 ns['txt'](s,10.02,.42,2.62,.28,'CAMP  /  组会汇报',11,'6A7973',align=ns['PP_ALIGN'].RIGHT)
 ns['line'](s,.65,.94,12.67,.94);ns['line'](s,.65,.94,1.40,.94,'00553B',2)
 ns['txt'](s,12.12,7.05,.52,.27,str(n),11,'00553B',True,ns['PP_ALIGN'].RIGHT)
 for sh in s.shapes:
  for e in sh._element.xpath('.//a:effectRef'):e.set('idx','0')
  for rp in sh._element.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr'):
   rp.set('lang','zh-CN')
   for tag,face in [('a:latin','Microsoft YaHei'),('a:ea','微软雅黑'),('a:cs','Microsoft YaHei')]:
    old=rp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag.split(':')[1])
    if old is not None:rp.remove(old)
    rp.append(ns['el'](tag,typeface=face))
assert len(R.slides)==25
PPT='CAMP_完整组会汇报_含三个案例_25页.pptx';R.save(O/PPT);shutil.copy2(O/PPT,V/PPT)
(O/'source_manifest.json').write_text(json.dumps(ns['SOURCES'],ensure_ascii=False,indent=2),encoding='utf-8')
(O/'append_audit.json').write_text(json.dumps(ns['AUDIT'],ensure_ascii=False,indent=2),encoding='utf-8')
print(O/PPT)
