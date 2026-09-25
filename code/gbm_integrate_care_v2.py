"""Replace only the GBM single-cell layer; preserve previous bulk statistics."""
from pathlib import Path
import json,hashlib,subprocess
import pandas as pd,numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.colors import Normalize
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill
R=Path(__file__).resolve().parents[1];RUN='20260925T151224Z_sc_replacement_v2';B=R/'results/GBM';S=B/'06_EXTERNAL'/RUN;O=B/'07_INTEGRATION'/RUN;OLD=B/'07_INTEGRATION/20260925T145000Z_integration_v1'
def load(p):return pd.read_csv(p,sep='\t')
def save(d,p):d.to_csv(p,sep='\t',index=False,na_rep='NA')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    O.mkdir(parents=True,exist_ok=False);F=O/'figures';F.mkdir()
    rank=load(S/'sc_source_stability.tsv');prof=load(S/'sc_celltype_profiles.tsv');registry=load(S/'sc_dataset_registry.tsv');comp=load(S/'sc_partition_comparison.tsv')
    drop=['study1_top','study2_top','study1_bootstrap','study2_bootstrap','same_top_all_categories','same_top_shared_categories','same_top_both_bootstrap_ge080','n_shared_categories']
    oldg=load(OLD/'candidate_genes_integrated.tsv');oldr=load(OLD/'candidate_relations_integrated.tsv');g=oldg.drop(columns=drop+['status']);r=oldr.drop(columns=drop+['status_cell_source'])
    for part,prefix in [('CARE2025_primary','SC_primary'),('CARE2025_recurrent','SC_recurrent')]:
        d=rank[rank.study.eq(part)].drop(columns=['study','stable_gene_id']).rename(columns={c:prefix+'_'+c for c in rank.columns if c not in ['study','stable_gene_id','gene']})
        g=g.merge(d,on='gene',validate='one_to_one');r=r.merge(d,on='gene',validate='many_to_one')
    for tab in [g,r]:
        tab['SC_dataset']='CARE2025_GSE274546';tab['SC_independent_replication']=False;tab['SC_scope']='primary main;recurrent same-cohort context;source expression only'
    assert len(g)==142 and len(r)==171
    comparisons={}
    for old,new,key,label in [(oldg,g,'gene','genes'),(oldr,r,'relation_id','relations')]:
        cols=[c for c in old if c.startswith(('RNA_','metabolite_','patient_'))]
        pd.testing.assert_frame_equal(old.set_index(key)[cols].sort_index(),new.set_index(key)[cols].sort_index(),check_exact=True)
        comparisons[label+'_bulk_columns_unchanged']=len(cols)
    save(g,O/'candidate_genes_integrated.tsv');save(r,O/'candidate_relations_integrated.tsv')
    ordered=sorted(g.gene);studies=['CARE2025_primary','CARE2025_recurrent'];types=['Neoplastic','Myeloid','Lymphoid','OPC','Oligodendrocyte','Astrocyte','Neuron','Vascular','Unresolved'];vmax=float(prof.effect.max());norm=Normalize(0,vmax);cmap=plt.get_cmap('viridis')
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'pdf.fonttype':42})
    with PdfPages(F/'GBM_all142_CARE_cell_sources.pdf') as pdf:
        for page,start in enumerate(range(0,len(ordered),48),1):
            gg=ordered[start:start+48];fig,axes=plt.subplots(1,2,figsize=(15,16),sharey=True)
            for ax,study in zip(axes,studies):
                df=prof[prof.cohort.eq(study)].set_index(['gene','celltype']);reg=registry[registry.study.eq(study)].iloc[0]
                for yi,gene in enumerate(gg):
                    for xi,t in enumerate(types):
                        if (gene,t) in df.index and df.loc[(gene,t),'status']=='DONE':
                            q=df.loc[(gene,t)];ax.scatter(xi,yi,s=8+130*q.mean_detection_fraction,c=[cmap(norm(q.effect))],edgecolors='none')
                        else:ax.scatter(xi,yi,s=12,c='#ccc',marker='x',linewidths=.7)
                ax.set_xticks(range(len(types)),types,rotation=60,ha='right');ax.set_yticks(range(len(gg)),gg);ax.set_ylim(len(gg)-.5,-.5);ax.set_xlim(-.7,len(types)-.3)
                ax.set_title(f'{study.replace("CARE2025_", "CARE 2025 ")} | {reg.patients} patients | {reg.cells:,} nuclei');ax.grid(axis='y',alpha=.13);ax.set_axisbelow(True);ax.spines[['top','right']].set_visible(False)
            fig.suptitle(f'GBM all 142 candidate genes | page {page}\nPrimary: main analysis; recurrence: same-cohort context, not independent replication\nColor: patient-equal mean log1p(CP10k); size: detection; unresolved not ranked',fontsize=12)
            fig.subplots_adjust(left=.12,right=.9,bottom=.12,top=.91,wspace=.12);ca=fig.add_axes([.92,.3,.015,.35]);fig.colorbar(plt.cm.ScalarMappable(norm=norm,cmap=cmap),cax=ca,label='Mean log1p(CP10k)');pdf.savefig(fig);fig.savefig(F/f'GBM_CARE_sources_page{page}.png',dpi=160);plt.close(fig)
    # Within-gene scaled heatmap: relative sources, not cross-gene abundance.
    mainprof=prof[prof.cohort.eq('CARE2025_primary')];mat=mainprof.pivot(index='gene',columns='celltype',values='effect').reindex(index=ordered,columns=types);scaled=mat.div(mat.max(axis=1).replace(0,np.nan),axis=0)
    fig,ax=plt.subplots(figsize=(9,28));im=ax.imshow(scaled,aspect='auto',vmin=0,vmax=1,cmap='viridis');ax.set_xticks(range(len(types)),types,rotation=60,ha='right');ax.set_yticks(range(len(ordered)),ordered,fontsize=6);ax.set_title('CARE 2025 primary | all 142 genes\nWithin-gene relative patient-equal expression');fig.colorbar(im,ax=ax,fraction=.025,pad=.02,label='Mean / gene maximum');fig.tight_layout();fig.savefig(F/'GBM_CARE_primary_all142_heatmap.pdf');fig.savefig(F/'GBM_CARE_primary_all142_heatmap.png',dpi=160);plt.close(fig)
    ly=prof[prof.gene.eq('LYPLA1')];save(ly,O/'LYPLA1_CARE_profiles.tsv')
    fig,axes=plt.subplots(1,2,figsize=(13,5),sharey=True,sharex=True)
    for ax,study in zip(axes,studies):
        z=ly[ly.cohort.eq(study)].set_index('celltype').reindex(types);ax.barh(types,z.effect,color=['#bf4c45' if x=='Neoplastic' else '#4389a5' for x in types]);ax.set_title(study);ax.set_xlabel('Patient-equal mean log1p(CP10k)');ax.spines[['top','right']].set_visible(False)
        for i,(_,row) in enumerate(z.iterrows()):
            if np.isfinite(row.effect):ax.text(row.effect+.015,i,f'n={row.n:.0f}; detected {row.mean_detection_fraction:.0%}',va='center',fontsize=8)
        ax.set_xlim(0,ly.effect.max()*1.65)
    axes[0].invert_yaxis();fig.suptitle('LYPLA1 | primary main analysis and same-cohort recurrence context');fig.tight_layout();fig.savefig(F/'LYPLA1_CARE_sources.png',dpi=180);fig.savefig(F/'LYPLA1_CARE_sources.pdf');plt.close(fig)
    notes=pd.DataFrame([('更新','单细胞替换为CARE2025；初发为主，复发为同队列补充，不称两队列独立验证'),('统计保持','此前代谢物、bulk RNA、肿瘤关联统计原样复用；142基因171关系均保留'),('单核数据','10x snRNA-seq；作者质控与注释；全基因库CP10k；患者等权'),('证据边界','表达来源不证明代谢物来源、酶活或机制；Other不参与明确来源排名'),('初发代表性','可再次手术纵向队列，有临床选择偏倚；复发多标本在患者内合并'),('历史','Darmanis2017/Neftel2019留在历史版本，不再作为当前主来源')],columns=['项目','说明'])
    workbook=O/'GBM_CARE2025_全候选142基因171关系.xlsx'
    with pd.ExcelWriter(workbook,engine='openpyxl') as wr:
        for name,d in [('阅读说明',notes),('全部142基因',g),('全部171关系',r),('CARE队列',registry),('来源排名',rank),('类别长表',prof),('初发复发描述比较',comp),('LYPLA1',ly),('基因覆盖',load(S/'sc_gene_coverage.tsv'))]:d.to_excel(wr,sheet_name=name,index=False)
    wb=load_workbook(workbook)
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True)
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(45,max(14,len(str(col[0].value))+2))
    wb['阅读说明'].column_dimensions['B'].width=100;wb.save(workbook)
    assert wb['全部142基因'].max_row==143 and wb['全部171关系'].max_row==172
    val=dict(status='DONE',scope='single-cell replacement and integration;biochemical mapping remains partial',genes=142,relations=171,source_partitions=registry.to_dict('records'),bulk_statistics_preserved=comparisons,independent_cohorts=1,ranked_primary=int(rank[rank.study.eq(studies[0])].status.eq('DONE').sum()),ranked_recurrent=int(rank[rank.study.eq(studies[1])].status.eq('DONE').sum()),LYPLA1=rank[rank.gene.eq('LYPLA1')].to_dict('records'))
    (O/'validation.json').write_text(json.dumps(val,indent=2));(O/'analysis_spec.json').write_text(json.dumps(dict(version='gbm_care_integration_v2',code_commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=R,text=True).strip(),bulk_recomputed=False,new_SC_sources=True,independent_cohorts=1,plots='all142;heatmap within-gene relative mean;no selected significant gene filter'),indent=2))
    src=[OLD/'candidate_genes_integrated.tsv',OLD/'candidate_relations_integrated.tsv']+list(S.glob('*.tsv'))+[Path(__file__)]
    save(pd.DataFrame([dict(source_id=p.name,path=str(p.relative_to(R)),sha256=sha(p)) for p in src]),O/'source_manifest.tsv');save(pd.DataFrame([dict(file=str(p.relative_to(O)),sha256=sha(p)) for p in O.rglob('*') if p.is_file()]),O/'checksums.tsv')
    print(json.dumps(val,indent=2))
if __name__=='__main__':main()
