"""Three figures from frozen aggregate evidence; no new tests or embedding."""
from pathlib import Path
import io,subprocess,json,hashlib
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager,colors,patches
import fitz
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
R=Path(__file__).resolve().parents[1];SHA='8952c54b61709c875ec83881b66ee43bd1de59f3';B='results/BRCA/07_INTEGRATION/20260922T140000Z_report_v1/appendix/'
O=R/'results/BRCA/07_INTEGRATION/20260923T010000Z_choline_ethanolamine_v1';assert not O.exists();O.mkdir(parents=True)
manifest=[]
def read(name):
 p=B+name+'.tsv';b=subprocess.check_output(['git','show',SHA+':'+p],cwd=R);manifest.append(dict(path=p,commit=SHA,sha256=hashlib.sha256(b).hexdigest()));return pd.read_csv(io.BytesIO(b),sep='\t',dtype=str,keep_default_na=False)
def num(x):
 try:return float(x)
 except:return np.nan
def fmt(x):return f'{num(x):.3g}' if np.isfinite(num(x)) else 'NA'
def write(d,n):d.to_csv(O/(n+'.tsv'),sep='\t',index=False,na_rep='NA')
names=['choline','phosphocholine',"cytidine 5'-diphosphocholine",'glycerophosphorylcholine (GPC)','ethanolamine','phosphoethanolamine','1-palmitoyl-GPC (16:0)','2-palmitoyl-GPC (16:0)','1-stearoyl-GPC (18:0)','1-oleoylglycerophosphocholine','2-oleoylglycerophosphocholine','1-linoleoyl-GPC (18:2)','2-linoleoylglycerophosphocholine','2-arachidonoylglycerophosphocholine','2-palmitoleoyl-GPC (16:1)','1-arachidonoyl-GPE (20:4n6)','1-stearoylglycerophosphoethanolamine','2-arachidonoylglycerophosphoethanolamine','2-docosahexaenoylglycerophosphoethanolamine','2-linoleoylglycerophosphoethanolamine','2-oleoylglycerophosphoethanolamine','1-palmitoylplasmenylethanolamine']
labels=['胆碱','磷酸胆碱','CDP-胆碱','游离GPC','乙醇胺','磷酸乙醇胺','1-palmitoyl-GPC 16:0','2-palmitoyl-GPC 16:0','1-stearoyl-GPC 18:0','1-oleoyl-GPC 18:1','2-oleoyl-GPC 18:1','1-linoleoyl-GPC 18:2','2-linoleoyl-GPC 18:2','2-arachidonoyl-GPC 20:4','2-palmitoleoyl-GPC 16:1','1-arachidonoyl-GPE 20:4','1-stearoyl-GPE 18:0','2-arachidonoyl-GPE 20:4','2-docosahexaenoyl-GPE 22:6','2-linoleoyl-GPE 18:2','2-oleoyl-GPE 18:1','1-palmitoylplasmenylethanolamine']
m0=read('metabolites');p=m0[m0.analysis_type.eq('paired_processed')].set_index('metabolite_name');assert set(names)<=set(p.index)
m=m0[m0.metabolite_name.isin(names)].copy();keys=set(m.metabolite_key);rel=read('relations');rel=rel[rel.metabolite_key.isin(keys)].copy();genes=sorted(rel.gene.unique());assert len(names)==22 and len(rel)==41 and len(genes)==20
a=read('association');a=a[a.relation_id.isin(rel.relation_key)].copy();rna=read('rna');rna=rna[rna.gene.isin(genes)].copy();source=read('source_status');source=source[source.gene.isin(genes)].copy()
sc=pd.concat([read(s) for s in ['Wu2021_old','Wu2021_new','Pal2021_reprocessed_old','Pal2021_reprocessed_new']],ignore_index=True);sc=sc[sc.gene.isin(genes)&sc.partition.eq('ALL')];assert not sc.duplicated(['gene','cohort','celltype']).any()
scope=pd.DataFrame(dict(order=range(1,23),metabolite_name=names,display_label=labels,group=['胆碱乙醇胺小分子']*6+['具体LPC']*9+['乙醇胺相关脂质']*7));scope['metabolite_key']=scope.metabolite_name.map(p.metabolite_key);scope['direct_relation_count']=scope.metabolite_key.map(rel.groupby('metabolite_key').size()).fillna(0).astype(int)
for d,n in [(scope,'scope22'),(m,'metabolites44'),(rel,'relations41_history'),(a,'CAMP_associations82'),(rna,'RNA20'),(sc,'single_cell_profiles'),(source,'source_status')]:write(d,n)
summary=rel.copy()
for fam,prefix in [('processed_Spearman262','primary'),('available_Spearman262','sensitivity')]:
 q=a[a.analysis_type.eq(fam)].set_index('relation_id')
 for col in ['n','effect','ci_lower','ci_upper','p_value','q_value','status','reason']:summary['review_'+prefix+'_'+col]=summary.relation_key.map(q[col])
summary['review_note']='Existing associations; not function, flux or mediation'
write(summary,'topic_summary41')
font_manager.fontManager.addfont('C:/Windows/Fonts/msyh.ttc');plt.rcParams.update({'font.family':'Microsoft YaHei','axes.unicode_minus':False,'font.size':10,'pdf.fonttype':42})
def save(fig,name):
 for ext in ['png','pdf','svg']:fig.savefig(O/(name+'.'+ext),dpi=160,bbox_inches='tight',facecolor='white')
 plt.close(fig)
# Fig1 preserves scale and displays both independent families.
fig,axs=plt.subplots(1,4,figsize=(21,13),gridspec_kw={'width_ratios':[1.3,1.5,1.65,1.85]});fig.subplots_adjust(left=.22,right=.99,top=.88,bottom=.09,wspace=.12)
main=m[m.analysis_type.eq('paired_processed')].set_index('metabolite_name').loc[names];sen=m[m.analysis_type.eq('paired_author_available')].set_index('metabolite_name').loc[names]
for y,name in enumerate(names):
 r=main.loc[name];s=sen.loc[name];lo,eq,hi=map(num,[r.pairs_lower,r.pairs_equal,r.pairs_higher]);axs[0].barh(y,lo,color='#087F8C');axs[0].barh(y,eq,left=lo,color='#E1E5EA');axs[0].barh(y,hi,left=lo+eq,color='#D56E36');axs[0].text(46,y,f'{int(lo)}/{int(eq)}/{int(hi)}',va='center',fontsize=8)
 for row,dy,c in [(r,-.14,'#213A57'),(s,.14,'#D56E36')]:
  e,l,u=map(num,[row.effect,row.ci_lower,row.ci_upper])
  if np.isfinite(e):axs[1].errorbar(e,y+dy,xerr=[[e-l],[u-e]],fmt='o',ms=3,capsize=2,color=c)
 axs[2].text(0,y,f'{fmt(r.p_value)} / {fmt(r.q_value)}',va='center');axs[3].text(0,y,f'{int(num(s.n)):2d}对   {fmt(s.p_value)} / {fmt(s.q_value)}',va='center')
for ax in axs:ax.set_ylim(21.7,-1);ax.set_yticks(range(22));ax.tick_params(axis='y',length=0);ax.spines[['top','right']].set_visible(False)
axs[0].set_yticklabels(labels,fontsize=9);axs[0].set_xlim(0,63);axs[0].set_xticks([0,15,30,45]);axs[0].set_xlabel('降低／相等／升高（45对）');axs[0].set_title('A  配对方向人数',loc='left')
axs[1].set_yticklabels([]);axs[1].axvline(0,color='gray',ls='--');axs[1].set_title('B  均值差及95%点区间',loc='left');axs[1].set_xlabel('作者处理尺度；不是浓度倍数')
for ax,title in zip(axs[2:],['C  主分析 P / q','D  可用值 n、P / q']):ax.axis('off');ax.set_title(title,loc='left')
fig.suptitle('图1｜胆碱／乙醇胺相关代谢物：全22项保留',fontsize=21,y=.96,fontweight='bold');fig.text(.22,.91,'深蓝效应＝主分析；橙色＝可用值敏感性。统计沿用原全量家族；不可评估保留NA。',fontsize=11)
fig.text(.22,.025,'游离GPC、具体LPC和GPE类分开解释；包括未显著项。LPC16:0（1-palmitoyl-GPC）敏感性仅28对。',fontsize=11);save(fig,'Figure1_metabolites')
# Fig2 matrix with gene RNA and external coverage summaries, exact external rows remain in table.
primary=a[a.analysis_type.eq('processed_Spearman262')].set_index(['gene','metabolite_name']);rnaidx=rna.set_index('gene');mapped=set(zip(rel.gene,rel.metabolite_name))
mat=np.full((20,22),np.nan)
fig,ax=plt.subplots(figsize=(24,12));fig.subplots_adjust(left=.08,right=.99,top=.86,bottom=.27)
for i,g in enumerate(genes):
 for j,n in enumerate(names):
  if (g,n) in primary.index:mat[i,j]=num(primary.loc[(g,n),'effect'])
cm=plt.get_cmap('RdBu_r').copy();cm.set_bad('#EFEFEF');im=ax.imshow(mat,cmap=cm,vmin=-.6,vmax=.6,aspect='auto')
for i,g in enumerate(genes):
 for j,n in enumerate(names):
  if (g,n) not in mapped:ax.text(j,i,'·',ha='center',va='center',color='#9AA3AB')
  elif not np.isfinite(mat[i,j]):ax.text(j,i,'NA',ha='center',va='center',fontsize=7)
  else:
   r=primary.loc[(g,n)];tag='**' if num(r.q_value)<.05 else ('*' if num(r.p_value)<.05 else '')
   ax.text(j,i,f'{mat[i,j]:.2f}{tag}',ha='center',va='center',fontsize=7,color='white' if abs(mat[i,j])>.4 else 'black')
  if (g,n) in [('LYPLA1','1-palmitoyl-GPC (16:0)'),('GPCPD1','glycerophosphorylcholine (GPC)'),('ETNK1','ethanolamine'),('PCYT2','phosphoethanolamine')]:ax.add_patch(patches.Rectangle((j-.48,i-.48),.96,.96,fill=False,edgecolor='#E9A400',lw=2))
 r=rnaidx.loc[g] if g in rnaidx.index else None
 rt='NA' if r is None or r.status!='DONE' else f'{num(r.effect):+.2f}; ↑{int(num(r.positive_pairs))}/↓{int(num(r.negative_pairs))}; q={fmt(r.q_value)}'
 ax.text(22,i,rt,va='center',fontsize=8)
 for x,co in [(28,'FUSCC'),(32,'Tang')]:
  rows=rel[rel.gene.eq(g)];done=rows[rows[co+'_status'].eq('DONE')];pp=(pd.to_numeric(done[co+'_p_value'],errors='coerce')<.05).sum();qq=(pd.to_numeric(done[co+'_q_value'],errors='coerce')<.05).sum();ax.text(x,i,f'{len(done)}/{len(rows)}; {pp}/{qq}',va='center',fontsize=8)
ax.set_xlim(-.5,35);ax.set_ylim(19.5,-1.8);ax.set_yticks(range(20),genes);ax.set_xticks(range(22),labels,rotation=65,ha='right',fontsize=8);ax.text(22,-1,'RNA均值差；方向人数；原q',fontsize=10);ax.text(28,-1,'FUSCC',fontsize=10);ax.text(32,-1,'Tang',fontsize=10)
ax.set_xticks(np.arange(-.5,22,1),minor=True);ax.set_yticks(np.arange(-.5,20,1),minor=True);ax.grid(which='minor',color='white',lw=.6);ax.tick_params(which='minor',length=0)
fig.suptitle('图2｜全部41条直接关系 × 20基因；不按显著性删行',fontsize=21,y=.96,fontweight='bold');fig.text(.08,.91,'格内为CAMP主ρ；** 原q<0.05，* 仅P<0.05；· 无当前直接映射，NA 已映射不可评估。金框为预定四条重点。',fontsize=11)
fig.text(.08,.055,'外部栏＝可评估/该基因专题关系数；其中P<0.05数/q<0.05数。不是同向复现计数；逐关系效应、方向、状态见完整汇总表。',fontsize=11)
fig.text(.08,.025,'RNA效应为作者尺度的配对均值差；相关矩阵留空不等于生物学无关。敏感性及旧外部数值完整保留，不新增检验。',fontsize=11);save(fig,'Figure2_relations_RNA')
# Fig3 equal-donor expression and detection, separate study color scale.
types=list(dict.fromkeys(sc.celltype));cn={'B_cells':'B细胞','Endothelial':'内皮','Fibroblasts':'成纤维','Malignant_epithelial':'恶性上皮','Mast_cells':'肥大细胞','Myeloid':'髓系','Nonmalignant_epithelial':'非恶性上皮','Perivascular':'血管周','Plasma_cells':'浆细胞','T_NK_cells':'T/NK细胞','Unresolved_stromal':'未定基质'}
fig,axs=plt.subplots(1,2,figsize=(19,13));fig.subplots_adjust(left=.08,right=.95,top=.86,bottom=.18,wspace=.25)
for ax,co in zip(axs,['Wu2021','Pal2021_reprocessed']):
 d=sc[sc.cohort.eq(co)];idx=d.set_index(['gene','celltype']);vals=pd.to_numeric(d.loc[d.status.eq('DONE'),'effect'],errors='coerce');norm=colors.Normalize(0,max(vals.max(),.01))
 for i,g in enumerate(genes):
  for j,t in enumerate(types):
   if (g,t) in idx.index and idx.loc[(g,t),'status']=='DONE':
    r=idx.loc[(g,t)];ax.scatter(j,i,s=15+350*num(r.mean_detection_fraction),c=[num(r.effect)],cmap='viridis',norm=norm,edgecolor='#7D8790',linewidth=.3)
   else:ax.text(j,i,'×',color='#B7BEC4',ha='center',va='center',fontsize=8)
 ax.set_xticks(range(len(types)),[cn.get(t,t) for t in types],rotation=55,ha='right');ax.set_yticks(range(20),genes);ax.set_ylim(19.6,-.6);ax.set_xlim(-.6,len(types)-.4);ax.set_title(co,fontsize=15);ax.grid(axis='y',alpha=.12);fig.colorbar(plt.cm.ScalarMappable(norm=norm,cmap='viridis'),ax=ax,shrink=.65,label='供者等权平均 log1p(counts/库大小×10⁴)')
fig.suptitle('图3｜对应全20基因的单细胞表达背景',fontsize=21,y=.96,fontweight='bold');fig.text(.08,.91,'颜色＝供者等权平均表达（每研究独立色标）；点面积＝15+350×平均检出比例；×＝未测或覆盖不足。',fontsize=11)
fig.text(.08,.075,'LYPLA1、PCYT2两研究偏恶性上皮；GPCPD1偏髓系；ETNK1最高类别不一致，且排名不够稳定。',fontsize=11);fig.text(.08,.04,'表达来源不等于代谢物生成位置或唯一作用细胞。来源状态、有效供者数、保持率全部保留在source_status.tsv。',fontsize=11);save(fig,'Figure3_cell_sources')
caps={1:'主/敏感性均沿用原318项计划检验族的统计；效应为作者尺度均值差，不可跨分子解释绝对浓度差；灰条为相等。',2:'CAMP主关联ρ矩阵，未映射与不可评估分开；原检验族q不重算。外部统计计数不等于同向验证；所有准确关系结果在topic_summary41.tsv。',3:'按供者等权的均值与检出比例；两研究表达色标独立；Pal为既有再注释版本；不新增聚类或UMAP。'}
for i,c in caps.items():(O/f'Figure{i}_caption_CN.md').write_text(c,encoding='utf-8')
book=Workbook();book.remove(book.active)
for name,d in [('先读我',pd.DataFrame({'说明':['固定8952c54汇总；22代谢物/41关系/20基因；未新增P/q','GPE及plasmenylethanolamine另组，不冒称LPC；不含游离脂肪酸或酰基肉碱','未映射不是无关联；外部缺测不作阴性','只有已有直接关系：不新增映射或搜新队列']})),('专题范围',scope),('专题汇总',summary),('配对代谢物',m),('全部CAMP关联',a),('配对RNA',rna),('单细胞表达',sc),('来源状态',source)]:
 ws=book.create_sheet(name);ws.append(list(d.columns))
 for row in d.itertuples(index=False,name=None):ws.append([None if pd.isna(x) else ("'"+x if isinstance(x,str) and x.startswith(('=','+','@')) else x) for x in row])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='087F8C')
 for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=23
book.save(O/'专题汇总.xlsx')
pdf=fitz.open()
for n in ['Figure1_metabolites','Figure2_relations_RNA','Figure3_cell_sources']:
 d=fitz.open(O/(n+'.pdf'));pdf.insert_pdf(d);d.close()
pdf.save(O/'三图专题.pdf',garbage=4,deflate=True);pdf.close()
write(pd.DataFrame(manifest),'source_manifest')
spec=dict(input_commit=SHA,scope=names,relationships=41,genes=genes,new_tests=0,new_P_q=0,new_mapping=0,new_external_cohorts=0,cooccurrence='NOT_RUN',font='system Microsoft YaHei;not distributed',source_status_used=True)
(O/'analysis_spec.json').write_text(json.dumps(spec,ensure_ascii=False,indent=2),encoding='utf-8')
(O/'README_CN.md').write_text('''# 乳腺癌胆碱／乙醇胺相关代谢分化：首轮三图

## 问题与范围
胆碱/游离GPC低，乙醇胺/磷酸乙醇胺/部分具体LPC高；现象对应哪些基因和表达背景？固定8952c54，科学数值源于既有版本。明确22项范围，分小分子、9项具体LPC、7项乙醇胺相关脂质；最后一组不能统称LPC。不扩展游离脂肪酸、酰基肉碱。

## 实际结果与阅读
图1含全部22项主/敏感性结果，LPC16:0敏感性仅28对。图2含41条现有关系、20基因；金框四条重点不是唯一保留结果，RNA与外部结果分别列出。图3为两研究供者等权表达/检出比例，无重聚类。完整原数值见专题汇总及原列TSV。

## 解释与限制
胆碱/GPC与乙醇胺/磷酸乙醇胺方向不同；不是所有膜脂一起变化。LYPLA1、PCYT2偏恶性上皮，GPCPD1偏髓系，ETNK1来源不一致；不能把组织关联拼成单细胞内连续反应链或细胞间输送机制。PCYT2 RNA未显著且外部关联弱仍保留；GPCPD1外部不足，LPC精确匹配缺项，ETNK1来源不稳定，不隐藏反证。
没有新增筛选、P/q、外部队列、映射、共变人数或机制推断。现有262关系入口没有覆盖所有22项的每一种可能生化关系，无当前映射不是无生物学关系。

## 当前决定与复现
到三张图和专题表停止，供选择下一项具体问题，不锁定LYPLA1。python code/brca_choline_ethanolamine_v1.py；输出目录存在则拒绝覆盖。单细胞色标分研究，跨分子效应尺度不是浓度倍数；历史列原样保留。
''',encoding='utf-8')
(O/'.gitattributes').write_bytes(b'* -text\n')
(O/'validation.json').write_text(json.dumps(dict(features=22,relations=41,genes=20,association_rows=len(a),sc_rows=len(sc),visual_review='PENDING',new_P_q=0),indent=2),encoding='utf-8')
write(pd.DataFrame([dict(path=p.name,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in sorted(O.iterdir())]),'checksums')
print(O);print(scope[['display_label','direct_relation_count']].to_string(index=False))
