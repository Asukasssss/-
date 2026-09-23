"""Render only public donor-level test summaries."""
from pathlib import Path
import json,hashlib
import numpy as np,pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
R=Path(__file__).resolve().parents[1];O=R/'results/BRCA/06_EXTERNAL/20260923T040012Z_epithelial_paired20_v1'
d=pd.read_csv(O/'results.tsv',sep='\t');assert len(d)==60
for fam,g in d.groupby('test_family'):
    assert len(g)==20
    ps=g.p_value.fillna(1).values;order=np.argsort(ps);q=np.empty(20)
    q[order]=np.minimum(1,np.minimum.accumulate((ps[order]*20/np.arange(1,21))[::-1])[::-1])
    ok=g.status.eq('DONE').values
    assert np.allclose(q[ok],g.q_value.values[ok])
    assert (g.positive_pairs+g.negative_pairs+g.equal_pairs==g.n).all()
x=d[d.cohort.eq('Wu2021')&d.analysis_type.eq('ALL')].copy()
assert (x.n==8).all() and x.p_value.lt(.05).sum()==5 and x.q_value.lt(.05).sum()==0
font_manager.fontManager.addfont('C:/Windows/Fonts/msyh.ttc');plt.rcParams.update({'font.family':'Microsoft YaHei','axes.unicode_minus':False,'pdf.fonttype':42})
fig,axs=plt.subplots(1,2,figsize=(14,10),gridspec_kw={'width_ratios':[1.2,1]});fig.subplots_adjust(left=.12,right=.98,top=.84,bottom=.20,wspace=.10)
for i,r in enumerate(x.itertuples()):
    if r.status=='DONE':
        axs[0].errorbar(r.effect,i,xerr=[[r.effect-r.ci_lower],[r.ci_upper-r.effect]],fmt='o',color='#C16A2D' if r.p_value<.05 else '#46778A',capsize=3)
        label=f'{int(r.positive_pairs)}/{int(r.n)}      {r.p_value:.4f}      {r.q_value:.4f}'
    else: label='低检出：不进行推断'
    axs[1].text(0,i,label,va='center',fontsize=11)
for a in axs:a.set_ylim(len(x)-.4,-1);a.spines[['top','right']].set_visible(False)
axs[0].set_yticks(range(len(x)),x.gene);axs[0].axvline(0,c='gray',ls='--');axs[0].set_xlabel('恶性 − 非恶性：供者均值差及95%点区间\nmean log1p(count/库大小×10⁴)，不是倍数')
axs[1].axis('off');axs[1].set_title('恶性较高人数      双侧P       BH20 q',fontsize=12,loc='left')
fig.suptitle('Wu｜20基因的供者内上皮表达比较',fontsize=20,y=.97)
fig.text(.12,.90,'8位共同供者；每位每类≥20细胞。5项P<0.05，0项BH20 q<0.05。',fontsize=12)
fig.text(.12,.055,'橙色＝原始P<0.05；灰虚线＝无均值差。区间未经多重校正，与秩检验回答的统计问题不同。\n参照是肿瘤标本内非恶性上皮，不是健康组织。Pal当前缺少参照，不能复现；已看过表达图，属于探索性检验。',fontsize=10)
for ext in ['png','pdf','svg']:fig.savefig(O/('donor_paired20.'+ext),dpi=170,bbox_inches='tight')
plt.close(fig)
w=Workbook();w.remove(w.active)
for name,t in [('全部60项',d),('覆盖情况',pd.read_csv(O/'coverage.tsv',sep='\t')),('结果计数',pd.read_csv(O/'summary_counts.tsv',sep='\t'))]:
    s=w.create_sheet(name);s.append(list(t.columns))
    for row in t.itertuples(index=False,name=None):s.append([None if pd.isna(v) else v for v in row])
    s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
    for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='226B7A')
    for c in s.columns:s.column_dimensions[c[0].column_letter].width=22
w.save(O/'供者配对表达结果.xlsx')
readme=(O/'README_CN.md').read_text(encoding='utf-8').split('\n## 本轮数值摘要')[0]
readme+='\n## 本轮数值摘要\nWu主分析20项，18项可检验，PCYT1B和SLC22A2因低检出不推断；5项原始P<0.05，0项BH20 q<0.05。ABHD12、LPCAT1、LYPLA2为8/8位较高，P=0.0078125，q=0.0520833；PCYT2为7/8，P=0.015625，q=0.078125；LYPLA1为7/8，P=0.0390625，q=0.15625。统一保留计划20项，未检验项仅校正内部以P=1占位；不改成仅校正18项来寻求过线。作者Naive7位敏感性也无FDR支持。\n\n## 解读修正\n两研究中最高类别为恶性上皮，不等于同供者恶性高于非恶性。ABHD12总体最高为髓系，也可在上皮内部呈现恶性高于非恶性；比较对象不同，不矛盾。当前不能宣称任何基因已通过本轮FDR证明恶性特异升高，更不能据此判断正常组织安全性。\n'
(O/'README_CN.md').write_text(readme,encoding='utf-8')
v=json.loads((O/'validation.json').read_text());v.update(BH20_independently_rechecked=True,direction_counts_checked=True,main_p_lt_005=5,main_q_lt_005=0)
(O/'validation.json').write_text(json.dumps(v,indent=2),encoding='utf-8')
(O/'.gitattributes').write_bytes(b'* -text\n')
pd.DataFrame([dict(path=str(p.relative_to(R)),sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in [R/'code/brca_epithelial_paired20_v1.py',Path(__file__)]]).to_csv(O/'code_manifest.tsv',sep='\t',index=False)
print(x[['gene','effect','p_value','q_value']].to_string(index=False))
