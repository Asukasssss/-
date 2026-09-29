"""Insert imagegen cutouts into three role panels; preserve all result objects."""
from pathlib import Path
import ast,json,shutil,hashlib
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
ROOT=Path(__file__).resolve().parents[1];RUN='20260929T_case_art_v4';O=ROOT/'results/BRCA/07_INTEGRATION'/RUN;A=O/'assets';A.mkdir(parents=True,exist_ok=True)
V=Path('D:/CodexData/visualizations/2026/09/29/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/CAMP_case_art_v4');V.mkdir(exist_ok=True)
OLD=ROOT/'results/BRCA/07_INTEGRATION/20260929T_complete_cases_v3';INPUT=OLD/'CAMP_完整组会汇报_含三个案例_25页.pptx'
SOURCE=Path('D:/CodexData/generated_images/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/exec-b3f59589-1a5b-4e21-8b27-83fc450c7fd3.png')
if not (A/'scientific_cutouts.png').exists():shutil.copy2(SOURCE,A/'scientific_cutouts.png')
R=Presentation(INPUT)
mod=ast.parse((ROOT/'code/ppt_editable_main_v1.py').read_text(encoding='utf-8'));ns=dict(G='00553B',INK='183D34',M='6A7973',PALE='F0F5F1',LINE='DDE7E0',FONT='Microsoft YaHei')
exec(compile(ast.Module(body=[n for n in mod.body if isinstance(n,(ast.Import,ast.ImportFrom,ast.FunctionDef))],type_ignores=[]),'<helpers>','exec'),ns)
txt,box,arrow,el=[ns[x] for x in ['txt','box','arrow','el']];CENTER=ns['PP_ALIGN'].CENTER;G='00553B';M='6A7973'
rects={'nucleoside':(65,40,450,470),'nucleotide':(520,45,995,470),'enzyme':(1045,20,1525,490),'metabolites':(12,550,510,965),'transporter':(530,530,1020,980),'atp':(1035,540,1536,955)}
def art(s,key,x,y,w,h):
 p=A/'scientific_cutouts.png';iw,ih=Image.open(p).size;l,t,r,b=rects[key];sw,sh=r-l,b-t;f=min(w/sw,h/sh);dw,dh=sw*f,sh*f
 z=s.shapes.add_picture(str(p),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh));z.crop_left=l/iw;z.crop_right=1-r/iw;z.crop_top=t/ih;z.crop_bottom=1-b/ih;z.name='Imagegen conceptual cutout: '+key
def flow(s,x,y,w):
 z=ns['line'](s,x,y+.075,x+w,y+.075,G,1.8);z._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(el('a:tailEnd',type='triangle',w='med',len='med'))
def label(s,x,y,w,t,size=19,c=G,b=True):return txt(s,x,y,w,.34,t,size,c,b,CENTER)
for n,gene in [(19,'ASNS'),(21,'SLC6A6'),(24,'UCKL1')]:
 s=R.slides[n-1]
 for sh in list(s.shapes):
  if 1.98<=sh.top/Inches(1)<3.70:s.shapes._spTree.remove(sh._element)
 box(s,.66,2.02,12.0,1.66,'F3F7F3')
 if gene=='UCKL1':
  art(s,'nucleoside',1.75,2.11,1.0,1.04);art(s,'nucleotide',9.10,2.11,1.15,1.04);art(s,'enzyme',6.15,2.08,.83,.68)
  label(s,2.94,2.44,2.0,'尿苷',22);label(s,10.1,2.44,1.7,'UMP',22)
  flow(s,5.03,2.99,3.22)
  label(s,7.00,2.32,1.10,'UCKL1',14);label(s,5.17,3.17,2.95,'ATP → ADP',13,M,False)
  label(s,1.6,3.19,3.48,'核苷底物',12,M,False);label(s,9.13,3.19,2.83,'增加一个磷酸基团',12,M,False)
  label(s,1.30,3.43,10.85,'尿苷 → UMP：进入嘧啶核苷补救合成；也可催化胞苷 → CMP。',10.4,M,False)
 elif gene=='ASNS':
  art(s,'metabolites',1.05,2.13,1.32,1.06);art(s,'enzyme',6.15,2.08,.83,.68);art(s,'metabolites',10.77,2.13,1.32,1.06)
  label(s,2.44,2.46,2.71,'谷氨酰胺 + 天冬氨酸',16);label(s,8.14,2.46,2.63,'谷氨酸 + 天冬酰胺',16)
  flow(s,5.36,2.99,2.62);label(s,7.00,2.32,1.10,'ASNS',14);label(s,5.1,3.17,3.2,'ATP → AMP + PPi',12,M,False)
  label(s,1.32,3.19,3.79,'谷氨酰胺提供酰胺氮',12,M,False);label(s,8.14,3.19,3.86,'产物为天冬酰胺',12,M,False)
  label(s,1.2,3.43,10.95,'利用谷氨酰胺的氮合成天冬酰胺；分子与蛋白均为概念示意。',10.4,M,False)
 else:
  art(s,'transporter',5.35,2.04,2.38,1.20)
  label(s,1.52,2.12,2.75,'细胞外',12,M,False);label(s,9.17,2.12,2.75,'细胞内',12,M,False)
  label(s,1.52,2.53,2.75,'牛磺酸',22);label(s,9.17,2.53,2.75,'牛磺酸',22)
  flow(s,4.02,2.71,1.17);flow(s,7.97,2.71,1.16)
  label(s,5.02,3.25,3.22,'SLC6A6 / TauT',14)
  label(s,1.45,3.43,10.49,'Na+ / Cl− 依赖性摄取；将细胞外牛磺酸转运入细胞，不改变化学结构。',10.4,M,False)
 for sh in s.shapes:
  for e in sh._element.xpath('.//a:effectRef'):e.set('idx','0')
  for rp in sh._element.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr'):
   rp.set('lang','zh-CN')
   for tag,face in [('a:latin','Microsoft YaHei'),('a:ea','微软雅黑'),('a:cs','Microsoft YaHei')]:
    old=rp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag.split(':')[1])
    if old is not None:rp.remove(old)
    rp.append(el(tag,typeface=face))
 s.notes_slide.notes_text_frame.text+='\n插图由内置image_gen生成，仅为概念图，不是实际分子结构、蛋白结构或本队列机制证据。图片保留透明度并使用PPT原生裁剪，文字与箭头独立可编辑。'
PPT='CAMP_完整组会汇报_生化插图精修_25页.pptx';R.save(O/PPT);shutil.copy2(O/PPT,V/PPT)
(O/'asset_manifest.json').write_text(json.dumps(dict(input=str(INPUT),input_sha256=hashlib.sha256(INPUT.read_bytes()).hexdigest(),asset=str(SOURCE),asset_sha256=hashlib.sha256((A/'scientific_cutouts.png').read_bytes()).hexdigest(),edited_slides=[19,21,24],crop_rectangles=rects,imagegen_mode='built-in',conceptual_only=True),ensure_ascii=False,indent=2),encoding='utf-8')
print(O/PPT)
