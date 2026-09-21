"""Aggregate-only evidence joins; preserve every historical column and value."""
from pathlib import Path
import json,hashlib,sys
import pandas as pd,numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill,Font,Alignment
from openpyxl.utils import get_column_letter
repo=Path(__file__).resolve().parents[1];run='20260921T154000Z_mapping262_integration_v1';R=repo/'results/BRCA/07_INTEGRATION'/run;R.mkdir(parents=True,exist_ok=True)
mp=repo/'results/BRCA/02_MAPPING/20260921T124136Z_mapping190_v2';cp=repo/'results/BRCA/03_PATIENT/20260921T150000Z_mapping262_patient_v1';ep=repo/'results/BRCA/06_EXTERNAL/20260921T151000Z_mapping262_external_v1';fp=repo/'results/BRCA/05_FUNCTION/20260921T152000Z_new39_literature_v1';oldp=repo/'results/BRCA/04_ROBUSTNESS/20260921T103911Z_camp_pair_sensitivity_v1/all117_comparison_identity_appended.tsv'
def read(p):return pd.read_csv(p,sep='\t')
def save(d,p):d.to_csv(p,sep='\t',index=False,na_rep='NA',lineterminator='\n')
rel=read(mp/'direct_relations.tsv');met=read(mp/'metabolite190_mapping_status.tsv');camp=read(cp/'CAMP_relations262.tsv');rna=read(cp/'paired_RNA150.tsv');ext=read(ep/'external_relations262.tsv');fun=read(fp/'new39_functional_summary.tsv');members=read(mp/'gene_old_new_comparison.tsv');old=pd.read_csv(oldp,sep='\t',dtype=str,keep_default_na=False)
assert len(old)==117 and old.gene.is_unique and len(rel)==262 and rel.relation_key.is_unique
out=rel.copy();out['stage_id']='07_INTEGRATION';out['run_id']=run;out['analysis_version']='mapping262_integration_v1'
# Carry original metabolite paired evidence without recomputing or changing its multiplicity.
cols=['metabolite_key','n','effect','p_value','q_value','pairs_higher','pairs_lower','pairs_equal','same_direction_fraction','majority_direction','q_value_available','n_available']
out=out.merge(met[cols].rename(columns={c:'metabolite_paired_'+c for c in cols if c!='metabolite_key'}),on='metabolite_key',validate='many_to_one')
statscols=['relation_id','n','effect','ci_lower','ci_upper','p_value','q_value','status','reason','statistics_reused','previous_q','family_n_evaluable']
families=[(camp,'processed_Spearman262','CAMP'),(camp,'available_Spearman262','CAMP_available'),(ext,'FUSCC_PRIMARY262','FUSCC'),(ext,'FUSCC_EXCLUDE_WARNING262','FUSCC_warning_excluded'),(ext,'Tang262','Tang')]
for data,f,prefix in families:
 a=data[data.test_family.eq(f)][statscols].rename(columns={'relation_id':'relation_key',**{c:prefix+'_'+c for c in statscols if c!='relation_id'}});out=out.merge(a,on='relation_key',validate='one_to_one');assert len(out)==262
for prefix in ['FUSCC','Tang']:
 ok=out.CAMP_effect.notna()&out[prefix+'_effect'].notna();out[prefix+'_direction_vs_CAMP']=np.where(ok,np.where(out.CAMP_effect*out[prefix+'_effect']>0,'SAME_SIGN','OPPOSITE_OR_ZERO'),'NOT_EVALUABLE')
out['new_relation_statistics_not_transferred_by_gene']=True;save(out,R/'relations262_comparison.tsv')
# Preserve exact original117 text values; no rewriting historical NA strings or formatted P/q.
union=old.set_index('gene').reindex(sorted(members.gene)).reset_index();union=union.merge(members.rename(columns={c:'v2_'+c for c in members if c!='gene'}),on='gene',validate='one_to_one')
assert union.set_index('gene').loc[old.gene,old.columns.drop('gene')].equals(old.set_index('gene')[old.columns.drop('gene')])
rc=['gene','n','effect','ci_lower','ci_upper','p_value','q_value','status','reason','positive_pairs','negative_pairs','equal_pairs'];union=union.merge(rna[rc].rename(columns={c:'v2_paired_RNA_'+c for c in rc if c!='gene'}),on='gene',how='left',validate='one_to_one')
union=union.merge(fun.rename(columns={c:'v2_new39_'+c for c in fun if c!='gene'}),on='gene',how='left',validate='one_to_one')
summ=[]
for g in sorted(members.gene):
 a=out[out.gene.eq(g)];b=dict(gene=g,v2_n_direct_relations=len(a),v2_relation_keys=';'.join(a.relation_key),v2_current_scope='CURRENT150' if len(a) else 'HISTORICAL6_ONLY')
 for pre in ['CAMP','CAMP_available','FUSCC','Tang']:
  b['v2_'+pre+'_n_evaluable']=int(a[pre+'_p_value'].notna().sum()) if len(a) else np.nan;b['v2_'+pre+'_q_lt005_relations']=';'.join(a.loc[a[pre+'_q_value'].lt(.05),'relation_key']) if len(a) else 'HISTORICAL_ONLY'
  b['v2_'+pre+'_n_access_blocked']=int(a[pre+'_status'].eq('ACCESS_BLOCKED').sum()) if len(a) else np.nan
 if g=='LPCAT4':decision='身份暂挂：稳定ID核实前不解释RNA或文献'
 elif not len(a):decision='历史保留：不在当前190入口，已有结果不删除'
 elif a.CAMP_q_value.lt(.05).any() or a.FUSCC_q_value.lt(.05).any() or a.Tang_q_value.lt(.05).any():decision='保留患者关联线索；逐关系比较效应、缺失及外部支持'
 else:decision='保留候选；患者未显著或缺测不取消功能证据'
 b['v2_current_decision_cn']=decision;summ.append(b)
union=union.merge(pd.DataFrame(summ),on='gene',validate='one_to_one');assert len(union)==156 and union.gene.is_unique
assert union.set_index('gene').loc[old.gene,old.columns.drop('gene')].equals(old.set_index('gene')[old.columns.drop('gene')])
hist=union.v2_current_scope.eq('HISTORICAL6_ONLY');union.loc[hist,'v2_paired_RNA_status']='NOT_IN_CURRENT_POOL_HISTORY_RETAINED';assert set(union.loc[hist,'gene'])=={'ACLY','CROT','CS','MAT1A','MAT2A','PAH'}
save(union,R/'genes156_all_history_comparison.tsv')
# Readable 156-row front page, without hiding disagreements or forcing one scalar score.
front=[]
for r in union.to_dict('records'):
 g=r['gene'];a=out[out.gene.eq(g)];n39=fun[fun.gene.eq(g)]
 front.append({'基因':g,'候选范围':r['v2_current_scope'],'是否新增':not bool(r['v2_in_original117']),'本轮直接关系数':len(a),'对应代谢物':';'.join(a.metabolite_name.str.strip().unique()),'配对RNA效应_作者尺度':r.get('v2_paired_RNA_effect'),'配对RNA新版q':r.get('v2_paired_RNA_q_value'),'肿瘤RNA较高对数':r.get('v2_paired_RNA_positive_pairs'),'肿瘤RNA较低对数':r.get('v2_paired_RNA_negative_pairs'),'CAMP本轮q小于005关系':r['v2_CAMP_q_lt005_relations'],'FUSCC本轮q小于005关系':r['v2_FUSCC_q_lt005_relations'],'Tang本轮q小于005关系':r['v2_Tang_q_lt005_relations'],'FUSCC下载受阻关系数':r['v2_FUSCC_n_access_blocked'],'功能证据':n39.iloc[0].functional_evidence_class if len(n39) else r.get('review_v4_functional_class','历史表保留'),'功能限制':n39.iloc[0].model_and_limits_cn if len(n39) else r.get('review_v4_counterevidence_and_limits',''),'当前决定':r['v2_current_decision_cn']})
front=pd.DataFrame(front);save(front,R/'genes156_reader_cn.tsv')
newrels=out[~out.CAMP_statistics_reused & out.CAMP_p_value.notna()];save(newrels,R/'new_patient_relations96.tsv')
overview=pd.DataFrame([{'项目':'当前入口','结果':'190项配对P<0.05；保留原q；不是190项均多重校正显著'}, {'项目':'映射','结果':'93项代谢物、262条关系、当前150基因、完整历史并集156'}, {'项目':'CAMP','结果':'60肿瘤病例；255条可评估；主分析14条、可用值9条新版q<0.05'}, {'项目':'配对RNA','结果':'45对；148基因可评估，79新版q<0.05；新增38可评估，其中22显著'}, {'项目':'FUSCC','结果':'新增RNA接口受阻；126条可评估；主分析13条新版q<0.05；阶段性结果'}, {'项目':'Tang','结果':'20例；235条可评估；MDH1—苹果酸1条新版q<0.05'}, {'项目':'功能','结果':'39基因78次检索；26条摘录涉及24基因；包括反证/模型待核对，不是24个已验证靶点'}, {'项目':'研究资源','结果':'ACSL4 GSE283282已核元数据，尚未下载计数或分析'}, {'项目':'阅读顺序','结果':'先156基因简表，再262关系对照，最后按来源查看功能证据及全部历史列'}, {'项目':'空值','结果':'未测/未算/不适用，不是零或阴性'}, {'项目':'排序','结果':'基因字母顺序；不是统一排名'}])
xlsx=R/'BRCA_156基因与262关系_新版比较.xlsx'
with pd.ExcelWriter(xlsx,engine='openpyxl') as w:
 for name,d in [('说明',overview),('156基因简表',front),('262关系对照',out),('150基因配对RNA',rna),('新增39功能',fun),('功能证据26条',read(fp/'functional_evidence.tsv')),('干预资源',read(fp/'intervention_resource_inventory.tsv')),('156基因全历史',union)]:d.to_excel(w,sheet_name=name,index=False)
wb=load_workbook(xlsx)
for ws in wb:
 ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions;ws.row_dimensions[1].height=34
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
 for j,c in enumerate(ws[1],1):ws.column_dimensions[get_column_letter(j)].width=min(45,max(14,len(str(c.value))*1.25))
wb.save(xlsx);assert load_workbook(xlsx,read_only=True)['156基因简表'].max_row==157
spec=dict(version='mapping262_integration_v1',source_commits=['9e4387f','809ef3c','05b937d','ba26058'],new_statistics=0,join_keys='exact metabolite_key|gene;gene-only joins restricted to gene-level evidence',historical_columns_preserved=len(old.columns),history_only6=sorted(union.loc[hist,'gene']),no_composite_q_or_score=True)
(R/'analysis_spec.json').write_text(json.dumps(spec,indent=2));val=dict(status='DONE',integration_only=True,relations=262,genes_union=156,current_genes=150,history_genes=117,new_genes=39,historical249_columns_exactly_preserved=True,unique_keys_checked=True,missing_values_not_zero=True,FUSCC_status='PARTIAL_API_ACCESS_BLOCKED',function_status='BOUNDED_FIRST_PASS_NOT_EXHAUSTIVE',new_patient_relations=96,Excel_rows_verified=True)
(R/'validation.json').write_text(json.dumps(val,indent=2));paths=[mp/'direct_relations.tsv',mp/'metabolite190_mapping_status.tsv',mp/'gene_old_new_comparison.tsv',cp/'CAMP_relations262.tsv',cp/'paired_RNA150.tsv',ep/'external_relations262.tsv',fp/'new39_functional_summary.tsv',fp/'functional_evidence.tsv',oldp,Path(__file__)];save(pd.DataFrame([dict(path=str(p.relative_to(repo)).replace('\\','/'),sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in paths]),R/'source_manifest.tsv')
text='''# BRCA新版候选比较：156基因、262关系

## 本轮问题
让新增39个基因获得患者计算和功能取证机会，同时保留旧117基因全部研究记录。没有把名单缩为少数显著项。

## 输入与范围
当前150基因/262直接关系，来自190项45对主分析P<0.05代谢物入口。加历史保留6基因，形成156行。历史117基因249列原值全部保留；只有新增列用于新版判断。

## 实际结果
|层次|本轮结果|
|---|---|
|CAMP患者关联|255/262可评估；主分析14条、可用值敏感性9条新版q<0.05。各复用159、计算96条|
|配对RNA|148/150可评估；复用110、新算38；79个新版q<0.05，其中新增基因22个|
|FUSCC|126/262暂可评估；13条新版q<0.05。新增RNA接口拒绝连接，不能视为阴性|
|Tang|235/262可评估；复用153、新算82；MDH1—苹果酸1条新版q<0.05|
|功能取证|39基因78次查询；26条摘录涉及24基因，含反证与待核模型，不是24个已验证靶点|

新增患者关系值得注意：LYPLA1—1-palmitoyl-GPC主分析ρ=0.482、q=0.017，可用值ρ=0.518、q=0.0255；ENPP2—1-stearoyl-GPC主分析ρ=-0.400、q=0.0371，但可用值q=0.1509。LYPLA1—油酸主分析q=0.0455，可用值q=0.0510，效应不变。不得只展示过线版本。

## 新手解释
配对代谢物差异是发现入口；RNA差异、肿瘤内相关、外部相关、模型干预分别提供不同证据。旧关系P值没改，但新版q随范围改变，因此不能把q变化当作新实验结果。完整表按字母排序，不是靶点排名。

## 限制与反证
LPCAT4历史名称歧义未解决，CHKB缺RNA。FUSCC新版覆盖受接口限制；公开表将受阻和缺测分开。Tang仅20例，新版未显著不能删除原线索。所有候选来自同CAMP选择入口，CAMP补算不是独立验证。
功能方面，ENPP2脂肪细胞敲除和全身药理结果不同；ACSL3/SDHC/SDHD不能统一建议抑制；SUCLA2在脱附条件下的功能并未被琥珀酸改变解释，SUCLG1/G2也不能按复合体成员机械继承功能。

## 当前决定
LYPLA1/ENPP2等作为新增线索进入讨论；原ASNS/GLS/SLC6A8/GPCPD1等资料继续保留。不存在“CAMP显著且外部显著且RNA显著才保留”的硬交集。条件性18关系不混入统计。

## 下一步
首选补FUSCC新增RNA缺口，并核对LPCAT4探针稳定身份。若选择新的干预分析，ACSL4已有GSE283282明确入口，下一轮只需先核对处理后计数、三组标签和作者批次独立性；本轮不自动启动更多机制分析。

## 复现
运行code/brca_mapping262_integrate_v1.py。输入是已公开汇总表，无患者矩阵。本脚本验证117历史249列逐值保留、156/262唯一键及Excel行数。上游代码、范围、来源、q家族见各阶段analysis_spec与source_manifest。所有患者级源值仍留server165。
'''
(R/'README_CN.md').write_text(text,encoding='utf-8');(R/'.gitattributes').write_text('* -text\n');print(json.dumps(val))
