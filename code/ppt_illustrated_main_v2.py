"""Hybrid editable PPT: native text and diagrams with isolated scientific artwork."""
from pathlib import Path
import ast,shutil,json,hashlib
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN
ROOT=Path(__file__).resolve().parents[1];BASE=Path('D:/CodexData/visualizations/2026/09/25/01a0d8cd-4545-75b0-ba43-922588deea1a');RUN='20260929T_illustrated_main_v2';O=ROOT/'results/BRCA/07_INTEGRATION'/RUN;A=O/'assets';A.mkdir(parents=True,exist_ok=True)
OLD=ROOT/'results/BRCA/07_INTEGRATION/20260929T_editable_main_v1';V=Path('D:/CodexData/visualizations/2026/09/29/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/CAMP_main_illustrated_v2');V.mkdir(parents=True,exist_ok=True)
files={'artwork.png':Path('D:/CodexData/generated_images/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/exec-c8651984-ddc3-44ad-a602-a36153525bfb.png'),'cover_reference.png':BASE/'CAMP_IMAGE_fivepages_v3/01_标题封面.png','matching_reference.png':BASE/'CAMP_IMAGE_fivepages_v3/03_CAMP优势与患者匹配.png','paper_real.png':BASE/'CAMP_IMAGE_fivepages_v3/paper_header_authors_real.png','sysu_logo.png':BASE/'CAMP_SYSU_style/image2.png'}
for name,p in files.items():shutil.copy2(p,A/name)
# Reuse only drawing helpers; do not execute or overwrite the prior deck.
source=(ROOT/'code/ppt_editable_main_v1.py').read_text(encoding='utf-8');module=ast.parse(source);wanted={'rgb','alpha','box','txt','line','arrow','notes','el'}
ns=dict(G='00553B',INK='183D34',M='6A7973',PALE='F0F5F1',LINE='DDE7E0',FONT='Microsoft YaHei');imports=[n for n in module.body if isinstance(n,(ast.Import,ast.ImportFrom))];defs=[n for n in module.body if isinstance(n,ast.FunctionDef) and n.name in wanted]
exec(compile(ast.Module(body=imports+defs,type_ignores=[]),'<drawing helpers>','exec'),ns)
ns.update(G='00553B',INK='183D34',M='6A7973',PALE='F0F5F1',LINE='DDE7E0',FONT='Microsoft YaHei')
rgb,box,txt,line,arrow,notes,el=[ns[k] for k in ['rgb','box','txt','line','arrow','notes','el']]
G='00553B';M='6A7973';INK='183D34';R=Presentation(OLD/'CAMP_泛癌发现与LYPLA1_可编辑汇报.pptx')
rects={'patient':(12,20,555,550),'molecule':(615,130,1008,490),'dna':(1110,60,1480,585),'enzyme':(10,587,498,1010),'transporter':(548,573,1025,1012),'cells':(1032,618,1532,1012)}
organs={'BRCA':(1285,275,1380,370),'COAD':(1470,275,1590,370),'GBM':(1280,425,1400,525),'PDAC':(1470,429,1600,515),'PRAD':(1300,575,1390,656),'ccRCC':(1494,575,1576,657)}
def crop(sl,path,bbox,x,y,w,h):
 iw,ih=Image.open(path).size;l,t,r,b=bbox;sw,sh=r-l,b-t;scale=min(w/sw,h/sh);dw,dh=sw*scale,sh*scale
 p=sl.shapes.add_picture(str(path),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh));p.crop_left=l/iw;p.crop_right=1-r/iw;p.crop_top=t/ih;p.crop_bottom=1-b/ih;p.name='Illustration crop: '+path.name;return p
def art(sl,key,x,y,w,h):return crop(sl,A/'artwork.png',rects[key],x,y,w,h)
def organ(sl,key,x,y,w,h):return crop(sl,A/'matching_reference.png',organs[key],x,y,w,h)
def new(n,title,sub):
 s=R.slides.add_slide(R.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('FFFFFF')
 s.shapes.add_picture(str(A/'sysu_logo.png'),Inches(.66),Inches(.13),width=Inches(1.93),height=Inches(.75));txt(s,10.1,.42,2.54,.26,'CAMP  /  组会汇报',11,M,align=PP_ALIGN.RIGHT)
 line(s,.65,.94,12.67,.94);line(s,.65,.94,1.40,.94,G,2);txt(s,.67,1.17,12,.58,title,27,G,True);txt(s,.68,1.85,11.9,.48,sub,13,M);txt(s,12.12,7.03,.52,.27,f'{n:02d}',11,G,True,PP_ALIGN.RIGHT);return s
def foot(s,t):line(s,.67,6.89,12.67,6.89);txt(s,.68,7.03,11.1,.24,t,9,M)
def tag(s,x,y,w,text):box(s,x,y,w,.33,'E7EFE8');txt(s,x+.08,y+.055,w-.16,.24,text,11,G,True)

# Cover: exact lower landscape from the approved image, editable title above.
s=new(1,'','')
for sh in list(s.shapes):
 if sh.has_text_frame and sh.text=='CAMP  /  组会汇报':s.shapes._spTree.remove(sh._element)
crop(s,A/'cover_reference.png',(0,689,1672,941),0,5.49,13.333,2.01)
crop(s,A/'cover_reference.png',(1190,0,1672,228),9.57,0,3.77,1.79)
tag(s,.74,1.65,2.12,'CAMP · 多组学研究')
txt(s,.74,2.33,11.7,1.45,'基于 CAMP 多组学的\n泛癌潜在代谢靶点筛选',35,G,True)
txt(s,.79,4.08,10.9,.56,'以 LYPLA1 为重点的跨癌种证据分析',23,INK)
txt(s,.80,4.90,10.7,.39,'患者匹配  /  代谢物—基因关系  /  细胞与空间背景',14,M)
notes(s,'插图背景沿用用户认可的IMAGE封面下方校景意象及叶片，为装饰性插画，不是真实校园摄影。全部标题、副标题和标签为可编辑文字。')

# Resource page: real paper is the evidence; illustrations only explain matched modalities.
s=new(2,'CAMP 是什么？','Cancer Atlas of Metabolic Profiles · 将代谢组与转录组放到同一研究框架')
box(s,.68,2.49,5.46,3.99,'F3F7F4');box(s,.68,2.49,.045,3.99,G)
crop(s,A/'paper_real.png',(0,0,*Image.open(A/'paper_real.png').size),.87,2.73,5.06,2.42)
txt(s,.92,5.45,4.97,.38,'Nature Metabolism · 2023',19,G,True);txt(s,.94,5.99,4.95,.30,'Benedetti et al. · 原始论文截图',12,M)
txt(s,6.58,2.53,5.74,.98,'整合肿瘤及参照组织的两种组学，\n研究代谢物与基因表达的共同变化。',20,INK)
for x,num,lab in [(6.62,'11','类癌症'),(8.62,'15','个数据集'),(10.65,'988','份组织标本')]:
 txt(s,x,3.87,1.6,.58,num,32,G,True);txt(s,x,4.52,1.7,.31,lab,13,M)
art(s,'molecule',6.75,5.08,1.55,1.30);art(s,'dna',10.77,5.02,1.12,1.40);line(s,8.40,5.72,10.48,5.72,'89AB96',1.5)
txt(s,8.44,5.14,2.14,.32,'对应标本匹配',15,G,True,PP_ALIGN.CENTER);txt(s,8.42,5.94,2.19,.31,'代谢水平 ↔ 基因表达',12,M,align=PP_ALIGN.CENTER)
foot(s,'来源：Nature Metabolism 5, 1029–1044 (2023) · doi:10.1038/s42255-023-00817-8')
notes(s,'CAMP包含11癌种15数据集988组织标本。组学标本匹配与同患者肿瘤—正常配对分开核对。左侧为原文截图，右侧分子与DNA仅为组学示意，不表示数据或特定分子结构。')

# Three-column matching story with extracted scientific figures.
s=new(3,'为什么选择 CAMP？','患者内配对减少个体差异；对应标本匹配连接代谢物与基因')
xs=[.68,4.89,9.10];ww=[3.55,3.55,3.55]
for x,w,n,t in zip(xs,ww,['01','02','03'],['患者内组织配对','对应标本组学匹配','跨癌种覆盖']):
 box(s,x,2.57,w,3.93,'F8FAF8');box(s,x,2.57,w,.58,'E7EFE8');txt(s,x+.15,2.70,.42,.3,n,18,G,True);txt(s,x+.75,2.72,w-.9,.3,t,17,G,True)
art(s,'patient',.83,3.30,3.19,2.28);txt(s,1.88,5.58,2.15,.30,'肿瘤组织 / 参照组织',12,M,align=PP_ALIGN.CENTER)
txt(s,.94,6.00,3.02,.32,'以同一患者为自身参照',15,G,True,PP_ALIGN.CENTER)
crop(s,A/'matching_reference.png',(820,328,991,439),6.02,3.35,1.33,.90)
line(s,6.69,4.24,5.83,4.58,G,1.2);line(s,6.69,4.24,7.56,4.58,G,1.2)
art(s,'molecule',5.19,4.55,1.34,1.11);art(s,'dna',7.11,4.45,.93,1.32)
txt(s,5.08,5.65,1.53,.3,'代谢组',12,G,True,PP_ALIGN.CENTER);txt(s,6.92,5.65,1.50,.3,'RNA',12,G,True,PP_ALIGN.CENTER)
txt(s,5.09,6.00,3.15,.32,'匹配肿瘤标本中的实测关联',14,G,True,PP_ALIGN.CENTER)
for k,g in enumerate(['BRCA','COAD','GBM','PDAC','PRAD','ccRCC']):
 xx=9.40+(k%2)*1.57;yy=3.26+(k//2)*.82;organ(s,g,xx,yy,1.13,.59);txt(s,xx,yy+.61,1.13,.22,g,10.8,G,True,PP_ALIGN.CENTER)
txt(s,9.29,6.00,3.12,.32,'比较共同方向与背景差异',14,G,True,PP_ALIGN.CENTER)
foot(s,'两类匹配分别核对；GBM 的跨来源参照单独解释。插图均为示意。')
notes(s,'患者组织配对、对应标本的组学匹配分别成立才使用相应分析。示意组织使用非写实绘图。')

# Cohorts: organs are cropped from the approved reference; native labels remain editable.
s=new(4,'六癌种、七队列：按真实设计分别分析','本项目使用 CAMP 子集；保留各队列的配对、参照与治疗背景')
coh=[('BRCA','乳腺癌','患者内配对'),('COAD','结肠癌','患者内配对'),('PDAC','胰腺导管癌','患者内配对'),('PRAD','前列腺癌','患者内配对'),('ccRCC','肾透明细胞癌','两个队列，分别分析'),('GBM','胶质母细胞瘤','非配对 · 跨来源参照')]
for i,(g,note,mode) in enumerate(coh):
 x=.69+(i%3)*4.22;y=2.62+(i//3)*1.85;box(s,x,y,3.67,1.57,'F5F8F5');organ(s,g,x+.18,y+.19,1.17,1.13)
 txt(s,x+1.53,y+.18,1.98,.39,g,22,G,True);txt(s,x+1.54,y+.68,1.93,.34,note,14,INK);txt(s,x+1.54,y+1.17,1.98,.30,mode,10.2,M)
foot(s,'ccRCC 两队列在癌种计数时仅计一次；LYPLA1 专题聚焦 BRCA、COAD、PDAC、PRAD。')
notes(s,'队列和研究范围不变。器官图仅示意癌种；未修改原始科学数据或肿瘤亚型判断。')

# Workflow, illustrated and parallel evidence, not an invented linear significance gate.
s=new(5,'从代谢物到候选基因：映射与证据并行','生化知识提出候选，患者数据与细胞背景检验其研究价值')
for i,(t,body) in enumerate([('差异代谢物','癌种内肿瘤—参照\n名义 P < 0.05'),('身份与注释核对','名称、化学标识对齐\n异构体 / 混合峰另列'),('直接生化关系','反应底物 / 产物 · 转运\n必要复合体成员另注'),('代谢物—基因候选池','按特征 × 基因去重\n条件性关系单列')]):
 x=.69+i*3.18;box(s,x,2.52,2.74,2.29,'F5F8F5');txt(s,x+.15,2.68,2.43,.36,t,16,G,True)
 if i==0:art(s,'molecule',x+.88,3.14,.93,.87)
 elif i==1:art(s,'molecule',x+.15,3.16,1.04,.89);tag(s,x+1.25,3.45,1.16,'化学标识')
 elif i==2:art(s,'enzyme',x+.08,3.15,1.16,.89);art(s,'transporter',x+1.39,3.13,1.22,.91)
 else:art(s,'molecule',x+.16,3.22,.87,.72);arrow(s,x+1.1,3.5,.35);art(s,'dna',x+1.63,3.09,.74,1.0)
 txt(s,x+.16,4.18,2.43,.56,body,11.7,M)
 if i<3:arrow(s,x+2.83,3.51,.26)
line(s,11.6,4.83,11.6,5.13,G,1.3);line(s,2.38,5.13,11.6,5.13,G,1.3)
for x,t,b,icon in [(.7,'RNA 是否改变？','组织表达 · 同患者方向','dna'),(4.97,'代谢与 RNA 是否关联？','匹配肿瘤标本 · 实测关联','molecule'),(9.24,'候选在哪里表达？','单细胞定位 · 空间背景','cells')]:
 line(s,x+1.66,5.13,x+1.66,5.38,G,1.3);box(s,x,5.39,3.40,1.03,'FAFBFA');art(s,icon,x+.1,5.52,.74,.69);txt(s,x+.96,5.57,2.37,.34,t,13.4,G,True);txt(s,x+.96,6.07,2.3,.27,b,10.5,M)
foot(s,'三类证据并行检查；不以相关显著作为唯一入场条件。插图为关系示意，不是实验数据。')
notes(s,'完整流程沿用前版：差异代谢物、身份核对、生化关系映射、候选池，后接三个并行证据模块。相关性不等于因果；同通路不等于直接关系。所有文字、箭头和卡片原生可编辑，插图为独立裁剪对象。')

# Move the newly built opening slides ahead of the unchanged scientific results.
ids=R.slides._sldIdLst
for _ in range(5):
 old=ids[0];R.part.drop_rel(old.rId);ids.remove(old)
newids=list(ids)[-5:]
for item in newids:ids.remove(item)
for j,item in enumerate(newids):ids.insert(j,item)
for slide in R.slides:
 for sh in slide.shapes:
  for e in sh._element.xpath('.//a:effectRef'):e.set('idx','0')
  for rp in sh._element.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr'):
   rp.set('lang','zh-CN')
   for tag,face in [('a:latin','Microsoft YaHei'),('a:ea','微软雅黑'),('a:cs','Microsoft YaHei')]:
    old=rp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag.split(':')[1])
    if old is not None:rp.remove(old)
    rp.append(el(tag,typeface=face))
assert len(R.slides)==18
PPT='CAMP_泛癌与LYPLA1_插图精修可编辑版.pptx';R.save(O/PPT)
manifest=[dict(asset=n,source=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest(),use='Independent cropped artwork; no evidence plots generated') for n,p in files.items()]
(O/'illustration_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
(O/'imagegen_prompt.txt').write_text('Mode: built-in image_gen, reference-based background extraction/redrawing.\nSix isolated illustrations, 3 columns x 2 rows, transparent background, no text/labels/data. Patient with simplified tissues; molecule; DNA; enzyme; membrane transporter; epithelial cells. Muted sage/forest green and pale pink. Preserve soft semi-3D scientific illustration style. No photoreal flesh, no fabricated data charts.\nReference files: CAMP_IMAGE_fivepages_v3/03_CAMP优势与患者匹配.png; CAMP_IMAGE_mapping_redesign_v3/05_生化映射与多层证据.png\nFull atlas kept with native PPT crop settings; no destructive pixel crop.\n',encoding='utf-8')
shutil.copy2(O/PPT,V/PPT)
print(O/PPT)
