"""Aggregate-only plots and integration for new39 single-cell profiles."""
from pathlib import Path
import json,hashlib
import numpy as np,pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from matplotlib.colors import Normalize
from matplotlib.cm import ScalarMappable
repo=Path(__file__).resolve().parents[1];run='20260922T023000Z_sc39_v1';R=repo/'results/BRCA/06_EXTERNAL'/run
def save(d,p):d.to_csv(p,sep='\t',index=False,na_rep='NA',lineterminator='\n')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
dist=pd.read_csv(R/'new39_source_distributions.tsv',sep='\t');cross=pd.read_csv(R/'new39_cross_study_comparison.tsv',sep='\t');stab=pd.read_csv(R/'new39_source_stability.tsv',sep='\t');v=json.loads((R/'validation.json').read_text())
cohorts=['Wu2021','Pal2021_reprocessed'];genes=sorted(cross.gene);assert len(genes)==39 and cross.gene.is_unique
types=['Malignant_epithelial','Nonmalignant_epithelial','Fibroblasts','Perivascular','Endothelial','Myeloid','T_NK_cells','B_cells','Plasma_cells','Mast_cells','Unresolved_stromal'];labels=['Malignant epi.','Nonmalignant epi.','Fibroblasts','Perivascular','Endothelial','Myeloid','T / NK','B cells','Plasma cells','Mast cells','Unresolved stroma']
tables={};errors=[]
for c in cohorts:
 d=dist[dist.cohort.eq(c)&dist.partition.eq('ALL')];profiles=pd.read_csv(R/(c+'_celltype_profiles.tsv'),sep='\t');p=profiles[profiles.partition.eq('ALL')];j=d.merge(p,on=['gene','celltype'],suffixes=('_a','_b'),validate='one_to_one');ok=j.status_a.eq('DONE');assert j.loc[ok,'status_b'].eq('DONE').all();assert j.loc[ok,'n_a'].eq(j.loc[ok,'n_b']).all();err=abs(j.loc[ok,'effect_a']-j.loc[ok,'effect_b']).max();assert err<1e-10;errors.append(float(err));assert d.loc[d.status.eq('DONE'),'n'].ge(3).all();tables[c]=d
def matrix(c,col):return tables[c].pivot(index='gene',columns='celltype',values=col).reindex(index=genes,columns=types).to_numpy(float)
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'figure.facecolor':'white'})
mx=max(np.nanmax(matrix(c,'effect')) for c in cohorts);norm=Normalize(0,mx);cmap=plt.get_cmap('viridis').copy();cmap.set_bad('#e6e6e6')
fig,axes=plt.subplots(1,2,figsize=(19,16),sharey=True)
for ax,c in zip(axes,cohorts):
 a=matrix(c,'effect');det=matrix(c,'mean_detection');yy,xx=np.indices(a.shape);ok=np.isfinite(a);ax.scatter(xx[ok],yy[ok],s=det[ok]*180,c=a[ok],cmap=cmap,norm=norm,edgecolors='none');ax.scatter(xx[~ok],yy[~ok],s=12,c='#bbbbbb',marker='x',linewidths=.7);ax.set_xticks(range(len(types)),labels,rotation=55,ha='right');ax.set_yticks(range(39),genes);ax.set_ylim(38.7,-.7);ax.set_xlim(-.6,len(types)-.4);ax.set_title(c+'\nEqual source-label weights');ax.grid(axis='y',alpha=.13);ax.set_axisbelow(True)
for frac in [.1,.5,1.]:axes[1].scatter([],[],s=180*frac,c='#555555',label=str(int(frac*100))+'%')
axes[1].legend(title='Mean detected cells',loc='upper left',bbox_to_anchor=(1.01,.55),frameon=False)
fig.subplots_adjust(left=.09,right=.85,bottom=.13,top=.93,wspace=.1);cb=fig.add_axes([.88,.64,.015,.21]);fig.colorbar(ScalarMappable(norm=norm,cmap=cmap),cax=cb,label='Mean log1p(counts per 10k)');fig.suptitle('BRCA new39: expression by cell type',fontsize=17);fig.text(.1,.02,'Dot size: mean donor detection fraction. Gray x: missing / insufficient coverage, not zero.\nAt least 20 cells per donor-type and 3 source labels per type. No reclustering or cell-level hypothesis tests.',fontsize=10)
for ext in ['png','pdf']:fig.savefig(R/('new39_expression_dotplot.'+ext),dpi=180)
plt.close(fig)
for scaled,name in [(False,'new39_expression_heatmap'),(True,'new39_relative_pattern_heatmap')]:
 fig,axes=plt.subplots(1,2,figsize=(18,15),sharey=True)
 for ax,c in zip(axes,cohorts):
  a=matrix(c,'effect')
  if scaled:
   mu=np.nanmean(a,axis=1,keepdims=True);sd=np.nanstd(a,axis=1,keepdims=True);a=np.divide(a-mu,sd,out=np.full_like(a,np.nan),where=sd>0);cm=plt.get_cmap('RdBu_r').copy();cm.set_bad('#e6e6e6');im=ax.imshow(np.ma.masked_invalid(a),aspect='auto',cmap=cm,vmin=-2,vmax=2)
  else:im=ax.imshow(np.ma.masked_invalid(a),aspect='auto',cmap=cmap,norm=norm)
  ax.set_xticks(range(len(types)),labels,rotation=55,ha='right');ax.set_yticks(range(39),genes);ax.set_title(c)
 fig.subplots_adjust(left=.1,right=.87,bottom=.14,top=.94,wspace=.08);cb=fig.add_axes([.9,.6,.015,.25]);fig.colorbar(im,cax=cb,label='Within-gene/study z-score (clipped +/-2)' if scaled else 'Mean log1p(counts per 10k)');fig.suptitle('BRCA new39: '+('relative expression patterns' if scaled else 'cell-type expression'),fontsize=16);fig.text(.1,.025,'Gray = not measured / insufficient coverage'+('. Row scaling shows relative patterns, not absolute expression.' if scaled else '. Shared color scale does not remove study/platform differences.')+'\nGene order is alphabetical, not a ranking. Annotation provenance differs between studies.',fontsize=10)
 for ext in ['png','pdf']:fig.savefig(R/(name+'.'+ext),dpi=180)
 plt.close(fig)
# Append new columns to latest156 master, preserving all original values.
src=repo/'results/BRCA/03_PATIENT/20260922T020919Z_pairedRNA_p005_view_v1/genes156_RNA_retention_appended.tsv';old=pd.read_csv(src,sep='\t',dtype=str,keep_default_na=False);cols=old.columns.tolist();a=old.merge(cross.rename(columns={c:'sc39_'+c for c in cross if c!='gene'}),on='gene',how='left',validate='one_to_one');assert len(a)==156 and a[cols].equals(old);a['sc39_scope']=np.where(a.gene.isin(genes),'NEW39_ANALYZED','OLD117_HISTORY_RETAINED_NOT_RERUN');save(a,R/'genes156_sc39_appended.tsv')
cn={'Malignant_epithelial':'恶性上皮','Nonmalignant_epithelial':'非恶性上皮','Fibroblasts':'成纤维细胞','Perivascular':'血管周细胞','Endothelial':'内皮细胞','Myeloid':'髓系细胞','T_NK_cells':'T/NK细胞','B_cells':'B细胞','Plasma_cells':'浆细胞','Mast_cells':'肥大细胞','Unresolved_stromal':'未细分基质'}
reader=cross.copy();reader['Wu最高类别']=reader.Wu_top.map(cn);reader['Pal最高类别']=reader.Pal_top.map(cn);save(reader,R/'new39_cell_source_reader_cn.tsv')
notes=pd.DataFrame([{'说明':'覆盖','内容':'全39基因；逐供者汇总；不足3个来源标签时不展示均值'}, {'说明':'最高表达','内容':'仅是所测细胞类别中平均表达最高，不是唯一来源或作用细胞'}, {'说明':'重采样','内容':'描述总体排名的稳定性，不是患者百分比或P值'}, {'说明':'图','内容':'点大小为供者等权检测比例；灰色为缺测/覆盖不足'}, {'说明':'历史','内容':'156表旧列原样保留，本轮新增信息使用sc39前缀'}, {'说明':'LPCAT4','内容':'稳定ID匹配只用于本次单细胞；CAMP微阵列仍独立待核'}])
book=R/'BRCA_新增39基因_单细胞来源.xlsx'
with pd.ExcelWriter(book,engine='openpyxl') as w:
 for name,d in [('说明',notes),('39基因跨研究',reader),('细胞类别表达',dist),('来源稳定性',stab),('LPCAT4身份',pd.read_csv(R/'LPCAT4_stable_identity.tsv',sep='\t')),('156基因完整保留',a)]:d.to_excel(w,sheet_name=name,index=False)
wb=load_workbook(book)
for ws in wb:
 ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D');cell.alignment=Alignment(wrap_text=True)
 for i,col in enumerate(ws[1],1):ws.column_dimensions[get_column_letter(i)].width=min(42,max(15,len(str(col.value))*1.1))
wb.save(book);assert load_workbook(book,read_only=True)['39基因跨研究'].max_row==40
v.update(independent_summary_max_error=max(errors),all156_prior_columns_preserved=True,plots_missing_values_not_zero=True,plot_order='alphabetical genes;fixed celltypes;no clustering',plot_software=dict(matplotlib=matplotlib.__version__,numpy=np.__version__,pandas=pd.__version__))
(R/'delivery_validation.json').write_text(json.dumps(v,indent=2));rows=[]
for g in ['LYPLA1','LPCAT1','ABHD12','LYPLA2','ENPP2','ACSL4','SUCLA2','LPCAT4']:
 t=reader[reader.gene.eq(g)].iloc[0];rows.append(f'|{g}|{t["Wu最高类别"]}|{t["Pal最高类别"]}|{t.Wu_top_frequency:.1%}|{t.Pal_top_frequency:.1%}|')
text=f'''# 新增39基因：单细胞表达位置与两研究对照

## 本轮问题
新增39基因在乳腺肿瘤的哪些细胞类型中表达？来源标签覆盖与两研究结果是否一致？

## 输入与范围
复用原Wu2021及Pal2021_reprocessed H5AD。Wu沿用作者celltype_major；Pal沿用Chen2026对Pal2021的重注释，筛选pal_2021子集。原矩阵与供者级表达只存server165；无重新聚类、无新P/q、无新增机制分析。
全39均保留。每个来源标签×细胞类别至少20细胞，类别至少3个合格来源标签；先在供者内求均值，再给供者等权。来源标签不等于新增核验的独立患者身份。

## 实际结果
两研究各自可评估类别范围内，{v['same_top']}个基因最高类别相同；其中{v['same_top_both_bootstrap080']}个在两边1000次来源标签重采样中的最高类别保持频率均≥80%。限制在两研究共同类别时一致数为{v['same_top_shared']}，两个口径不能混用。

|基因|Wu最高类别|Pal最高类别|Wu保持频率|Pal保持频率|
|---|---|---|---:|---:|
'''+ '\n'.join(rows)+'''

## 新手解释
点图颜色是表达水平，点大小是供者等权的表达细胞比例；灰叉/灰格是缺测或覆盖不足，不是零。表达热图使用原汇总尺度；relative_pattern热图逐基因逐研究标准化，只显示相对分布，不代表绝对表达高低。
“最高类别”不是唯一表达来源，也不是功能验证。保持频率是总体排序的描述性稳定性，不能理解成该比例的患者均符合；配对供者方向比例另列在完整表中。

## 限制与反证
研究的细胞类别范围、注释与技术不同，不能用最高类别差异直接证明癌种或亚型机制。Wu另保留治疗未暴露Naive子集敏感性结果；未把多个细胞当成独立患者进行显著性检验。
LPCAT4只按Q643R3对应稳定Ensembl基因ID识别，另列MBOAT2等候选条目；若唯一匹配成功仅解决本次单细胞身份，不自动解除CAMP微阵列探针的历史身份待核项。
RNA表达位置不能证明对应CAMP代谢物在哪种细胞产生或被消耗；单细胞和CAMP不是同一批病例。

## 当前决定
把新增39表达背景附加到完整156表；旧117单细胞/功能/统计字段保留原值。历史NOT_RUN列是以前阶段记录，本轮状态看sc39列。

## 下一步
结合适用功能研究选择有限问题，不仅取表达最高的细胞，不按表达位置淘汰其他候选。本轮不启动细胞通讯、拟时序或新的干预计算。

## 复现
server165独占新运行目录放source/genes_unique.tsv、两个冻结配置及旧profile/stability脚本，.running内容sc39_v1。执行`python3 brca_sc39_v1.py <run_directory>`。脚本生成39基因适配器，并流式读取既有H5AD；实际适配器也随交付保存。公开汇总取回后运行code/brca_sc39_deliver_v1.py生成图和整合表。source_manifest记录输入与脚本SHA256。
'''
(R/'README_CN.md').write_text(text,encoding='utf-8');(R/'.gitattributes').write_text('* -text\n')
save(pd.DataFrame([dict(path=str(src.relative_to(repo)).replace('\\','/'),sha256=sha(src)),dict(path='code/brca_sc39_deliver_v1.py',sha256=sha(Path(__file__))),dict(path='reference/brca_mapping190_sources_v1/LPCAT4.json',sha256=sha(repo/'reference/brca_mapping190_sources_v1/LPCAT4.json'))]),R/'integration_source_manifest.tsv')
save(pd.DataFrame([dict(file=p.name,sha256=sha(p)) for p in sorted(R.iterdir()) if p.is_file() and p.name!='checksums.tsv']),R/'checksums.tsv')
print(reader[['gene','Wu最高类别','Pal最高类别','same_top_and_both_bootstrap_ge080']].to_string(index=False))
