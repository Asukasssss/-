"""Configuration-driven, read-only scientific report renderer from frozen Git aggregates.
No patient fitting, imputation, statistical tests, P adjustment or UMAP computation.
"""
from pathlib import Path
import argparse,csv,hashlib,io,json,subprocess,sys,zipfile,platform,shutil
import numpy as np
import pandas as pd
import yaml

ROOT=Path(__file__).resolve().parents[1]
def git_bytes(revision,path):
 return subprocess.check_output(['git','show',f'{revision}:{path}'],cwd=ROOT)
def num(series):return pd.to_numeric(series,errors='coerce')
def truth(x):return str(x).lower()=='true'
def dump(path,obj):path.write_bytes((json.dumps(obj,ensure_ascii=False,indent=2)+'\n').encode())
def tsv(path,df):df.to_csv(path,sep='\t',index=False,encoding='utf-8',lineterminator='\n',na_rep='NA')
def tiers(df):
 return np.select([df['status'].ne('DONE')|num(df['p_value']).isna(),num(df['q_value']).lt(.05),num(df['p_value']).lt(.05)],['不可评估','q<0.05','仅P<0.05'],default='P>=0.05')
def load(config):
 c=yaml.safe_load(Path(config).read_text(encoding='utf-8'))
 rev=c['input_commit'];resolved=subprocess.check_output(['git','rev-parse',rev],cwd=ROOT,text=True).strip()
 if rev!=resolved:raise ValueError('input_commit must be an immutable complete commit SHA')
 data={};blobs={};manifest=[]
 for key,path in c['files'].items():
  b=git_bytes(rev,path);blobs[path]=b
  df=pd.read_csv(io.BytesIO(b),sep='\t',dtype=str,keep_default_na=False)
  rename={actual:canonical for canonical,actual in c.get('columns',{}).items() if canonical in ['gene','metabolite_key','metabolite_name','effect','p_value','q_value','status','n','analysis_type']}
  # Explicit arbitrary source-field aliases may be supplied by another cancer adapter.
  rename.update(c.get('field_aliases',{}));data[key]=df.rename(columns=rename)
  manifest.append(dict(key=key,path=path,input_commit=rev,sha256=hashlib.sha256(b).hexdigest(),rows=len(df)))
 for asset in c.get('assets',[]):
  for path in [asset['path'],asset.get('source_table')]:
   if path and path not in blobs:
    b=git_bytes(rev,path);blobs[path]=b;manifest.append(dict(key='asset',path=path,input_commit=rev,sha256=hashlib.sha256(b).hexdigest(),rows='NA'))
 if c.get('design_source'):
  p=c['design_source'];b=git_bytes(rev,p);blobs[p]=b;manifest.append(dict(key='design',path=p,input_commit=rev,sha256=hashlib.sha256(b).hexdigest(),rows='NA'))
 return c,data,blobs,pd.DataFrame(manifest)
def prepare(c,d):
 fam=c['families'];m=d['metabolites'];a=d['association']
 take=lambda df,key:df[df.analysis_type.eq(fam.get(key,'__MISSING__'))].copy()
 z=dict(metabolites=take(m,'metabolite_primary'),sensitivity=take(m,'metabolite_sensitivity'),association=take(a,'association_primary'),association_sensitivity=take(a,'association_sensitivity'),genes=d['genes'].copy(),relations=d['relations'].copy(),mapping=d['mapping'].copy(),rna=d.get('rna',pd.DataFrame()).copy())
 z['profiles']=pd.concat([d[x] for x in c.get('sc',{}).get('profile_sources',[])],ignore_index=True) if c.get('sc',{}).get('profile_sources') else pd.DataFrame()
 z['stability']=pd.concat([d[x] for x in c.get('sc',{}).get('stability_sources',[])],ignore_index=True) if c.get('sc',{}).get('stability_sources') else pd.DataFrame()
 reader=d['reader'].copy();status=[]
 if len(z['stability']):
  sc=c['sc'];s=z['stability'];s=s[s[sc['partition_field']].eq(sc['partition']) & s[sc['cohort_field']].isin(sc['cohorts'])]
  assert not s.duplicated([sc['cohort_field'],'gene']).any()
  p=z['profiles'];p=p[p[sc['partition_field']].eq(sc['partition'])]
  for _,r in s.iterrows():
   top=r[sc['stability_top_field']];prof=p[p.gene.eq(r.gene)&p[sc['cohort_field']].eq(r[sc['cohort_field']])&p[sc['type_field']].eq(top)]
   good=r.status=='DONE' and len(prof)==1 and prof.iloc[0].status=='DONE'
   status.append(dict(gene=r.gene,cohort=r[sc['cohort_field']],source_status=r.status,reason=r.reason,display_top=top if good else '暂不可定位',mechanical_top=top,top_n=prof.iloc[0]['n'] if len(prof)==1 else 'NA',top_detection=prof.iloc[0][sc['detection_field']] if len(prof)==1 else 'NA',bootstrap=r.get('bootstrap_top_frequency','NA')))
  # Neutral study-prefixed display columns; never trust baseline easy-table mechanical ranks.
  reader=reader.drop(columns=[x for x in c.get('legacy_source_display_columns',[]) if x in reader])
  for cohort in sc['cohorts']:
   t=pd.DataFrame(status);t=t[t.cohort.eq(cohort)].drop(columns='cohort').set_index('gene').add_prefix(cohort+'__')
   reader=reader.merge(t,left_on='gene',right_index=True,how='left',validate='one_to_one')
 z['reader']=reader;z['source_status']=pd.DataFrame(status)
 selections=[]
 for e in c.get('examples',[]):
  rows=z['relations'][z['relations'][c['columns']['relation_id']].eq(e['relation_id'])]
  if len(rows)!=1 or rows.iloc[0].gene!=e['gene']:raise ValueError('Example relation is absent, ambiguous or assigned to another gene: '+e['relation_id'])
  selections.append(dict(e,metabolite_key=rows.iloc[0].metabolite_key,metabolite_name=rows.iloc[0].metabolite_name,selection_scope='DISPLAY_EXAMPLE_NOT_RANK',source_commit=c['input_commit']))
 z['selection']=pd.DataFrame(selections)
 return z
def validate(c,d,z):
 assert not z['genes'].gene.duplicated().any()
 assert not z['relations'][c['columns']['relation_id']].duplicated().any()
 counts=dict(metabolites=len(z['metabolites']),workpool=int(num(z['metabolites'].p_value).lt(.05).sum()),relations=len(z['relations']),genes=len(z['genes']),current_genes=z['relations'].gene.nunique())
 for k,v in c.get('expected',{}).items():
  if counts[k]!=v:raise ValueError(f'Expected {k}={v}; actual {counts[k]}')
 for kind in ['metabolites','association','rna']:
  t=z[kind]
  if t.empty:continue
  for col in ['effect','p_value','q_value','n','status']:
   if col not in t:raise ValueError(kind+' missing '+col)
  for col in ['p_value','q_value']:
   vals=num(t[col]).dropna();assert vals.between(0,1).all()
 if len(z['source_status']):
  t=z['source_status'];assert t[t.source_status.ne('DONE')].display_top.eq('暂不可定位').all()
 return dict(status='INPUTS_VALIDATED_READ_ONLY',counts=counts,input_commit=c['input_commit'],new_P_q=0,raw_patient_data_used=False)

def build(c,d,z,blobs,manifest,out,config_path):
 if out.exists() and any(out.iterdir()):raise FileExistsError('Refusing nonempty output directory: '+str(out))
 import matplotlib
 matplotlib.use('Agg')
 import matplotlib.pyplot as plt
 from matplotlib import font_manager,colors
 from matplotlib.backends.backend_pdf import PdfPages
 import pymupdf as fitz
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 font=next((p for p in c['font_candidates'] if Path(p).exists()),None)
 if not font:raise RuntimeError('No CJK font; configure installed font_candidates (font files not bundled)')
 font_manager.fontManager.addfont(font);family=font_manager.FontProperties(fname=font).get_name()
 plt.rcParams.update({'font.family':family,'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'axes.unicode_minus':False,'pdf.fonttype':42,'svg.fonttype':'path','figure.facecolor':'white','axes.facecolor':'white'})
 out.mkdir(parents=True,exist_ok=True)
 for folder in ['figures','figure_data','appendix','appendix/existing_figures','rendered_pages','repro']: (out/folder).mkdir(parents=True,exist_ok=True)
 figures=[];pages=[];foot='描述性展示；不新增P/q；输入 '+c['input_commit'][:12]
 palette={'q<0.05':'#087F8C','仅P<0.05':'#DE9B34','P>=0.05':'#97A4AE','不可评估':'#D6DADF'}
 def save(ident,title,frame,draw,narrative,sources,transform='原统计值；仅布局与选择',main=True):
  fig=plt.figure(figsize=(11.6,4.8));draw(fig)
  fig.savefig(out/'figures'/f'{ident}.png',dpi=180,bbox_inches='tight',facecolor='white')
  fig.savefig(out/'figures'/f'{ident}.pdf',bbox_inches='tight')
  fig.savefig(out/'figures'/f'{ident}.svg',bbox_inches='tight');plt.close(fig)
  tsv(out/'figure_data'/f'{ident}.tsv',frame)
  ns=c['narratives'].get(narrative,['问题见标题','数据来自已冻结汇总。','不新增推断。'])
  caption=title+'。'+ns[1]+' '+ns[2]
  (out/'figures'/f'{ident}_caption_CN.md').write_bytes((caption+'\n数据变换：'+transform+'\n').encode())
  figures.append(dict(figure_id=ident,title=title,format='PNG;PDF;SVG',source_commit=c['input_commit'],sources=';'.join(sources),source_run=';'.join(sorted({str(x) for x in frame.get('run_id',pd.Series(dtype=str)).unique()})),relation_keys=';'.join(frame.get('relation_id',frame.get('relation_key',pd.Series(dtype=str))).astype(str).unique()),filter='config examples' if len(frame)<len(z['relations']) else 'full specified scope',transform=transform,data_tsv='figure_data/'+ident+'.tsv',caption='figures/'+ident+'_caption_CN.md'))
  if main:pages.append(dict(title=title,question=ns[0],result=ns[1],limit=ns[2],figure='figures/'+ident+'.pdf',id=ident))
 def counts_text(t):
  return f'计划 {len(t)}；可评估 {int(t.status.eq("DONE").sum())}；P<0.05 {int(num(t.p_value).lt(.05).sum())}；q<0.05 {int(num(t.q_value).lt(.05).sum())}'
 def overview(t,xlabel):
  def draw(fig):
   ax=fig.add_subplot(111);t1=t[t.status.eq('DONE')];p=num(t1.p_value)
   for label,color in palette.items():
    idx=tiers(t1)==label;ax.scatter(num(t1.loc[idx,'effect']),-np.log10(p[idx].clip(lower=1e-300)),s=24,c=color,label=label,alpha=.8)
   ax.axvline(0,color='#D6DADF');ax.axhline(-np.log10(.05),color='#DE9B34',ls=':',lw=.8)
   ax.set(xlabel=xlabel,ylabel='−log10(已保存的 P)',title=counts_text(t));ax.legend(fontsize=9);fig.tight_layout()
  return draw
 def directions(t,labels,up,down,equal):
  def draw(fig):
   ax=fig.add_subplot(111);y=np.arange(len(t));left=np.zeros(len(t))
   for field,label,color in [(down,'降低','#337EAA'),(equal,'相等','#CAD2D8'),(up,'升高','#C76A40')]:
    v=num(t[field]).fillna(0).values;ax.barh(y,v,left=left,color=color,label=label);left+=v
   ax.set_yticks(y,labels);ax.invert_yaxis();ax.set_xlabel('配对数（不是逐患者显著次数）');ax.legend(loc='lower right');fig.tight_layout()
  return draw
 sel=z['selection'];names=sel.gene.tolist()
 def selected(t):return t[t.metabolite_key.isin(sel.metabolite_key)].drop_duplicates('metabolite_key').copy()
 def fig_design(fig):
  ax=fig.add_subplot(111);ax.axis('off');design=c.get('design',[])
  for i,r in enumerate(design):
   x=.03+i*.96/max(len(design),1);w=.9/max(len(design),1)
   ax.add_patch(plt.Rectangle((x,.56),w,.35,facecolor='#EAF3F5',edgecolor='#087F8C'))
   ax.text(x+w/2,.78,str(r['value']),ha='center',fontsize=27,color='#087F8C');ax.text(x+w/2,.61,r['label'],ha='center',fontsize=10)
  chain=['配对代谢物','直接映射','肿瘤内相关','配对RNA','细胞来源','候选整合']
  for i,label in enumerate(chain):ax.text(.06+i*.16,.28,label,fontsize=13,ha='center')
  ax.text(.5,.08,'连接依据来自作者信息；不同统计单位不合并成一个样本量',ha='center',fontsize=12)
 save('01_design','研究设计与数据范围',pd.DataFrame(c.get('design',[])),fig_design,'design',[c.get('design_source','config')])
 # Each cohort remains a separate plot, never combined into a new test.
 for j,(cohort,t) in enumerate(z['metabolites'].groupby('cohort',sort=False)):
  suff='' if j==0 else '_'+str(j+1)
  save('02_metabolites'+suff,'配对代谢物：全量效应概览 / '+cohort,t,overview(t,'肿瘤−自身正常：均值差（作者处理尺度）'),'metabolites',[c['files']['metabolites']])
  t1=selected(t);save('03_metabolite_counts'+suff,'示例代谢物：配对方向人数',t1,directions(t1,t1.metabolite_name.tolist(),'pairs_higher','pairs_lower','pairs_equal'),'paired_counts',[c['files']['metabolites']])
  sens=z['sensitivity'];sens=sens[sens.cohort.eq(cohort)] if len(sens) else sens
  if len(sens):
   q=t[['metabolite_key','metabolite_name','effect','n','q_value','status']].merge(sens[['metabolite_key','effect','n','q_value','status']],on='metabolite_key',suffixes=('_primary','_available'),validate='one_to_one')
   def draw_s(fig,q=q):
    ax=fig.add_subplot(111);ok=q.status_available.eq('DONE');ax.scatter(num(q.loc[ok,'effect_primary']),num(q.loc[ok,'effect_available']),c='#087F8C',s=18,alpha=.65)
    lo=min(ax.get_xlim()[0],ax.get_ylim()[0]);hi=max(ax.get_xlim()[1],ax.get_ylim()[1]);ax.plot([lo,hi],[lo,hi],ls=':',c='#AAB3BD')
    ax.set(xlabel='主分析均值差',ylabel='可用值子集均值差',title=f'可比较 {int(ok.sum())} 项；其余保留缺项；每点是一项代谢物')
    fig.tight_layout()
   save('04_sensitivity'+suff,'缺失敏感性：汇总效应对照',q,draw_s,'sensitivity',[c['files']['metabolites']])
 def mapping(fig):
  ax=fig.add_subplot(121);v=z['mapping'].mapping_state.value_counts();ax.barh([c.get('mapping_labels',{}).get(x,x) for x in v.index],v.values,color='#087F8C')
  for i,vv in enumerate(v):ax.text(vv+.5,i,str(vv),va='center')
  ax.set_xlabel('入口代谢物项数');ax=fig.add_subplot(122)
  now=set(z['relations'].gene);allg=set(z['genes'].gene);ax.bar(['直接关系','当前基因','历史保留'],[len(z['relations']),len(now),len(allg-now)],color=['#087F8C','#337EAA','#9AA6B2']);ax.set_ylabel('数量（不同计数单位）');fig.tight_layout()
 mapping_counts=pd.DataFrame([dict(label=k,count=int(v),unit='入口代谢物') for k,v in z['mapping'].mapping_state.value_counts().items()]+[dict(label='直接关系',count=len(z['relations']),unit='关系'),dict(label='当前基因',count=z['relations'].gene.nunique(),unit='基因'),dict(label='历史保留',count=len(set(z['genes'].gene)-set(z['relations'].gene)),unit='基因')])
 save('05_mapping','直接生化映射与保留范围',mapping_counts,mapping,'mapping',[c['files']['mapping'],c['files']['relations'],c['files']['genes']],transform='按状态计数；关系行数；去重基因及集合差；无统计检验')
 for j,(cohort,t) in enumerate(z['association'].groupby('cohort',sort=False)):
  suff='' if j==0 else '_'+str(j+1)
  save('06_associations'+suff,'肿瘤内部关系：全量概览 / '+cohort,t,overview(t,'Spearman ρ（无量纲）'),'associations',[c['files']['association']])
  a=t.merge(sel[['relation_id','gene','role']],on=['relation_id','gene'],validate='one_to_one');b=z['association_sensitivity'];b=b[b.cohort.eq(cohort)] if len(b) else b
  b=b.merge(sel[['relation_id','gene','role']],on=['relation_id','gene'],validate='one_to_one') if len(b) else b
  f=pd.concat([a.assign(display_family='主分析'),b.assign(display_family='可用值')],ignore_index=True)
  def forest(fig,f=f):
   ax=fig.add_subplot(111)
   for i,gene in enumerate(names):
    for off,(family,color) in enumerate([('主分析','#087F8C'),('可用值','#DE9B34')]):
     row=f[f.gene.eq(gene)&f.display_family.eq(family)]
     if row.empty:continue
     r=row.iloc[0];y=i+(off-.5)*.24
     if r.status!='DONE':ax.text(0,y,'NA',fontsize=8);continue
     x=float(r.effect);lo=pd.to_numeric(r.ci_lower,errors='coerce');hi=pd.to_numeric(r.ci_upper,errors='coerce')
     if np.isfinite(lo) and np.isfinite(hi):ax.plot([lo,hi],[y,y],color=color,lw=1.3)
     ax.plot(x,y,'o',color=color,ms=4,label=family if i==0 else None)
   labels=[]
   for g in names:
    vals=[]
    for family in ['主分析','可用值']:
     rr=f[f.gene.eq(g)&f.display_family.eq(family)];vals.append(f'{float(rr.iloc[0].n):g}' if len(rr) and rr.iloc[0].status=='DONE' else 'NA')
    labels.append(g+'  n='+ '/'.join(vals))
   ax.set_yticks(range(len(names)),labels);ax.invert_yaxis();ax.axvline(0,color='#ABB4BD',ls=':');ax.set(xlim=(-1,1),xlabel='Spearman ρ及已保存的点95%区间；n=主分析/可用值');ax.legend();fig.tight_layout()
  save('07_association_intervals'+suff,'固定示例：主分析与可用值相关',f,forest,'forest',[c['files']['association']])
 if c.get('modules',{}).get('rna',True) and len(z['rna']):
  for j,(cohort,t) in enumerate(z['rna'].groupby('cohort',sort=False)):
   suff='' if j==0 else '_'+str(j+1)
   save('08_RNA'+suff,'配对RNA：全量效应概览 / '+cohort,t,overview(t,'RNA平均肿瘤−正常差（作者表达尺度）'),'rna',[c['files']['rna']])
   t1=t[t.gene.isin(names)].set_index('gene').reindex(names).reset_index();t1=t1[t1.status.notna()]
   save('09_RNA_counts'+suff,'示例基因：配对RNA方向人数',t1,directions(t1,t1.gene.tolist(),'positive_pairs','negative_pairs','equal_pairs'),'rna_counts',[c['files']['rna']])
 if len(z['profiles']) and c.get('modules',{}).get('single_cell',True):
  sc=c['sc'];labels=sc.get('labels',{});allprofiles=z['profiles'];allprofiles=allprofiles[allprofiles[sc['partition_field']].eq(sc['partition'])]
  for study in sc['cohorts']:
   p=allprofiles[allprofiles[sc['cohort_field']].eq(study)]
   if p.empty:continue
   present=set(p[sc['type_field']]);cats=[x for x in sc['celltypes'] if x in present]+sorted(present-set(sc['celltypes']))
   chosen=p[p.gene.isin(names)].copy()
   def dots(fig,p=chosen,cats=cats,study=study):
    ax=fig.add_subplot(111);v=p[p.status.eq('DONE')];maxv=num(v.effect).max();norm=colors.Normalize(0,maxv if maxv>0 else 1)
    for yi,gene in enumerate(names):
     for xi,ct in enumerate(cats):
      row=p[p.gene.eq(gene)&p[sc['type_field']].eq(ct)]
      if len(row)!=1 or row.iloc[0].status!='DONE':ax.plot(xi,yi,'x',c='#999999',ms=4);continue
      r=row.iloc[0];det=float(r[sc['detection_field']]);ax.scatter(xi,yi,s=220*det,c=[float(r.effect)],cmap='viridis',norm=norm,edgecolors='#888888',linewidths=.15)
    ax.set_xticks(range(len(cats)),[labels.get(x,x) for x in cats],rotation=25,ha='right');ax.set_yticks(range(len(names)),names);ax.invert_yaxis()
    sm=plt.cm.ScalarMappable(norm=norm,cmap='viridis');fig.colorbar(sm,ax=ax,label='供者等权平均 log1p(counts/10k)')
    for s0 in [.1,.5,1]:ax.scatter([],[],s=220*s0,c='#999999',label=f'检出 {s0:.0%}')
    coverage=num(v.n)
    ax.legend(loc='upper left',bbox_to_anchor=(1.2,1),fontsize=8);ax.set_title(study+f'；有效来源标签n={coverage.min():g}–{coverage.max():g}；叉=覆盖不足');fig.subplots_adjust(left=.13,right=.72,bottom=.23,top=.89)
   save('10_source_'+study,'单细胞来源 / '+study,chosen,dots,'dot',[c['files'][x] for x in sc['profile_sources']])
   # Full appendix heatmap pagination, independent scaling within each gene/study.
   allnames=sorted(z['genes'].gene);pdf=PdfPages(out/'appendix'/('all_genes_heatmap_'+study+'.pdf'))
   for start in range(0,len(allnames),sc['heatmap_page_rows']):
    ns=allnames[start:start+sc['heatmap_page_rows']];chunk=p[p.gene.isin(ns)].copy();chunk['row_scaled']=np.nan
    for g in ns:
     mask=chunk.gene.eq(g)&chunk.status.eq('DONE');vals=num(chunk.loc[mask,'effect']);mx=vals.max()
     chunk.loc[mask,'row_scaled']=vals/mx if mx>0 else 0
    def heat(fig,chunk=chunk,ns=ns):
     ax=fig.add_subplot(111);mat=chunk.pivot(index='gene',columns=sc['type_field'],values='row_scaled').reindex(index=ns,columns=cats)
     cm=plt.get_cmap('viridis').copy();cm.set_bad('#BFC5CB');im=ax.imshow(mat.values.astype(float),aspect='auto',vmin=0,vmax=1,cmap=cm)
     ss=z['source_status'];bad=set(ss[(ss.cohort==study)&ss.source_status.ne('DONE')].gene)
     ax.set_yticks(range(len(ns)),[g+' ·暂不可定位' if g in bad else g for g in ns],fontsize=7);ax.set_xticks(range(len(cats)),[labels.get(x,x) for x in cats],rotation=30,ha='right',fontsize=8);fig.colorbar(im,ax=ax,label='同一基因内相对均值（0–1）');ax.set_title(study+'；灰色=类别覆盖不足；低检出基因不作定位结论');fig.tight_layout()
    ident='appendix_heat_'+study+'_'+str(start//sc['heatmap_page_rows']+1)
    save(ident,'全候选来源热图 / '+study,chunk,heat,'dot',[c['files'][x] for x in sc['profile_sources']],transform='每基因每研究均值除该行最大值；真零保持0；缺测NaN；不重排聚类',main=False)
    temp=plt.figure(figsize=(11.6,7.5));heat(temp);pdf.savefig(temp);plt.close(temp)
   pdf.close()
 if c.get('modules',{}).get('subtypes',True) and c.get('subtypes') and 'subtypes' in d:
  t=d['subtypes'];t=t[t.gene.isin(c.get('subtype_examples',names))].copy();display=[]
  for _,r in t.iterrows():
   for sub in c['subtypes']:
    display.append(dict(gene=r.gene,subtype=sub,own_top=r[sub+'_top'] if r[sub+'_status']=='DONE' else '暂不可定位',own_status=r[sub+'_status'],shared_top=r[sub+'_shared_top'] if r[sub+'_shared_status']=='DONE' else '暂不可定位',shared_status=r[sub+'_shared_status'],pooled_top=r.get('equal_subtype_top','NA'),n=r[sub+'_n'],gap=r[sub+'_gap']))
  frame=pd.DataFrame(display)
  def subtypeplot(fig):
   ax=fig.add_subplot(111);ax.axis('off');table=[]
   for gene in t.gene:
    ss=frame[frame.gene.eq(gene)];table.append([gene]+[' / '.join([str(r.own_top),str(r.shared_top)]) for _,r in ss.iterrows()])
   tb=ax.table(cellText=table,colLabels=['基因']+[s+'：自身 / 共同类别' for s in c['subtypes']],cellLoc='center',loc='center',colWidths=[.13]+[.87/len(c['subtypes'])]*len(c['subtypes']));tb.auto_set_font_size(False);tb.set_fontsize(8);tb.scale(1,2.6)
   ax.set_title('三种口径分别保留；各亚型最高类别 ≠ 等权汇总最高类别',pad=18)
  save('12_subtype','已有分型背景：描述性排名',frame,subtypeplot,'subtypes',[c['files']['subtypes']])
 # Existing figures are copied exactly. No invented data reconstruction or vectorization.
 for i,asset in enumerate(c.get('assets',[])):
  path=asset['path'];dest=('figures/' if asset['role']=='main_umap' else 'appendix/existing_figures/')+Path(path).name
  (out/dest).write_bytes(blobs[path]);index=asset.get('source_table');ident='reused_'+str(i+1)
  frame=pd.read_csv(io.BytesIO(blobs[index]),sep='\t',dtype=str,keep_default_na=False) if index else pd.DataFrame([{'source_image':path,'raw_coordinates':'NOT_EXPORTED'}])
  tsv(out/'figure_data'/f'{ident}.tsv',frame)
  (out/'figures'/f'{ident}_caption_CN.md').write_bytes((asset['caption']+'\n逐细胞数据未导出；本TSV为原图页码/尺度索引，不是重建坐标。\n').encode())
  figures.append(dict(figure_id=ident,title=asset['caption'],format='REUSED_'+Path(path).suffix,source_commit=c['input_commit'],sources=path+';'+str(index),source_run=Path(path).parent.name,relation_keys='',filter='EXACT_EXISTING_ASSET',transform='字节复用；不重算嵌入；不伪造矢量',data_tsv='figure_data/'+ident+'.tsv',caption='figures/'+ident+'_caption_CN.md'))
  if asset['role']=='main_umap':
   ns=c['narratives']['umap'];pages.append(dict(title='表达与细胞类型：同坐标对照',question=ns[0],result=ns[1],limit=ns[2],figure=dest,id=ident))
 # The evidence matrix encodes separate existing tests, never a new combined score.
 r=z['relations'].merge(sel[['relation_id','role']],left_on=c['columns']['relation_id'],right_on='relation_id',validate='one_to_one');r=r.set_index('gene').reindex(names).reset_index()
 rr=z['rna'].set_index('gene') if len(z['rna']) else pd.DataFrame()
 em=[]
 for _,x in r.iterrows():
  row=dict(gene=x.gene,relation_id=x.relation_id)
  for prefix,title in [('metabolite_paired','代谢物'),('CAMP','内部相关'),('CAMP_available','相关敏感性')]:
   row[title]=tiers(pd.DataFrame([dict(status=x.get(prefix+'_status','DONE'),p_value=x.get(prefix+'_p_value','NA'),q_value=x.get(prefix+'_q_value','NA'))]))[0]
  row['配对RNA']=tiers(rr.loc[[x.gene]])[0] if len(rr) and x.gene in rr.index else '不可评估';em.append(row)
 em=pd.DataFrame(em)
 def evidence(fig):
  ax=fig.add_subplot(111);cols=[x for x in em if x not in ('gene','relation_id')];codes={k:i for i,k in enumerate(palette)};mat=np.array([[codes[x] for x in em[col]] for col in cols]).T
  ax.imshow(mat,aspect='auto',cmap=colors.ListedColormap(list(palette.values())),vmin=-.5,vmax=len(palette)-.5)
  ax.set_yticks(range(len(em)),em.gene);ax.set_xticks(range(len(cols)),cols)
  for i in range(len(em)):
   for j,col in enumerate(cols):ax.text(j,i,em.iloc[i][col],ha='center',va='center',fontsize=8,color='white' if mat[i,j]==0 else '#26384A')
  ax.set_title('各列独立检验集合；不相加为总分，不以全列显著作入场门槛');fig.tight_layout()
 save('14_evidence','候选证据矩阵：支持与缺项并列',em,evidence,'evidence',[c['files']['relations']]+([c['files']['rna']] if 'rna' in c['files'] else []))
 def discuss(fig):
  ax=fig.add_subplot(111);ax.axis('off');table=[[e['gene'],e['role'],e['limitation']] for e in c['examples']]
  tab=ax.table(cellText=table,colLabels=['示例','展示理由类别','主要限制 / 反证'],cellLoc='left',loc='center',colWidths=[.13,.18,.69]);tab.auto_set_font_size(False);tab.set_fontsize(9);tab.scale(1,1.75)
 save('15_discussion','候选解释：不是最终靶点名单',sel,discuss,'discussion',[c['files']['relations'],c['files']['reader']])
 # Report assembly with explicit text bounding boxes; figures remain native PDF when possible.
 report=fitz.open();markdown=['# '+c['cancer']+' 首版结果展示','',f'输入提交：{c["input_commit"]}；单细胞仅至来源。示例不是最终排名。','']
 text_boxes=[]
 for i,pg in enumerate(pages):
  page=report.new_page(width=842,height=595);page.draw_rect(fitz.Rect(0,0,842,8),color=None,fill=(.03,.5,.55))
  def text(rect,s,size,color=(.12,.19,.25)):
   fit=page.insert_textbox(fitz.Rect(*rect),s,fontname='china-s',fontsize=size,color=color)
   if fit<0:raise RuntimeError(f'Text overflow page {i+1}: '+s[:35])
   text_boxes.append(dict(page=i+1,text=s,rect=list(rect),font_size=size))
  text((30,20,812,52),f'{i+1:02d}  {pg["title"]}',20)
  text((30,59,812,87),'问题｜'+pg['question'],12)
  image_rect=fitz.Rect(30,94,812,438);f=out/pg['figure']
  if f.suffix=='.pdf':
   doc=fitz.open(f);page.show_pdf_page(image_rect,doc,0,keep_proportion=True);doc.close()
  else:page.insert_image(image_rect,filename=str(f),keep_proportion=True)
  text((30,451,410,546),'结果与解释｜'+pg['result'],11)
  text((438,451,812,546),'限制｜'+pg['limit'],11,color=(.48,.28,.16))
  text((30,560,812,584),foot+'  |  '+pg['id']+f'  |  {i+1}/{len(pages)}',8)
  markdown += ['## '+pg['title'],'','**问题：**'+pg['question'],'',f'![{pg["title"]}]({pg["figure"].replace(".pdf",".png")})','','**结果与解释：**'+pg['result'],'','**限制：**'+pg['limit'],'']
 report.save(out/(c['cancer']+'_主报告.pdf'),garbage=4,deflate=True);report.close()
 (out/(c['cancer']+'_主报告.md')).write_bytes(('\n'.join(markdown)+'\n').encode())
 # Complete source appendix; preserve all original strings, including historical versions.
 for key,df in d.items():tsv(out/'appendix'/(key+'.tsv'),df)
 tsv(out/'appendix'/'genes_display_corrected.tsv',z['reader']);tsv(out/'appendix'/'source_status.tsv',z['source_status']);tsv(out/'display_selection.tsv',sel)
 fm=pd.DataFrame(figures);fm['configured_analysis_families']=json.dumps(c['families'],ensure_ascii=False);tsv(out/'figure_manifest.tsv',fm);tsv(out/'source_manifest.tsv',manifest)
 wb=Workbook();wb.remove(wb.active)
 book={'先读我':pd.DataFrame({'说明':['所有统计来自固定提交 '+c['input_commit'],'每层P/q分开；NA不是零；来源状态优先机械最高类别','示例不是排名；完整候选与未显著结果保留','原输入表在附录TSV中逐值保留；原始矩阵未读取','单细胞先按供者汇总；最高表达不证明功能','输入哈希与版本见来源清单']}),'基因易读表':z['reader'],'基因完整历史':d['genes'],'完整关系':d['relations'],'代谢物全量':d['metabolites'],'配对RNA':d.get('rna',pd.DataFrame({'状态':['无可用正常RNA']})),'单细胞全量':z['profiles'] if len(z['profiles']) else pd.DataFrame({'状态':['未提供']}),'来源状态':z['source_status'] if len(z['source_status']) else pd.DataFrame({'状态':['未提供']}),'分型全量':d.get('subtype_profiles',pd.DataFrame({'状态':['无分型']})),'分型排名':d.get('subtypes',pd.DataFrame({'状态':['无分型']})),'展示选择':sel,'图源索引':fm,'输入来源':manifest}
 for key in ['normal_reference_profiles','normal_reference_coverage','old_source_comparison','new_source_comparison','subtype_coverage','subtype_within_celltype']:
  if key in d:book[key[:31]]=d[key]
 for name,df in book.items():
  ws=wb.create_sheet(name);ws.append(list(df.columns))
  for row in df.itertuples(index=False,name=None):
   vals=[]
   for v in row:
    if pd.isna(v):vals.append('NA')
    elif isinstance(v,str) and v.startswith(('=','+','@')):vals.append("'"+v)
    else:vals.append(v)
   ws.append(vals)
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='087F8C');cell.alignment=Alignment(wrap_text=True)
  ws.row_dimensions[1].height=32
  for col in ws.columns:
   letter=col[0].column_letter;length=max(len(str(cell.value or '')) for cell in list(col)[:30]);ws.column_dimensions[letter].width=min(42,max(13,length*.85))
 wb.save(out/(c['cancer']+'_完整结果册.xlsx'))
 # Self-contained code/config package with explicit repo prerequisite for immutable Git reads.
 shutil.copyfile(__file__,out/'repro'/'camp_results_report.py');shutil.copyfile(config_path,out/'repro'/Path(config_path).name)
 test=ROOT/'tests/test_camp_results_report.py'
 if test.exists():shutil.copyfile(test,out/'repro'/'test_camp_results_report.py')
 versions=dict(python=platform.python_version(),numpy=np.__version__,pandas=pd.__version__,matplotlib=matplotlib.__version__,pymupdf=fitz.VersionBind)
 dump(out/'software_versions.json',versions);dump(out/'analysis_spec.json',c)
 dep=['numpy','pandas','matplotlib','PyYAML','openpyxl','PyMuPDF']
 import importlib.metadata
 (out/'repro'/'requirements.txt').write_text('\n'.join(x+'=='+importlib.metadata.version(x) for x in dep)+'\n',encoding='utf-8')
 (out/'repro'/'README_CN.md').write_text('将通用脚本放入仓库code，配置放入configs。必须保留配置指向的Git提交。运行 --validate-only只读；--out指定全新目录。缺失的可选RNA/分型/第二单细胞研究在配置关闭或省略，不补跑分析。通用代码不计算P/q。字体由本机提供，不随包分发。\n',encoding='utf-8')
 # Render every main page and create contact sheets for visual review.
 pdf=fitz.open(out/(c['cancer']+'_主报告.pdf'));from PIL import Image,ImageOps,ImageDraw
 for i,page in enumerate(pdf):page.get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(out/'rendered_pages'/f'page_{i+1:02d}.png')
 for start in range(0,len(pdf),6):
  sheet=Image.new('RGB',(1260,960),'#DDE3E8')
  for j in range(min(6,len(pdf)-start)):
   im=Image.open(out/'rendered_pages'/f'page_{start+j+1:02d}.png');im.thumbnail((620,300));sheet.paste(im,((j%2)*630,(j//2)*320+18));ImageDraw.Draw(sheet).text(((j%2)*630+8,(j//2)*320),str(start+j+1),fill='black')
  sheet.save(out/'rendered_pages'/f'contact_{start//6+1}.png')
 actual_pages=len(pdf);pdf.close();dump(out/'layout_boxes.json',text_boxes)
 for key,df in d.items():
  again=pd.read_csv(out/'appendix'/(key+'.tsv'),sep='\t',dtype=str,keep_default_na=False);assert df.equals(again),key
 modules=dict(rna='DONE' if len(z['rna']) else 'NOT_EVALUABLE',subtypes='DONE' if 'subtypes' in d and c.get('subtypes') else 'NOT_RUN',single_cell='DONE' if len(z['profiles']) else 'NOT_RUN')
 validation=validate(c,d,z);validation.update(status='BUILT_NUMERICALLY_VALIDATED',report_pages=actual_pages,new_figures=sum(x['format']=='PNG;PDF;SVG' for x in figures),reused_assets=sum(x['format'].startswith('REUSED') for x in figures),appendix_tables_preserved=True,text_boxes_fit=True,modules=modules,visual_review='PENDING_HUMAN_OR_AGENT_RENDER_REVIEW',cross_cancer_real_data_validation='NOT_RUN')
 dump(out/'validation.json',validation)
 (out/'README_CN.md').write_text(f'# {c["cancer"]}首版结果展示\n\n## 本轮问题\n将现有结果生成可读报告和完整附录，不重算。\n## 输入与范围\n固定提交 {c["input_commit"]}。\n## 实际结果\n主报告{actual_pages}页；新图{validation["new_figures"]}张；复用{validation["reused_assets"]}件。详见validation与清单。\n## 新手解释\n先读主报告，再按完整关系ID查Excel和附录；精选不是排名。\n## 限制/反证\n来源状态优先机械排名。外部未支持、身份歧义和覆盖缺项保留；其他癌种未实际运行。\n## 当前决定\n单细胞停在来源，不新增机制分析。\n## 下一步\n导师讨论；不自动加统计。\n## 复现\npython code/camp_results_report.py --config configs/{Path(config_path).name} --validate-only\n\npython code/camp_results_report.py --config configs/{Path(config_path).name} --out <全新目录>\n\nZIP见同级目录。附录原图册在本地完整包；GitHub按来源提交链接复用，避免重复上传大文件。\n',encoding='utf-8')
 return validation

def package(out):
 files=sorted(p for p in out.rglob('*') if p.is_file() and p.name!='checksums.tsv')
 manifest=pd.DataFrame([dict(path=p.relative_to(out).as_posix(),bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in files]);tsv(out/'checksums.tsv',manifest)
 target=out.parent/(out.name+'_完整展示包.zip')
 with zipfile.ZipFile(target,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=5) as z:
  for p in files+[out/'checksums.tsv']:z.write(p,arcname=out.name+'/'+p.relative_to(out).as_posix())
 with zipfile.ZipFile(target) as z:
  assert z.testzip() is None
  for _,r in manifest.iterrows():assert hashlib.sha256(z.read(out.name+'/'+r.path)).hexdigest()==r.sha256
 return target
def main():
 ap=argparse.ArgumentParser();ap.add_argument('--config',required=True);ap.add_argument('--out');ap.add_argument('--validate-only',action='store_true');a=ap.parse_args()
 c,d,b,m=load(a.config);z=prepare(c,d);v=validate(c,d,z)
 if a.validate_only:print(json.dumps(v,ensure_ascii=False));return
 if not a.out:ap.error('--out is required for rendering')
 out=Path(a.out);result=build(c,d,z,b,m,out,a.config);target=package(out);print(json.dumps(dict(result,zip=str(target)),ensure_ascii=False))
if __name__=='__main__':main()
