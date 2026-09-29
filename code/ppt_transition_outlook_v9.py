from pathlib import Path
import ast,shutil,json,hashlib
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
ROOT=Path(__file__).resolve().parents[1];RUN='20260929T_transition_outlook_v9';O=ROOT/'results/BRCA/07_INTEGRATION'/RUN;O.mkdir(exist_ok=True)
V=Path('D:/CodexData/visualizations/2026/09/29/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/CAMP_transition_outlook_v9');V.mkdir(exist_ok=True)
SRC=V.parent/'CAMP_reaction_inplace_v8/CAMP_反应示意原位替换_25页.pptx';R=Presentation(SRC)
mod=ast.parse((ROOT/'code/ppt_editable_main_v1.py').read_text(encoding='utf-8'));ns=dict(G='00553B',INK='183D34',M='6A7973',PALE='F0F5F1',LINE='DDE7E0',FONT='Microsoft YaHei')
exec(compile(ast.Module(body=[n for n in mod.body if isinstance(n,(ast.Import,ast.ImportFrom,ast.FunctionDef))],type_ignores=[]),'<helpers>','exec'),ns)
txt,box,line,el,notes=[ns[n] for n in ['txt','box','line','el','notes']];C=ns['PP_ALIGN'].CENTER;G='00553B';M='6A7973'
LOGO=ROOT/'results/BRCA/07_INTEGRATION/20260929T_illustrated_main_v2/assets/sysu_logo.png'
ART=ROOT/'results/BRCA/07_INTEGRATION/20260929T_case_art_v4/assets/scientific_cutouts.png'
def new(title,sub):
 s=R.slides.add_slide(R.slide_layouts[6]);s.shapes.add_picture(str(LOGO),Inches(.65),Inches(.13),width=Inches(1.93),height=Inches(.75));txt(s,10.02,.42,2.62,.28,'CAMP  /  组会汇报',11,M,align=ns['PP_ALIGN'].RIGHT)
 line(s,.65,.94,12.67,.94);line(s,.65,.94,1.40,.94,G,2);txt(s,.67,1.18,12,.61,title,27,G,True);txt(s,.69,1.92,12,.38,sub,14,M);return s
def art(s,rect,x,y,w,h):
 iw,ih=Image.open(ART).size;l,t,r,b=rect;f=min(w/(r-l),h/(b-t));dw,dh=(r-l)*f,(b-t)*f;p=s.shapes.add_picture(str(ART),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh));p.crop_left=l/iw;p.crop_right=1-r/iw;p.crop_top=t/ih;p.crop_bottom=1-b/ih
s=new('从跨癌共同候选，走向癌种内的具体问题','LYPLA1 展示跨癌重复；接下来关注候选在什么代谢关系、什么细胞背景中出现')
box(s,.69,2.59,11.97,.82,'EAF2EC');txt(s,.93,2.85,2.95,.34,'跨癌发现 → 癌种背景',14,G,True)
txt(s,4.03,2.84,8.28,.38,'同样的筛选框架，可能指向不同的生化作用与细胞来源。',17,G)
cards=[('BRCA','ASNS','氨基酸合成','谷氨酰胺 → 天冬酰胺','上皮表达是否支持组织层线索？',(12,550,510,965)),('PDAC / PRAD','SLC6A6','牛磺酸转运','同一转运体，组织变化方向不同','细胞来源能否帮助理解差异？',(530,530,1020,980)),('COAD','UCKL1','核苷补救合成','尿苷 → UMP','代谢关联与上皮定位是否衔接？',(520,45,995,470))]
for i,(c,g,role,desc,q,rect) in enumerate(cards):
 x=.69+i*4.10;box(s,x,3.77,3.76,2.64,'F5F8F5');txt(s,x+.22,3.98,2.5,.32,c,13,M,True);txt(s,x+.22,4.48,2.10,.49,g,27,G,True);art(s,rect,x+2.69,4.09,.83,.94)
 txt(s,x+.22,5.10,3.30,.32,role,17,G,True);txt(s,x+.22,5.59,3.3,.32,desc,11.9,M);txt(s,x+.22,6.08,3.32,.32,q,11.8,G)
line(s,.69,6.83,12.66,6.83);txt(s,.69,7.01,11.54,.31,'下面三个案例用于展示癌种背景；尚不能称为已证实的“癌种特异靶点”。',11,M)
notes(s,'衔接LYPLA1证据汇总与三个代表案例。ASNS-BRCA、SLC6A6-PDAC/PRAD、UCKL1-COAD是癌种背景下的候选解释，并非已建立癌种特异机制。延用已展示结果，没有新增分析。')
ids=R.slides._sldIdLst;it=ids[-1];ids.remove(it);ids.insert(18,it)
s=new('下一步：从候选线索走向可检验的机制假说','已经连接代谢异常、基因关系与细胞背景；后续重点是独立重复与功能方向')
box(s,.69,2.60,11.97,.80,'EAF2EC');txt(s,.94,2.82,11.4,.38,'优先推进 LYPLA1 的 BRCA 线索；其他候选按癌种与细胞背景分别评估。',19,G,True)
items=[('01','聚焦优先关系','LYPLA1 · BRCA','围绕 LPC / GPC / 油酸梳理脂质线索','明确底物、产物和需要验证的作用方向'),('02','补强独立证据','独立队列 · 患者重复','检验效应方向能否跨队列重复','区分癌细胞表达与细胞组成的影响'),('03','检验功能联系','扰动证据 · 代谢读出','优先检索可用的公开功能扰动数据','具备实验条件后测脂质变化与酶活')]
for i,(num,title,tag,a,b) in enumerate(items):
 x=.69+i*4.10;txt(s,x,3.91,.66,.57,num,30,'A2B9AC',True);txt(s,x+.82,3.98,2.95,.42,title,21,G,True);line(s,x,4.61,x+3.76,4.61)
 txt(s,x,4.91,3.76,.36,tag,15,G,True);txt(s,x,5.49,3.70,.52,a,14,M);txt(s,x,6.09,3.70,.53,b,14,M)
line(s,.69,6.94,12.66,6.94);txt(s,.69,7.11,11.6,.25,'需要回答的核心问题：改变候选基因，能否引起预期的代谢变化？',12,G,True)
notes(s,'展望为未来计划，不表示新增分析已完成。不假设现有实验条件，优先公开功能数据；实验性验证为条件具备后的方向。RNA变化不能替代蛋白、酶活、通量；关联不代表因果。')
for n,sl in enumerate(R.slides,1):
 if n<19:continue
 for sh in list(sl.shapes):
  if sh.has_text_frame and sh.text.strip().isdigit() and sh.left>Inches(11.9) and sh.top>Inches(6.9):sl.shapes._spTree.remove(sh._element)
 txt(sl,12.12,7.03,.52,.27,str(n),11,G,True,ns['PP_ALIGN'].RIGHT)
 for sh in sl.shapes:
  for e in sh._element.xpath('.//a:effectRef'):e.set('idx','0')
  for rp in sh._element.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr'):
   for tag,face in [('a:latin','Microsoft YaHei'),('a:ea','微软雅黑')]:
    old=rp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag.split(':')[1])
    if old is not None:rp.remove(old)
    rp.append(el(tag,typeface=face))
assert len(R.slides)==27
PPT='CAMP_过渡与展望完整版_27页.pptx';R.save(O/PPT);shutil.copy2(O/PPT,V/PPT)
(O/'source_manifest.json').write_text(json.dumps(dict(input=str(SRC),sha256=hashlib.sha256(SRC.read_bytes()).hexdigest(),added_slides=[19,27],no_new_statistics=True),ensure_ascii=False,indent=2),encoding='utf-8')
print(O/PPT)
