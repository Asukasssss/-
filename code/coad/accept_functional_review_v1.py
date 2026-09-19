"""Validate supplied evidence, append reviewed annotations, preserve every source cell."""
import argparse,csv,copy,hashlib,io,json,platform,subprocess,zipfile
from collections import Counter
from datetime import datetime,timezone
from pathlib import Path
from openpyxl import load_workbook
from integrate_functional_review_nominal35_v1 import enrich,read_tsv
ROOT=Path(__file__).resolve().parents[2]
def sha(data):return hashlib.sha256(data).hexdigest()
def write(p,data):
    p.parent.mkdir(parents=True,exist_ok=True)
    with p.open('xb') as f:f.write(data if isinstance(data,bytes) else data.encode('utf-8'))
def dump(p,x):write(p,json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
def table(p,fields,rows):
    text=io.StringIO(newline='');w=csv.DictWriter(text,fields,delimiter='\t',lineterminator='\n');w.writeheader();w.writerows(rows);write(p,text.getvalue())
def blob(b):return hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()
def key(r):return r['metabolite_key'],r['gene']
def main():
    ap=argparse.ArgumentParser();ap.add_argument('--delivery-zip',type=Path,required=True);ap.add_argument('--workbook',type=Path,required=True);ap.add_argument('--metadata-cache',type=Path,required=True);ap.add_argument('--run-id',required=True);a=ap.parse_args()
    out=(ROOT/'results/COAD/05_FUNCTION'/a.run_id).resolve();assert out.parent==(ROOT/'results/COAD/05_FUNCTION').resolve()
    out.mkdir(parents=True,exist_ok=False);write(out/'.running',a.run_id+'\n')
    config=ROOT/'config/coad_function_review_acceptance_v1.json';accept=json.loads(config.read_text(encoding='utf-8'))
    z=zipfile.ZipFile(a.delivery_zip);prefix='COAD_35_functional_review_delivery/'
    members={i.filename[len(prefix):]:z.read(i) for i in z.infolist() if not i.is_dir() and i.filename.startswith(prefix)}
    for name in members:assert not Path(name).is_absolute() and '..' not in Path(name).parts
    manifest=list(csv.DictReader(io.StringIO(members['MANIFEST_SHA256.tsv'].decode('utf-8-sig')),delimiter='\t'))
    for r in manifest:assert sha(members[r['path']])==r['sha256'] and len(members[r['path']])==int(r['bytes'])
    assert len(manifest)==17 and members[a.workbook.name]==a.workbook.read_bytes()
    source_spec=json.loads(members['analysis_spec.json']);original=json.loads(members['results/gene_function_review.json']);studies=json.loads(members['results/study_evidence.json'])
    assert len(original)==35 and len({r['gene'] for r in original})==35 and len(studies)==40
    np,fp=ROOT/source_spec['source_nominal_path'],ROOT/source_spec['source_full_catalog_path']
    assert blob(np.read_bytes())==source_spec['source_nominal_git_blob_sha1'] and blob(fp.read_bytes())==source_spec['source_full_catalog_git_blob_sha1']
    nf,nr=read_tsv(np);ff,fr=read_tsv(fp);assert len(nr)==39 and len(fr)==974
    genes={r['gene']:r for r in original};assert set(genes)=={r['gene'] for r in nr}
    wb=load_workbook(io.BytesIO(members[a.workbook.name]),read_only=True,data_only=True)
    grow=list(wb['35基因比较'].iter_rows(min_row=4,values_only=True));assert len(grow)==35
    for r in grow:assert r[0] in genes and r[1]==genes[r[0]]['research_decision_cn'] and r[2]==genes[r[0]]['evidence_class_cn']
    erow=list(wb['文献与反证'].iter_rows(min_row=4,values_only=True));assert len(erow)==40 and {r[0] for r in erow}=={r['evidence_id'] for r in studies}
    display=list(wb['39关系摘要'].iter_rows(min_row=4,values_only=True));assert len(display)==39
    pairs={key(r):r for r in nr}
    for r in display:
        src=pairs[(r[0],r[2])];assert int(r[3])==int(src['n'])
        for i,k in [(4,'effect'),(5,'p_value'),(6,'q_value')]:assert abs(float(r[i])-float(src[k]))<=.000051
    wb.close()
    meta_path=a.metadata_cache/'europepmc_core.json';meta=json.loads(meta_path.read_text(encoding='utf-8'))['resultList']['result'];assert len(meta)==40
    verified=[];detailed={'E03','E04','E08','E18','E21','E23'}
    for r in studies:
        match=[m for m in meta if (r['PMID'] and m['id']==r['PMID']) or (not r['PMID'] and m.get('doi','').lower()==r['DOI'].lower())];assert len(match)==1
        m=match[0];assert m.get('abstractText') and str(m['pubYear'])==str(r['year'])
        if r['DOI']:assert m['doi'].lower()==r['DOI'].lower()
        verified.append(dict(evidence_id=r['evidence_id'],gene=r['gene'],verified_PMID=m['id'],verified_DOI=m.get('doi','NA'),verified_title=m['title'],publication_year=m['pubYear'],source_url=r['source_url'],pmcid=m.get('pmcid','NA'),independent_read_depth='ABSTRACT_AND_SELECTED_FULLTEXT' if r['evidence_id'] in detailed else 'ABSTRACT',assessment='INTENTIONAL_OTHER_GENE_EXCLUSION_CONFIRMED' if r['evidence_id'] in ['E39','E40'] else 'CATEGORY_CONSISTENT_WITH_READ_SCOPE',supplement_locator=accept['supplements'].get(r['gene'],{}).get('locator','NA'),linked_comments=json.dumps(m.get('commentCorrectionList',{}),ensure_ascii=False),integrity_limit='Index metadata checked; not a systematic retraction/image audit'))
    # Keep the supplied text package as provenance. The workbook/archive and fetched articles stay local.
    for name,data in members.items():
        if Path(name).suffix.lower() in ['.md','.tsv','.json','.py']:write(out/'imported_delivery'/name,data)
    active=copy.deepcopy(original)
    for r in active:
        supplement=accept['supplements'].get(r['gene'],{})
        if r['gene']=='BCAT2':
            for k in ['functional_conclusion_cn','counterevidence_or_limit_cn','metabolic_link_cn','next_action_cn']:r[k]=supplement[k]
        r['acceptance_supplement_cn']=supplement.get('supplement_cn','详见BCAT2条件性补充。' if r['gene']=='BCAT2' else '')
        r['acceptance_review_status']='ABSTRACT_RECORDS_CHECKED' if r['evidence_ids'] else 'IMPORTED_SCOPED_SEARCH_NOT_INDEPENDENTLY_REPEATED'
        r['acceptance_version']=accept['version']
    review={r['gene']:r for r in active};selected=set(pairs)
    nh,nout=enrich(nf,nr,review,selected);fh,fout=enrich(ff,fr,review,selected)
    newfields=['func35_acceptance_version','func35_acceptance_review_status','func35_acceptance_supplement_cn']
    for rows in [nout,fout]:
        for r in rows:
            gene=review.get(r['gene']);r['func35_acceptance_version']=accept['version'];r['func35_acceptance_review_status']=gene['acceptance_review_status'] if gene else 'NOT_REVIEWED_THIS_BATCH';r['func35_acceptance_supplement_cn']=gene['acceptance_supplement_cn'] if gene else ''
    table(out/'nominal39_annotated_exact.tsv',nh+newfields,nout);table(out/'catalog974_annotated_exact.tsv',fh+newfields,fout)
    table(out/'gene_function_review.tsv',list(active[0]),[{k:'' if v is None else v for k,v in r.items()} for r in active]);dump(out/'gene_function_review.json',active)
    table(out/'source_verification.tsv',list(verified[0]),verified)
    queue=dict(Counter(r['research_decision_cn'] for r in active));assert queue==accept['retained_research_queue']
    summary=dict(run_id=a.run_id,version=accept['version'],genes=35,nominal_relations=39,full_catalog_rows=974,source_records=40,source_identities_matched=40,abstracts_read=40,selected_fulltext_records_checked=6,source_class_A_genes=10,additional_conditional_CRC_gene='BCAT2',broad_CRC_tumor_gene_perturbation_records_unique_genes=11,research_queue=queue,gene_level_annotations_in_full_catalog=sum(r['gene'] in review for r in fr),pair_level_review_scope=39,patient_tests_recomputed=0,depmap='NOT_RUN',external_patient_validation='NOT_RUN',stage_05='PARTIAL',stage_07='PARTIAL')
    dump(out/'summary.json',summary)
    sources=[dict(source_path='user_attachment/'+a.delivery_zip.name,sha256=sha(a.delivery_zip.read_bytes())),dict(source_path='user_attachment/'+a.workbook.name,sha256=sha(a.workbook.read_bytes()))]
    input_files=[np,fp,config,Path(__file__),Path(__file__).with_name('integrate_functional_review_nominal35_v1.py')]
    for p in input_files:sources.append(dict(source_path=p.relative_to(ROOT).as_posix(),sha256=sha(p.read_bytes())))
    fetch=json.loads((a.metadata_cache/'request.json').read_text(encoding='utf-8'));sources.append(dict(source_path=fetch['url'],sha256=sha(meta_path.read_bytes())))
    for r in json.loads((a.metadata_cache/'fulltext_fetch.json').read_text(encoding='utf-8')):sources.append(dict(source_path=r['url'],sha256=r.get('sha256','NA')))
    sources.append(dict(source_path='https://www.nature.com/articles/srep23642;web-tool selected Results/Fig5;no-byte-cache',sha256='NA'))
    table(out/'source_manifest.tsv',['source_path','sha256'],sources)
    spec=dict(accept,run_id=a.run_id,created_utc=datetime.now(timezone.utc).isoformat(),code_commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),source_base_commit=source_spec['base_commit'],source_package_sha256=sha(a.delivery_zip.read_bytes()),source_workbook_sha256=sha(a.workbook.read_bytes()),source_package_version=source_spec['analysis_version'],execution_location='local public aggregates and public literature only',software={'python':platform.python_version()},statistics='All source cells retained as strings, including source run/stage/version/status; current delivery stage recorded in index and func35 columns; no recomputation')
    dump(out/'analysis_spec.json',spec)
    # Validate saved files, all input columns and hash stability; no fixture-only claim.
    for name,fields,source in [('nominal39_annotated_exact.tsv',nf,nr),('catalog974_annotated_exact.tsv',ff,fr)]:
        h,rows=read_tsv(out/name);assert h[:len(fields)]==fields and len(rows)==len(source)
        assert [{k:r[k] for k in fields} for r in rows]==source
        assert len({(r['cohort'],r['metabolite_key'],r['human_gene_id']) for r in rows})==len(rows)
        assert all(r['func35_depmap_score']==r['func35_depmap_n_models']=='' for r in rows)
    assert all(r['func35_pair_scope']=='INITIAL39_NOMINAL_PAIR' for r in nout)
    assert sum(r['func35_pair_scope']=='INITIAL39_NOMINAL_PAIR' for r in fout)==39
    for s in sources[2:2+len(input_files)]:assert sha((ROOT/s['source_path']).read_bytes())==s['sha256']
    public=[p for p in out.rglob('*') if p.is_file() and p.name!='.running']
    dump(out/'validation.json',dict(status='PASS',checks=['17 original manifest entries','separate workbook identical to archive workbook','35 gene/category/queue cells match','40 source IDs and metadata','39 workbook keys/n/rounded statistics match fixed repository','exact pinned source git blobs','every original column unchanged in39 and974 real rows','unique relation keys and original order','gene-level versus39-pair scope distinguished','DepMap remains missing not zero','inputs unchanged'],limitations=['Original workbook not edited or independently visually re-rendered','No systematic re-search of35 genes','No full audit of every figure/replicate','No independent patient validation','Imported script/provenance status snapshots are historical'],public_output_sha256={p.relative_to(out).as_posix():sha(p.read_bytes()) for p in public}))
    (out/'.running').unlink();write(out/'DATA_JOIN_DONE',datetime.now(timezone.utc).isoformat()+'\n');print(json.dumps(summary,ensure_ascii=False))
if __name__=='__main__':main()
