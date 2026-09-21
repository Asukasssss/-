"""Acquire official supplements on server165; export metadata only."""
import pathlib,urllib.request,json,hashlib,subprocess,sys
root=pathlib.Path('/public3/xuzx/Cancer/pancancer_metabolomics_direct_matrix_20260716')
run=pathlib.Path(sys.argv[1]).resolve()
assert run.parent==root/'results/collaborative/COAD/B' and (run/'.running').is_dir()
dest=root/'data/candidates/coad_integral_20260921';dest.mkdir(exist_ok=True)
records=[]
for article in [28156574,28156568]:
 url='https://api.figshare.com/v2/articles/'+str(article)
 with urllib.request.urlopen(url,timeout=40)as h:meta=json.load(h)
 assert meta['resource_doi']=='10.1021/acs.analchem.4c04421'
 (dest/(str(article)+'.json')).write_text(json.dumps(meta))
 for f in meta['files']:
  p=dest/f['name'];assert not p.exists()
  with urllib.request.urlopen(f['download_url'],timeout=120)as h:p.write_bytes(h.read())
  b=p.read_bytes();assert len(b)==f['size'] and hashlib.md5(b).hexdigest()==f['computed_md5']
  records.append(dict(article=article,doi=meta['doi'],url=f['download_url'],path=str(p),bytes=len(b),md5=f['computed_md5'],sha256=hashlib.sha256(b).hexdigest()))
(run/'source_files.json').write_text(json.dumps(records,indent=2))
subprocess.run(['pdftotext','-layout',str(dest/'ac4c04421_si_001.pdf'),str(dest/'methods.txt')],check=True)
import openpyxl
w=openpyxl.load_workbook(dest/'ac4c04421_si_003.xlsx',read_only=True,data_only=True)
summary=[dict(sheet=s.title,rows=s.max_row,columns=s.max_column)for s in w]
(run/'workbook_structure.json').write_text(json.dumps(summary,indent=2))
print(json.dumps(dict(files=records,sheets=summary,openpyxl=openpyxl.__version__)),flush=True)
