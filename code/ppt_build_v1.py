"""Regenerate all quantitative charts and assemble editable-text PPT from frozen sources."""
from pathlib import Path
import json,hashlib,math
import numpy as np,pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.colors import ListedColormap
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
ROOT=Path(__file__).resolve().parents[1];RUN='20260928T150000Z_ppt_redraw_v1';O=ROOT/'results/BRCA/07_INTEGRATION'/RUN
F=O/'figures';D=O/'figure_data';F.mkdir(exist_ok=True);D.mkdir(exist_ok=True)
font_manager.fontManager.addfont(str(ROOT/'runtime/ppt_fonts/NotoSansCJK-Regular.ttc'))
RED='#C94B54';BLUE='#3276A5';TEAL='#258F88';GRAY='#D7DCE2';INK='#172A3A';MUTED='#657583'
plt.rcParams.update({'font.family':'Noto Sans CJK JP','font.size':12,'axes.titlesize':14,'axes.labelsize':12,'xtick.labelsize':11,'ytick.labelsize':11,'pdf.fonttype':42,'ps.fonttype':42,'axes.spines.top':False,'axes.spines.right':False,'axes.unicode_minus':False,'savefig.facecolor':'white'})
S={p.stem:pd.read_csv(p,sep='\t') for p in (O/'sources').glob('*.tsv')};manifest=[]
def pf(x):
 if pd.isna(x):return '未检验'
 return f'{x:.1e}' if x<.001 else f'{x:.4f}'.rstrip('0').rstrip('.')
def save(fig,name,data,sources):
 for ext in ['png','pdf']:fig.savefig(F/(name+'.'+ext),dpi=260,bbox_inches='tight',pad_inches=.12)
 plt.close(fig);pd.DataFrame(data).to_csv(D/(name+'.tsv'),sep='\t',index=False)
 manifest.append(dict(figure=name,source_ids=';'.join(sources),plot_data='figure_data/'+name+'.tsv',new_test=False))
def clean(ax):ax.spines[['top','right']].set_visible(False);ax.set_axisbelow(True);ax.grid(axis='x',color='#EEF1F4',lw=.7)
def rna(c,g):
 d=S[c.lower()+'_gene'];z=d[d.gene==g]
 if z.empty:return dict(cancer=c,gene=g,effect=np.nan,p=np.nan,n=np.nan)
 r=z.iloc[0];prefix='v2_paired_RNA_' if c=='BRCA' else 'RNA_'
 return dict(cancer=c,gene=g,effect=r.get(prefix+'effect',np.nan),p=r.get(prefix+'p_value',r.get(prefix+'p',np.nan)),n=r.get(prefix+'n',np.nan))
def rel(c,g):
 d=S[c.lower()+'_rel'];d=d[d.gene==g].copy();out=[]
 pref={'BRCA':'CAMP_','COAD':'CAMP_','PDAC':'association_','PRAD':'patient_primary_'}[c]
 for _,r in d.iterrows():
  if c=='PDAC' and r.get('mapping_status')!='DIRECT':continue
  out.append(dict(cancer=c,gene=g,metabolite=r.metabolite_name,key=r.metabolite_key,n=r.get(pref+'n',np.nan),rho=r.get(pref+'effect',r.get(pref+'rho',np.nan)),p=r.get(pref+'p_value',r.get(pref+'p',np.nan)),lo=r.get(pref+'ci_lower',np.nan),hi=r.get(pref+'ci_upper',np.nan)))
 return pd.DataFrame(out)
cs=['BRCA','COAD','PDAC','PRAD']
# 01 Current cohort coverage, no new inference.
counts=[];met=[]
for c,name in [('BRCA','brca_met'),('COAD','coad_met'),('GBM','gbm_met'),('PDAC','pdac_met'),('PRAD','prad_met'),('ccRCC','ccrcc_met')]:
 d=S[name].copy();d['cancer_display']=c
 if c=='ccRCC':d['cancer_display']=d.cohort
 for co,z in d.groupby('cancer_display',sort=False):
  e=pd.to_numeric(z.effect,errors='coerce');p=pd.to_numeric(z.p_value,errors='coerce');ok=e.notna()&p.notna()
  counts.append(dict(cohort=co,evaluable=int(ok.sum()),up=int(((p<.05)&(e>0)).sum()),down=int(((p<.05)&(e<0)).sum()),ns=int((ok&(p>=.05)).sum())))
 met.append(d)
counts=pd.DataFrame(counts);allmet=pd.concat(met,ignore_index=True)
fig,ax=plt.subplots(figsize=(10.4,3.8),constrained_layout=True);y=np.arange(len(counts));left=np.zeros(len(counts))
for k,col,lab in [('up',RED,'升高，P<0.05'),('down',BLUE,'降低，P<0.05'),('ns',GRAY,'其余可评估')]:ax.barh(y,counts[k],left=left,color=col,height=.60,label=lab);left+=counts[k]
for i,r in counts.iterrows():ax.text(r.evaluable+12,i,f'{int(r.up+r.down)}/{r.evaluable}',va='center',fontsize=11)
ax.set_yticks(y,counts.cohort);ax.invert_yaxis();ax.set_xlim(0,counts.evaluable.max()*1.16);ax.set_xlabel('当前版本可评估代谢特征数');ax.legend(ncol=3,frameon=False,loc='lower center',bbox_to_anchor=(.5,1.01));clean(ax)
save(fig,'01_coverage',counts,['brca_met','coad_met','pdac_met','prad_met','gbm_met','ccrcc_met'])
# 02 Curated identity-exact molecules, no aggregation of colliding keys.
keys={'KEGG:C00245':'牛磺酸','KEGG:C00670':'游离GPC','KEGG:C04102':'LPC(16:0)','KEGG:C00712':'油酸','KEGG:C00064':'谷氨酰胺','KEGG:C00025':'谷氨酸','KEGG:C00106':'尿嘧啶','KEGG:C00299':'尿苷','KEGG:C00065':'丝氨酸','KEGG:C00037':'甘氨酸','KEGG:C00157':'磷脂酰胆碱','KEGG:C00300':'肌酸','KEGG:C00791':'肌酐','KEGG:C00114':'胆碱','KEGG:C00189':'乙醇胺'}
keys.pop('KEGG:C00157')  # lipid class is not a single molecular species
cols=counts.cohort.tolist();rows=[];cells=np.full((len(keys),len(cols)),np.nan)
for i,(key,label) in enumerate(keys.items()):
 for j,c in enumerate(cols):
  z=allmet[(allmet.cancer_display==c)&(allmet.metabolite_key==key)]
  if len(z)!=1:rows.append(dict(metabolite=label,key=key,cohort=c,status='MISSING_OR_ID_COLLISION',p=np.nan,effect=np.nan));continue
  r=z.iloc[0];p=float(r.p_value);ef=float(r.effect);cells[i,j]=2 if p<.05 and ef>0 else 0 if p<.05 and ef<0 else 1
  rows.append(dict(metabolite=label,key=key,cohort=c,status='EVALUABLE',p=p,effect=ef))
fig,ax=plt.subplots(figsize=(10.4,5.1),constrained_layout=True);cm=ListedColormap([BLUE,'#F0F2F5',RED]);cm.set_bad('#C1C7CF');ax.imshow(np.ma.masked_invalid(cells),cmap=cm,vmin=0,vmax=2,aspect='auto')
for i in range(len(keys)):
 for j in range(len(cols)):
  v=cells[i,j];ax.text(j,i,'↑' if v==2 else '↓' if v==0 else '·' if v==1 else 'NA',ha='center',va='center',color='white' if v in [0,2] else MUTED,fontsize=12)
ax.set_xticks(range(len(cols)),cols);ax.set_yticks(range(len(keys)),list(keys.values()));ax.tick_params(length=0);ax.spines[:].set_visible(False)
save(fig,'02_pan_metabolites',rows,['brca_met','coad_met','pdac_met','prad_met','gbm_met','ccrcc_met'])
# 03 Candidate RNA evidence matrix; fixed discussion panel, not rank.
genes=['LYPLA1','SLC7A11','ASNS','SLC6A6','AHCY'];rr=pd.DataFrame([rna(c,g) for g in genes for c in cs]);fig,ax=plt.subplots(figsize=(10.2,3.8),constrained_layout=True)
a=np.array([0 if pd.isna(r.p) else 1 if r.p>=.05 else 2 if r.effect>0 else -2 for r in rr.itertuples()]).reshape(5,4)
display=np.select([a==-2,a==0,a==1,a==2],[0,1,2,3])
ax.imshow(display,cmap=ListedColormap([BLUE,'#C1C7CF','#F4F6F8',RED]),vmin=0,vmax=3,aspect='auto')
for i,g in enumerate(genes):
 for j,c in enumerate(cs):
  r=rr[(rr.gene==g)&(rr.cancer==c)].iloc[0];label='未纳入当前表' if pd.isna(r.p) else ('↑' if r.effect>0 else '↓')+'  P='+pf(r.p)
  ax.text(j,i,label,ha='center',va='center',color='white' if pd.notna(r.p) and r.p<.05 else INK,fontsize=12)
ax.set_xticks(range(4),cs);ax.set_yticks(range(5),genes);ax.tick_params(length=0);ax.spines[:].set_visible(False)
save(fig,'03_candidate_rna',rr,[c.lower()+'_gene' for c in cs])
# 04 Focus RNA p and direction, avoid incomparable effect axes.
z=rr[rr.gene=='LYPLA1'].copy();fig,ax=plt.subplots(figsize=(10.2,3.5),constrained_layout=True)
ax.barh(z.cancer,-np.log10(z.p),color=[RED if p<.05 else '#E9B8BC' for p in z.p],height=.55);ax.invert_yaxis();ax.axvline(-np.log10(.05),color=MUTED,ls='--',lw=1)
for i,r in enumerate(z.itertuples()):ax.text(-np.log10(r.p)+.08,i,f'↑  P={pf(r.p)}  |  {int(r.n)}对',va='center')
ax.set_xlim(0,7.3);ax.set_xlabel('−log10(P)；虚线为名义 P=0.05');clean(ax)
save(fig,'04_lypla1_rna',z,[c.lower()+'_gene' for c in cs])
# 05 Same metabolite versus LYPLA1 across cancers.
lr=pd.concat([rel(c,'LYPLA1') for c in cs],ignore_index=True);ms=[('KEGG:C04102','LPC(16:0)'),('KEGG:C00670','游离GPC'),('KEGG:C00712','油酸')]
fig,axs=plt.subplots(1,3,figsize=(11.2,3.6),sharey=True,constrained_layout=True)
for ax,(key,label) in zip(axs,ms):
 ax.axvline(0,color=GRAY,lw=1);ax.set_title(label,loc='left');ax.set_xlim(-.75,.75);ax.set_yticks(range(4),cs);ax.set_ylim(3.6,-.6);ax.set_xlabel('Spearman ρ');clean(ax)
 for i,c in enumerate(cs):
  q=lr[(lr.cancer==c)&(lr.key==key)]
  if len(q)!=1:ax.text(-.68,i,'当前关系池未覆盖',va='center',color=MUTED,fontsize=9);continue
  r=q.iloc[0];col=RED if r.rho>0 else BLUE
  if pd.notna(r.lo) and pd.notna(r.hi):ax.plot([r.lo,r.hi],[i,i],color=col,lw=1.5)
  ax.scatter(r.rho,i,s=65,facecolors=col if r.p<.05 else 'white',edgecolors=col,zorder=3)
  ax.text(-.70,i+.30,f'ρ={r.rho:.2f}  P={pf(r.p)}  n={int(r.n)}',fontsize=9)
save(fig,'05_lypla1_relations',lr,[c.lower()+'_rel' for c in cs])
# 06 paired change sensitivity, original CI.
z=S['brca_delta'];fig,ax=plt.subplots(figsize=(10.2,3.4),constrained_layout=True);lab=[]
for i,r in enumerate(z.itertuples()):
 col=RED if r.effect>0 else BLUE;ax.plot([r.ci_lower,r.ci_upper],[i,i],c=col,lw=2);ax.scatter(r.effect,i,s=65,facecolors=col if r.p_value<.05 else 'white',edgecolors=col,zorder=3)
 lab.append(('LPC(16:0)' if r.metabolite_key=='KEGG:C04102' else '游离GPC')+' | '+('主分析' if r.analysis_type=='processed' else '可用值敏感性'))
 ax.text(.62,i,f'n={r.n}  P={pf(r.p_value)}',va='center')
ax.set_yticks(range(len(z)),lab);ax.set_ylim(3.5,-.5);ax.set_xlim(-.5,.98);ax.axvline(0,c=GRAY);ax.set_xlabel('同患者肿瘤−正常变化量的 Spearman ρ（原95%区间）');clean(ax)
save(fig,'06_delta_sensitivity',z,['brca_delta'])
# 07 paired patient single-cell effect; no pooling with descriptive cohorts.
br=S['brca_sc'];br=br[(br.gene=='LYPLA1')&(br.cohort=='Wu2021')].copy();pdac=S['pdac_sc'];pdac=pdac[(pdac.unit=='paired_author_pid')&(pdac.treatment=='Untreated')&(pdac.reference=='Ductal')].copy()
z=pd.concat([br.assign(label='BRCA Wu | '+br.analysis_type),pdac.assign(label='PDAC | '+pdac.population)],ignore_index=True)
fig,ax=plt.subplots(figsize=(10.5,3.6),constrained_layout=True)
for i,r in enumerate(z.itertuples()):
 col=RED if r.effect>0 else BLUE;ax.plot([r.ci_lower,r.ci_upper],[i,i],c=col,lw=2);ax.scatter(r.effect,i,s=60,facecolors=col if r.p_value<.05 else 'white',edgecolors=col,zorder=3);ax.text(.50,i,f'{int(r.n)}对  P={pf(r.p_value)}',va='center',fontsize=11)
labs=[s.replace('ALL','全部供者').replace('AUTHOR_NAIVE','作者Naive标签').replace('All_nuclei','全部细胞核').replace('LYPLA1_positive_nuclei','仅阳性细胞核') for s in z.label]
ax.set_yticks(range(len(z)),labs);ax.set_ylim(len(z)-.5,-.5);ax.axvline(0,c=GRAY);ax.set_xlim(-.28,.94);ax.set_xlabel('恶性−非恶性：患者平均 log1p(CP10K) 差（原95%区间）');clean(ax)
save(fig,'07_sc_paired',z,['brca_sc','pdac_sc'])
# 08 Detection decomposition, each dataset independently described.
pr=S['prad_expr'];pp=S['prad_positive'];co=S['coad_sc'];pdex=S['pdac_expr'];decomp=[]
for _,r in pr.iterrows():decomp.append(dict(cohort='PRAD',group='癌上皮' if 'Malignant' in r.celltype else '非恶性上皮',mean=r.pooled_mean,detect=r.positive_fraction,positive=pp.loc[pp.celltype==r.celltype,'positive_mean_log1p_CP10K'].iloc[0]))
for group,lab in [('normal_reference','正常组织上皮'),('tumor_CNA','CNA上皮')]:
 r=co[co.group==group].iloc[0];decomp.append(dict(cohort='COAD',group=lab,mean=r.effect,detect=r.detection_fraction,positive=r.positive_cell_mean_log1p_CP10K))
for group,lab in [('Ductal','非恶性导管'),('Malignant','恶性上皮')]:
 r=pdex[(pdex.treatment=='Untreated')&(pdex.celltype==group)&(pdex.population=='All_nuclei')].iloc[0];rp=pdex[(pdex.treatment=='Untreated')&(pdex.celltype==group)&(pdex.population=='LYPLA1_positive_nuclei')].iloc[0];decomp.append(dict(cohort='PDAC',group=lab,mean=r.pooled_mean,detect=r.detection_fraction,positive=rp.pooled_mean))
decomp=pd.DataFrame(decomp);fig,axs=plt.subplots(1,3,figsize=(11,3.7),constrained_layout=True)
for ax,(metric,title) in zip(axs,[('mean','全部细胞 / 核均值'),('detect','检出比例（%）'),('positive','仅阳性细胞 / 核均值')]):
 for i,c in enumerate(['COAD','PDAC','PRAD']):
  d=decomp[decomp.cohort==c];normal=d[d.group.str.contains('非恶性|正常')].iloc[0];mal=d[~d.group.str.contains('非恶性|正常')].iloc[0]
  for off,r,col in [(-.17,normal,BLUE),(.17,mal,RED)]:
   val=r[metric]*(100 if metric=='detect' else 1);ax.bar(i+off,val,width=.30,color=col);ax.text(i+off,val+(.025 if metric!='detect' else 1),f'{val:.2f}' if metric!='detect' else f'{val:.1f}',ha='center',fontsize=9)
 ax.set_xticks(range(3),['COAD','PDAC','PRAD']);ax.set_title(title,loc='left');ax.spines[['top','right']].set_visible(False);ax.set_ylim(0,100 if metric=='detect' else 1 if metric=='mean' else 1.9)
save(fig,'08_detection',decomp,['coad_sc','pdac_expr','prad_expr','prad_positive'])
# 09 Spatial patient summary, distinct scales are not pooled.
sp=S['pdac_spatial'];z=sp[(sp.contrast=='PDAC_vs_ND')].copy();fig,ax=plt.subplots(figsize=(10.4,3.4),constrained_layout=True)
for i,r in enumerate(z.itertuples()):
 ax.plot([r.ci_lower,r.ci_upper],[i,i],color=RED,lw=2);ax.scatter(r.effect,i,s=70,facecolors=RED if r.p_value<.05 else 'white',edgecolors=RED);ax.text(1.5,i,f'{int(r.positive_patients)}/6方向为正  P={pf(r.p_value)}',va='center')
ax.set_yticks(range(3),['Q3标准化（主分析）','总量标准化（敏感性）','作者VST（敏感性）']);ax.set_ylim(2.5,-.5);ax.set_xlim(0,2.25);ax.axvline(0,c=GRAY);ax.set_xlabel('癌上皮−正常导管：各处理尺度差值（不比较幅度大小）');clean(ax)
save(fig,'09_geomx',z,['pdac_spatial'])
# 10 Cell-state expression, source-defined labels, no new clusters.
z=S['brca_states'];z=z[(z.cohort=='Wu2021')&(z.status=='DONE')].sort_values('mean_log',ascending=False);fig,axs=plt.subplots(1,2,figsize=(10.5,3.6),constrained_layout=True)
labs=z.state.str.replace('Cancer ','').str.replace(' SC','')
axs[0].barh(labs,z.mean_log,color=RED,height=.58);axs[0].invert_yaxis();axs[0].set_xlabel('患者等权均值 log1p(CP10K)')
axs[1].barh(labs,z.fraction_above_own_malignant_mean*100,color=TEAL,height=.58);axs[1].invert_yaxis();axs[1].set_xlim(0,115);axs[1].set_xlabel('高于自身癌细胞平均值的患者比例（%）')
for i,r in enumerate(z.itertuples()):axs[1].text(r.fraction_above_own_malignant_mean*100+2,i,f'n={r.n_patients_centered}',va='center',fontsize=10)
for ax in axs:clean(ax)
save(fig,'10_states',z,['brca_states'])
# 11 UCKL1 correlation and marginal directions from exact published relation.
u=rel('COAD','UCKL1');u=u[u.key=='KEGG:C00299'];assert len(u)==1
fig,axs=plt.subplots(1,2,figsize=(10.5,3.3),constrained_layout=True);r=u.iloc[0];axs[0].plot([r.lo,r.hi],[0,0],color=BLUE,lw=2);axs[0].scatter(r.rho,0,color=BLUE,s=85);axs[0].axvline(0,c=GRAY);axs[0].set_xlim(-.85,.3);axs[0].set_yticks([0],['UCKL1—尿苷']);axs[0].set_xlabel('Spearman ρ');axs[0].set_title(f'ρ={r.rho:.3f}  P={pf(r.p)}  n={int(r.n)}',loc='left');clean(axs[0])
cr=S['coad_rel'];ur=cr[(cr.gene=='UCKL1')&(cr.metabolite_key=='KEGG:C00299')].iloc[0]
directions=pd.DataFrame([dict(label='尿苷',up=ur.metabolite_paired_pairs_higher,down=ur.metabolite_paired_pairs_lower),dict(label='UCKL1 RNA',up=ur.RNA_n_up,down=ur.RNA_n_down)])
axs[1].barh(directions.label,directions.up,color=RED,label='升高');axs[1].barh(directions.label,directions.down,left=directions.up,color=BLUE,label='降低');axs[1].set_xlim(0,33);axs[1].set_xlabel('各自方向人数（不是共同变化人数）');axs[1].legend(ncol=2,frameon=False,loc='lower center',bbox_to_anchor=(.5,1.01))
save(fig,'11_uckl1',u,['coad_rel'])
# 12 Taurine example, full primary relation comparison.
ta=pd.concat([rel(c,'SLC6A6') for c in ['PDAC','PRAD']],ignore_index=True);ta=ta[ta.key=='KEGG:C00245'];fig,ax=plt.subplots(figsize=(10.5,3),constrained_layout=True)
for i,r in enumerate(ta.itertuples()):
 ax.plot([r.lo,r.hi],[i,i],color=RED,lw=2);ax.scatter(r.rho,i,s=75,edgecolors=RED,facecolors=RED if r.p<.05 else 'white');ax.text(.91,i,f'n={int(r.n)}  P={pf(r.p)}',va='center')
ax.set_yticks(range(len(ta)),ta.cancer);ax.set_ylim(1.6,-.6);ax.set_xlim(-.3,1.36);ax.set_xticks([-.2,0,.2,.4,.6,.8]);ax.axvline(0,c=GRAY);ax.set_xlabel('牛磺酸—SLC6A6：肿瘤内 Spearman ρ（原95%区间）');clean(ax)
save(fig,'12_taurine',ta,['pdac_rel','prad_rel'])
# 13 External glutamine panel.
rows=[]
for g in ['ASNS','GLS']:
 r=S['brca_rel'].query('gene==@g');r=r[r.metabolite_key=='KEGG:C00064'].iloc[0]
 for co in ['CAMP','FUSCC','Tang']:rows.append(dict(gene=g,cohort=co,rho=r[co+'_effect'],p=r[co+'_p_value'],n=r[co+'_n'],lo=r[co+'_ci_lower'],hi=r[co+'_ci_upper']))
z=pd.DataFrame(rows);fig,axs=plt.subplots(1,2,figsize=(10.8,3.6),sharex=True,constrained_layout=True)
for ax,g in zip(axs,['ASNS','GLS']):
 zz=z[z.gene==g];ax.set_title(g+'—谷氨酰胺',loc='left')
 for i,r in enumerate(zz.itertuples()):
  ax.plot([r.lo,r.hi],[i,i],c=BLUE,lw=2);ax.scatter(r.rho,i,s=65,edgecolors=BLUE,facecolors=BLUE if r.p<.05 else 'white');ax.text(-.75,i+.28,f'n={int(r.n)}  P={pf(r.p)}',fontsize=10)
 ax.set_yticks(range(3),zz.cohort);ax.set_ylim(2.6,-.5);ax.set_xlim(-.8,.45);ax.axvline(0,c=GRAY);ax.set_xlabel('Spearman ρ');clean(ax)
save(fig,'13_glutamine',z,['brca_rel'])
pd.DataFrame(manifest).to_csv(O/'figure_manifest.tsv',sep='\t',index=False)

# Native PPT text/shapes; quantitative figures regenerated above.
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5);prs.core_properties.title='CAMP代谢相关候选与四癌LYPLA1';prs.core_properties.author='CAMP project'
slides=[];FONT='Noto Sans CJK SC'
def rgb(s):return RGBColor.from_string(s.strip('#'))
def box(sl,x,y,w,h,text,size=20,color=INK,bold=False,fill=None):
 if fill:
  sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background();sh._element.spPr.append(OxmlElement('a:effectLst'));tf=sh.text_frame
 else:sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
 tf.clear();tf.word_wrap=True;tf.margin_left=Inches(.09);tf.margin_right=Inches(.09);tf.margin_top=Inches(.04);tf.margin_bottom=Inches(.03)
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(8)
 return sh
def new(title,kicker,claim,source='',note=''):
 sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb('#FFFFFF');num=len(prs.slides)
 box(sl,.48,.16,11,.27,kicker,10,TEAL,True);box(sl,.48,.58,12.3,.68,title,28,INK,True)
 box(sl,.48,1.40,12.3,.63,claim,17,INK,False,'#F2F5F7')
 box(sl,.50,6.91,11.75,.36,source,8,MUTED);box(sl,12.30,6.91,.5,.32,f'{num:02d}',10,MUTED)
 sl.notes_slide.notes_text_frame.text=note+'\n\n来源：'+source+'\n统计：复用冻结结果；名义P<0.05；不新增患者检验。'
 slides.append(dict(number=num,title=title,claim=claim,source=source,note=note));return sl
def pic(sl,name,x=.65,y=2.18,w=12.0,h=4.45,server=False):
 p=(O/'server_figures' if server else F)/(name+'.png');im=Image.open(p);rat=im.width/im.height;ww=min(w,h*rat);hh=ww/rat;sl.shapes.add_picture(str(p),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))
def cards(sl,items,y=2.35):
 w=(12.0-.25*(len(items)-1))/len(items)
 for i,(head,body) in enumerate(items):
  x=.65+i*(w+.25);box(sl,x,y,w,.55,head,22,TEAL,True);box(sl,x,y+.72,w,3.2,body,19,INK,False,'#F6F8FA')

sl=new('从CAMP代谢差异，聚焦四癌LYPLA1','CAMP / RESEARCH UPDATE','以代谢物—基因关系为起点，结合患者关联、单细胞来源与空间定位。','2026-09-28｜公开汇总冻结重绘｜LYPLA1：BRCA · COAD · PDAC · PRAD')
cards(sl,[('发现','CAMP六癌代谢差异\n精确代谢物直接映射\n保留全部候选背景'),('聚焦','四癌LYPLA1统计线索\nRNA差异与代谢关联分列\n从候选进入定向深入'),('定位','单细胞中的细胞来源\n患者层面的表达比较\n病理区域中的空间分布')])
sl=new('统一分析框架，区分每一层证据回答的问题','01 / PROJECT LOGIC','表达差异、患者相关、细胞定位相互补充；不相互替代。','来源：六癌主线整合报告；各队列原方法保留')
cards(sl,[('代谢层','肿瘤与参照有何不同？\n同一分子是否跨癌同向？\n缺测与不显著分开'),('基因层','哪些酶/转运体直接相关？\nRNA是否发生变化？\n肿瘤内是否与代谢物相关？'),('细胞与空间层','表达主要位于哪些细胞？\n患者方向是否重复？\n病理癌区是否有定位支持？')])
sl=new('六癌主线已有结果，覆盖与比较设计并不相同','02 / CAMP PANORAMA','柱尾为P<0.05数 / 可评估特征数；数量差异不能直接排名癌种异常程度。','图源01｜BRCA45对；COAD33对；PDAC11对；PRAD43对；GBM74/6非配对；ccRCC17/12对');pic(sl,'01_coverage')
sl=new('精确分子对照显示：共同变化与相反方向并存','02 / CAMP PANORAMA','红↑与蓝↓：名义P<0.05；浅灰点：未显著；NA：未覆盖或化学标识不唯一。','图源02｜固定14个讨论分子；不是全分子显著性排名；ccRCC两队列分开');pic(sl,'02_pan_metabolites',h=4.65)
sl=new('在CAMP关系候选中，LYPLA1出现四癌关注线索','03 / CANDIDATE FOCUS','BRCA、COAD、PRAD：RNA显著上调；PDAC：GPC—LYPLA1关联显著。','图源03｜固定讨论面板，不是全转录组排名；PDAC的LYPLA1 RNA P=0.0777',note='四癌都有统计线索采用不同证据类型的并列概括，不是四癌RNA均显著，也不是同一关系四癌复现。');pic(sl,'03_candidate_rna')
sl=new('LYPLA1 RNA在四癌同向，三癌达到名义显著','03 / LYPLA1 · RNA','保留PDAC的上调趋势与不显著结果，不将其写成四癌RNA全部显著。','图源04｜各队列配对RNA原结果；P值不是效应大小，也不代表患者逐人显著');pic(sl,'04_lypla1_rna')
sl=new('LYPLA1代谢联系必须逐一比较同一个分子','03 / LYPLA1 · METABOLITE','BRCA的LPC(16:0)正关联突出；PDAC的游离GPC负关联回答另一项关系。','图源05｜实心：P<0.05；空心：未显著；横线：原95%区间；未覆盖不等于阴性');pic(sl,'05_lypla1_relations')
sl=new('BRCA配对变化提供线索，但敏感性支持有限','03 / LYPLA1 · PATIENT','LPC(16:0)变化量主分析P=0.0312；可用值子集P=0.2647；游离GPC接近零相关。','图源06｜45对主分析 / 28对LPC可用值；同患者变化量不同于肿瘤间相关');pic(sl,'06_delta_sensitivity')
for cancer,cohort,caution in [('BRCA','Wu2021','既有分析者上皮UMAP；原有细胞身份'),('COAD','Uhlitz','CNA/CNN为原作者分类；CNN不等于非恶性；排除细胞仅作灰色背景'),('PDAC','GSE202051 未治疗','单核数据；原作者UMAP与细胞身份'),('PRAD','PRAD24 肿瘤内上皮','原作者UMAP；正常上皮与恶性上皮')]:
 sl=new(cancer+'：细胞身份与LYPLA1表达并列查看','04 / LYPLA1 · SINGLE CELL','统一点大小、绘制顺序和0–3色标；平台与取材不同，跨队列颜色不能直接比较。','图源：server redraw｜'+cohort+'｜'+caution)
 pic(sl,'umap_'+cancer,x=.55,y=2.16,w=12.1,h=4.48,server=True)
sl=new('患者层面支持应与整体细胞图分开呈现','04 / LYPLA1 · PATIENT REPEATABILITY','BRCA主分析7/8方向升高；PDAC未治疗9/9均值升高，仅阳性细胞核未获支持。','图源07｜患者是统计单位；BRCA Naive为标签敏感性；PDAC另一CRA001160肿瘤内14对P=0.6045');pic(sl,'07_sc_paired')
sl=new('均值升高，可能主要体现为更多细胞被检出','04 / LYPLA1 · EXPRESSION PATTERN','PRAD阳性细胞均值几乎相同；检出比例不同，不能解释为每个癌细胞都更强。','图源08｜蓝：参照上皮；红：恶性/CNA上皮；细胞等权描述，无新P值；检出受深度影响');pic(sl,'08_detection')
# Spatial examples: explicit lexical selection, preserve all 21 in supplemental atlas.
inv=pd.read_csv(O/'server_figures/spatial_display_inventory.tsv',sep='\t');brfirst=inv[inv.cancer=='BRCA'].sort_values('sample').iloc[0]['plot'];prfirst=inv[(inv.cancer=='PRAD')&(inv.primary_evaluable.astype(str).str.lower()=='true')].sort_values('sample').iloc[0]['plot']
sl=new('空间图统一为H&E—病理区域—LYPLA1表达','05 / LYPLA1 · SPATIAL','BRCA：癌区对免疫区3/3患者方向为正，P=0.25；未提供正常上皮对照。','CTA2025｜14切片 / 3患者｜正文按样本名排序取首片，非按结果选片；全部切片另附');pic(sl,brfirst,x=.60,y=2.15,w=12.1,h=4.20,server=True)
sl=new('PRAD病理癌区高于良性腺体，但仅来自一位患者','05 / LYPLA1 · SPATIAL','5张可比较切片方向一致；深度校正相对丰度比1.37–1.84。','Erickson2022｜7切片中5张可比较 / 1患者｜CC BY-NC 3.0；按字典序首张合格片展示');pic(sl,prfirst,x=.60,y=2.15,w=12.1,h=4.20,server=True)
sl=new('PDAC空间结果支持较高表达，也存在处理依赖','05 / LYPLA1 · SPATIAL','GeoMx主分析6/6癌上皮更高；两种敏感性为5/6，P=0.0625。','图源09｜Bell2025，6配对患者；正常导管均低于全局LOQ；各标准化尺度不合并');pic(sl,'09_geomx')
sl=new('癌细胞内部存在差异，尚不能概括为共同程序','06 / LYPLA1 · HETEROGENEITY','Wu Cycling在15/17供者中相对自身癌细胞均值较高；不同队列亚群不能直接对应。','图源10｜既有作者亚群；BRCA两队列未得到满足固定效应与患者方向规则的共同伴随基因');pic(sl,'10_states')
sl=new('COAD代表案例：UCKL1与尿苷负相关','07 / CANCER CONTEXT · COAD','代谢关联与上皮表达背景形成线索，但不能直接解释为UCKL1消耗尿苷。','图源11｜33个肿瘤；CNA上皮对正常上皮7/9较高但P=0.1797；不是已证实癌特异机制');pic(sl,'11_uckl1')
sl=new('牛磺酸—SLC6A6在PDAC和PRAD呈不同背景','07 / CANCER CONTEXT · PDAC / PRAD','PDAC两者升高且相关；PRAD两者下降，但共同方向不等于变化幅度相关。','图源12｜PDAC21肿瘤 / PRAD91肿瘤；PRAD配对变化ρ=−0.040、P=0.803；未做癌种交互检验');pic(sl,'12_taurine')
sl=new('BRCA代表案例：ASNS负关联获得外部队列支持','07 / CANCER CONTEXT · BRCA','ASNS在CAMP与FUSCC均与谷氨酰胺负相关；Tang小队列未显著，完整保留。','图源13｜CAMP60 / FUSCC258 / Tang20；GLS在CAMP较弱；条件关联不是通量或代谢介导');pic(sl,'13_glutamine')
sl=new('目前可讲清的是候选与表达背景，机制仍需独立证据','08 / TAKEAWAYS','LYPLA1来自CAMP关系候选，在四癌获得不同层面的线索，值得有边界地深入。','本版：所有定量图代码重绘；原P、效应和范围保留；单细胞与空间原始数据留server165')
cards(sl,[('已形成','六癌代谢与候选框架\n四癌LYPLA1定向证据\n三个癌种代表关系'),('已看清','RNA差异不等于代谢相关\n细胞检出与强度需分开\n空间区域不等于纯癌细胞'),('下一步','优先复核同一精确关系\n补足可靠患者与正常上皮对照\n不以更多图替代关键证据')])
source_links=pd.read_csv(O/'source_manifest.tsv',sep='\t').set_index('id').url.to_dict()
for i,sl in enumerate(prs.slides):
 text=slides[i]['source'];ids=[]
 for m in manifest:
  if ('图源'+m['figure'][:2]) in text:ids+=m['source_ids'].split(';')
 if 'server redraw' in text:ids=[{'BRCA':'brca_sc','COAD':'coad_sc','PDAC':'pdac_sc','PRAD':'prad_expr'}[slides[i]['title'].split('：')[0]]]
 if 'CTA2025' in text:ids=['brca_spatial']
 if 'Erickson2022' in text:ids=['prad_spatial','prad_regions']
 sl.notes_slide.notes_text_frame.text+='\n\n冻结汇总表链接：\n'+'\n'.join(source_links[k] for k in dict.fromkeys(ids))
prs.save(O/'CAMP_LYPLA1_统一重绘版.pptx')
(O/'slide_manifest.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding='utf8')
style=dict(background='white',font=FONT,up_or_malignant=RED,down_or_reference=BLUE,context=TEAL,missing=GRAY,nominal_p=.05,figure_formats=['PNG 260dpi','PDF fonttype42'],UMAP='1pt dots, fixed random seed, original coordinates, gray zero, 0–3 shared log1pCP10K colorbar',spatial='H&E / author pathology / measured LYPLA1; no smoothing; same sequential scale; all source sections retained')
(O/'STYLE_STANDARD.json').write_text(json.dumps(style,ensure_ascii=False,indent=2),encoding='utf8')
checks={'slides':len(prs.slides),'quantitative_figures':len(manifest),'new_patient_tests':0,'new_embeddings':0,'out_of_bounds_shapes':[]}
for i,sl in enumerate(prs.slides,1):
 for sh in sl.shapes:
  if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+100 or sh.top+sh.height>prs.slide_height+100:checks['out_of_bounds_shapes'].append([i,sh.name])
assert not checks['out_of_bounds_shapes'];assert len(z)==6
(O/'build_validation.json').write_text(json.dumps(checks,indent=2))
print(json.dumps(checks))
