"""Public summaries only: version-aware integration, PDF, Excel, and auditable figures."""
import argparse,json,hashlib,subprocess,zipfile,textwrap
from pathlib import Path
import numpy as np,pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from openpyxl.styles import Font,PatternFill,Alignment
R=Path(__file__).resolve().parents[1];C=R/'results/ccRCC'
COHORTS={'ccRCC3':('20260925T151000Z_ccrcc3_v1','20260925T155500Z_ccrcc3_therapy_v2'),'ccRCC4':('20260925T155000Z_ccrcc4_histology_v2','20260925T155000Z_ccrcc4_histology_v2')}
SC=C/'06_EXTERNAL/20260925T153000Z_source_union_v1'
def save(d,p):d.to_csv(p,sep='\t',index=False,na_rep='NA',lineterminator='\n')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main(run,commit):
 out=C/'07_INTEGRATION'/run;out.mkdir(parents=True,exist_ok=False);figdir=out/'figures';figdir.mkdir()
 inputs=[]
 def read(p):inputs.append(p);return pd.read_csv(p,sep='\t')
 datasets={};relations=[];progress=[];long=[]
 sc=read(SC/'sc_celltype_profiles.tsv');rank=read(SC/'sc_source_stability.tsv');union=read(SC/'gene_membership.tsv')
 for co,(dr,pr) in COHORTS.items():
  base=C/'04_ROBUSTNESS'/dr;met=read(base/'metabolite_paired.tsv');available=read(base/'metabolite_available_sensitivity.tsv');mapping=read(C/'02_MAPPING'/dr/'direct_relations.tsv');coverage=read(C/'02_MAPPING'/dr/'metabolite_mapping_status.tsv');assoc=read(C/'03_PATIENT'/pr/'association_results.tsv');rna=read(C/'03_PATIENT'/pr/'rna_results.tsv')
  audit=json.loads((C/'01_CAMP'/dr/'validation.json').read_text());inputs.append(C/'01_CAMP'/dr/'validation.json')
  datasets[co]=dict(met=met,available=available,mapping=mapping,coverage=coverage,assoc=assoc,rna=rna,audit=audit)
  d=mapping.copy()
  for label,table,key in [('metabolite',met,'feature_id'),('metabolite_available',available,'feature_id'),('RNA',rna,'gene')]:
   cols=[key]+[c for c in ['effect','n','ci_lower','ci_upper','p_value','q_value','status','reason'] if c in table]
   d=d.merge(table[cols].rename(columns={c:label+'_'+c for c in cols if c!=key}),on=key,how='left',validate='many_to_one')
  for fam,label in [('REL_TUMOR_PRIMARY','association'),('REL_TUMOR_AVAILABLE','association_available'),('REL_TUMOR_TREATMENT_STRATIFIED','association_treatment')]:
   z=assoc[assoc.test_family.eq(fam)];cols=['relation_id','effect','n','ci_lower','ci_upper','p_value','q_value','status','reason']
   d=d.merge(z[cols].rename(columns={c:label+'_'+c for c in cols if c!='relation_id'}),on='relation_id',how='left',validate='one_to_one')
  d=d.merge(rank.rename(columns={c:'sc_'+c for c in rank if c not in ['gene','stable_gene_id']}),on=['gene','stable_gene_id'],how='left',validate='many_to_one')
  d['patient_association_evidence']=np.select([d.association_q_value.lt(.05),d.association_p_value.lt(.05),d.association_p_value.isna()],['FDR_SUPPORTED','NOMINAL_EXPLORATORY','NOT_EVALUABLE'],default='NOT_SUPPORTED_THIS_TEST')
  d['interpretation_boundary']='same_cohort_postselection;RNA_not_activity;rho_not_flux;treated_context;no_independent_validation'
  relations.append(d)
  for stage,fam,z in [('04_ROBUSTNESS','METAB_PAIRED_PRIMARY',met),('04_ROBUSTNESS','METAB_PAIRED_AVAILABLE',available),('03_PATIENT','REL_TUMOR_PRIMARY',assoc[assoc.test_family.eq('REL_TUMOR_PRIMARY')]),('03_PATIENT','REL_TUMOR_TREATMENT_STRATIFIED',assoc[assoc.test_family.eq('REL_TUMOR_TREATMENT_STRATIFIED')]),('03_PATIENT','RNA_PAIRED_CURRENT',rna[rna.test_family.eq('RNA_PAIRED_CURRENT')])]:
   progress.append(dict(cohort=co,stage_id=stage,analysis=fam,planned=len(z),evaluable=int(z.p_value.notna().sum()),p005=int(z.p_value.lt(.05).sum()),q005=int(z.q_value.lt(.05).sum()),status='DONE',reason='small treated cohort;see version decisions'))
  progress.append(dict(cohort=co,stage_id='02_MAPPING',analysis='bounded_direct_mapping',planned=len(coverage),evaluable=mapping.feature_id.nunique(),p005=np.nan,q005=np.nan,status='PARTIAL',reason=f'{int(coverage.n_direct_relations.eq(0).sum())} features still NEEDS_REVIEW;not exhaustive'))
  long.extend([met,available,assoc,rna])
 rel=pd.concat(relations,ignore_index=True);assert not rel.duplicated(['cohort','relation_id']).any();save(rel,out/'candidate_relations_integrated.tsv')
 gene=union[['gene','stable_gene_id']].copy()
 for co,z in datasets.items():
  gene['current_in_'+co]=gene.gene.isin(z['mapping'].gene)
  gene['n_relations_'+co]=gene.gene.map(z['mapping'].groupby('gene').size()).fillna(0).astype(int)
  cols=['gene','effect','n','p_value','q_value','status','reason'];rr=z['rna'][cols].rename(columns={c:co+'_RNA_'+c for c in cols if c!='gene'});gene=gene.merge(rr,on='gene',how='left',validate='one_to_one')
  missing=gene[co+'_RNA_status'].isna();gene.loc[missing,co+'_RNA_status']='NOT_RUN';gene[co+'_RNA_reason']=gene[co+'_RNA_reason'].astype(object);gene.loc[missing,co+'_RNA_reason']='not_in_this_cohort_current_or_legacy_RNA_family'
 gene['current_pool']=gene.current_in_ccRCC3|gene.current_in_ccRCC4;gene['history_only']=~gene.current_pool
 gene=gene.merge(rank,on=['gene','stable_gene_id'],validate='one_to_one');gene['cell_background']=np.where(gene.status.eq('DONE'),gene.top_celltype,'暂不可定位')
 gene['interpretation']='描述性细胞来源；不是酶活、代谢通量或功能验证'
 save(gene,out/'candidate_genes_integrated.tsv');save(pd.DataFrame(progress),out/'progress_summary.tsv');save(pd.concat(long+[sc],ignore_index=True),out/'statistics_long.tsv')
 common=relations[0][['metabolite_key','metabolite_name','gene','association_effect','association_p_value','association_q_value']].merge(relations[1][['metabolite_key','metabolite_name','gene','association_effect','association_p_value','association_q_value']],on=['metabolite_key','metabolite_name','gene'],suffixes=('_ccRCC3','_ccRCC4'),validate='one_to_one')
 common['same_rho_direction']=np.sign(common.association_effect_ccRCC3).eq(np.sign(common.association_effect_ccRCC4));common['both_nominal_p005']=common.association_p_value_ccRCC3.lt(.05)&common.association_p_value_ccRCC4.lt(.05);common['independent_replication_claim']=False;save(common,out/'cross_cohort_descriptive.tsv')
 pd.set_option('display.max_colwidth',60)
 workbook=out/'ccRCC_完整分析与候选比较.xlsx'
 with pd.ExcelWriter(workbook,engine='openpyxl') as w:
  tables={'阶段进度':pd.DataFrame(progress),'全关系比较':rel,'全基因比较':gene,'共同关系描述':common,'单细胞供者等权':sc,'单细胞来源稳定性':rank}
  for co,z in datasets.items():
   for k in ['met','available','mapping','coverage','assoc','rna']:tables[co+'_'+k]=z[k]
  for name,d in tables.items():
   safe=d.copy()
   for col in safe.select_dtypes('object'):
    safe[col]=safe[col].map(lambda v:"'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v)
   safe.to_excel(w,sheet_name=name,index=False,na_rep='NA');ws=w.book[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
   for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24496A');cell.alignment=Alignment(wrap_text=True)
   ws.row_dimensions[1].height=42
   for cells in ws.columns:
    ws.column_dimensions[cells[0].column_letter].width=min(38,max(14,len(str(cells[0].value))+2))
 plt.rcParams.update({'font.family':'Microsoft YaHei','axes.unicode_minus':False,'figure.facecolor':'white','axes.spines.top':False,'axes.spines.right':False,'font.size':10})
 figures=[]
 def figsave(fig,name,data,caption):
  for ext in ['png','pdf']:fig.savefig(figdir/(name+'.'+ext),dpi=170,bbox_inches='tight')
  plt.close(fig);save(data,figdir/(name+'_source.tsv'));(figdir/(name+'_caption_CN.md')).write_text(caption,encoding='utf-8');figures.append(dict(figure_id=name,png='figures/'+name+'.png',vector='figures/'+name+'.pdf',source='figures/'+name+'_source.tsv',caption=caption));return figdir/(name+'.png')
 fig,ax=plt.subplots(figsize=(9,4));counts=pd.DataFrame([dict(cohort=k,肿瘤标本=v['audit']['tumor_specimens'],正常标本=v['audit']['normal_specimens'],患者配对=v['audit']['paired_patients']) for k,v in datasets.items()]);counts.set_index('cohort').plot.bar(ax=ax,color=['#477B9E','#A5BCCB','#D38F58'],rot=0);ax.set_ylabel('数量');ax.set_title('多区域标本先合并到作者患者：标本数不作为独立n');samplefig=figsave(fig,'01_sample_design',counts,'患者身份来自作者SUBJECT_ID；ccRCC4采用原研究病理交叉核实后的修正版。')
 fig,axs=plt.subplots(1,2,figsize=(11,4));fd=[]
 for ax,(co,z) in zip(axs,datasets.items()):
  d=z['met'].copy();d['level']=np.select([d.q_value.lt(.05),d.p_value.lt(.05)],['q<0.05','仅P<0.05'],default='未过阈值');colors=d.level.map({'q<0.05':'#BD654D','仅P<0.05':'#D7AC57','未过阈值':'#ACB7BD'});ax.scatter(d.effect,-np.log10(d.p_value),c=colors,s=9,alpha=.7);ax.set_title(co);ax.set_xlabel('患者配对均值差（作者处理尺度）');ax.set_ylabel('−log10(P)');fd.append(d)
 metfig=figsave(fig,'02_metabolite_discovery',pd.concat(fd),'全量检验后分别BH；颜色分开FDR支持、仅名义P支持和未过阈值。效应不是浓度倍数。')
 fig,axs=plt.subplots(1,2,figsize=(11,4));fd=[]
 for ax,(co,z) in zip(axs,datasets.items()):
  d=z['met'][['feature_id','metabolite_name','effect']].merge(z['available'][['feature_id','effect','n','status']],on='feature_id',suffixes=('_primary','_available'),validate='one_to_one');d['cohort']=co;d=d[d.status.eq('DONE')];ax.scatter(d.effect_primary,d.effect_available,s=8,color='#42748E',alpha=.6);lo=min(d.effect_primary.min(),d.effect_available.min());hi=max(d.effect_primary.max(),d.effect_available.max());ax.plot([lo,hi],[lo,hi],color='#999999',linestyle='--');ax.set_title(co);ax.set_xlabel('主分析均值差');ax.set_ylabel('全部区域有作者可用值的患者子集');fd.append(d)
 sensfig=figsave(fig,'03_available_sensitivity',pd.concat(fd),'只比较相同feature_id；敏感性有效患者减少，缺项不当作阴性。')
 fig,ax=plt.subplots(figsize=(9,4));covcounts=pd.DataFrame([dict(cohort=co,有直接关系=int(z['coverage'].n_direct_relations.gt(0).sum()),仍待审=int(z['coverage'].n_direct_relations.eq(0).sum())) for co,z in datasets.items()]);covcounts.set_index('cohort').plot.barh(stacked=True,ax=ax,color=['#4E8295','#CCD3D6']);ax.set_xlabel('探索工作池代谢特征数');mapfig=figsave(fig,'04_mapping_coverage',covcounts,'有界精确生化映射；未映射不代表没有关系，不能称为穷尽性注释。')
 fig,axs=plt.subplots(1,2,figsize=(11,4));fd=[]
 for ax,(co,z) in zip(axs,datasets.items()):
  d=z['assoc'][z['assoc'].test_family.eq('REL_TUMOR_PRIMARY')];ax.scatter(d.effect,-np.log10(d.p_value),s=17,c=np.where(d.p_value.lt(.05),'#BD654D','#ADB9BF'));ax.axhline(-np.log10(.05),color='#888888',linestyle='--');ax.set_title(co);ax.set_xlabel('Spearman rho（患者均值）');ax.set_ylabel('−log10(P)');fd.append(d)
 assocfig=figsave(fig,'05_all_associations',pd.concat(fd),'全部关系保留。两队列均没有主关联通过q<0.05；名义P线索不等于稳健或因果证据。')
 fig,axs=plt.subplots(1,2,figsize=(11,4));fd=[]
 for ax,(co,z) in zip(axs,datasets.items()):
  d=z['assoc'];p=d[d.test_family.eq('REL_TUMOR_PRIMARY')][['relation_id','p_value']].merge(d[d.test_family.eq('REL_TUMOR_TREATMENT_STRATIFIED')][['relation_id','p_value']],on='relation_id',suffixes=('_primary','_stratified'));p['cohort']=co;ax.scatter(-np.log10(p.p_value_primary),-np.log10(p.p_value_stratified),s=15,color='#477B9E');ax.set_title(co);ax.set_xlabel('未调整 −log10(P)');ax.set_ylabel('治疗分层 −log10(P)');fd.append(p)
 therapyfig=figsave(fig,'06_treatment_sensitivity',pd.concat(fd),'分层置换用于敏感性；不是治疗效应比较。ccRCC3使用两个作者治疗字段组合，避免NO组中的Cabozantinib暴露被掩盖。')
 fig,axs=plt.subplots(1,2,figsize=(11,4));fd=[]
 for ax,(co,z) in zip(axs,datasets.items()):
  d=z['rna'][z['rna'].test_family.eq('RNA_PAIRED_CURRENT')];ax.scatter(d.effect,-np.log10(d.p_value),s=17,c=np.where(d.q_value.lt(.05),'#BD654D','#ADB9BF'));ax.set_title(co);ax.set_xlabel('患者配对均值差：log2(1+TPM)');ax.set_ylabel('−log10(P)');fd.append(d)
 rnafig=figsave(fig,'07_RNA_background',pd.concat(fd),'RNA采用作者TPM加1取log2后患者区域均值配对t检验；不是RNA原始计数模型，也不是酶活。')
 qualified=rank[rank.status.eq('DONE')];sourcecounts=qualified.top_celltype.value_counts().rename_axis('top_celltype').reset_index(name='n_genes');fig,ax=plt.subplots(figsize=(9,4));ax.barh(sourcecounts.top_celltype,sourcecounts.n_genes,color='#598C96');ax.set_xlabel('可定位基因数');scsumfig=figsave(fig,'08_cell_source_summary',sourcecounts,'全历史并集157基因的描述性最高类别；未达到覆盖/检出阈值的基因不强排第一，ua不参与排名。')
 selected=sorted(gene.loc[gene.current_pool,'gene'])[:16];sub=sc[sc.gene.isin(selected)];matrix=sub.pivot(index='gene',columns='celltype',values='effect').reindex(selected);v=matrix.to_numpy();mu=np.nanmean(v,1,keepdims=True);sd=np.nanstd(v,1,keepdims=True);zz=(v-mu)/np.where(sd>0,sd,1);cmap=plt.get_cmap('RdBu_r').copy();cmap.set_bad('#D9D9D9');fig,ax=plt.subplots(figsize=(10,5.5));im=ax.imshow(zz,aspect='auto',cmap=cmap,vmin=-2,vmax=2);ax.set_yticks(range(len(selected)));ax.set_yticklabels(selected,fontsize=10);ax.set_xticks(range(len(matrix.columns)));ax.set_xticklabels(matrix.columns,rotation=45,ha='right');fig.colorbar(im,ax=ax,label='行z分数（仅展示）');scfig=figsave(fig,'09_cell_source_alphabetical_examples',sub,'当前基因按字母序取前16个，非按显著性或表达强度挑选；完整157基因图位于06_EXTERNAL。行缩放只比较同一基因跨类别，不比较不同基因绝对丰度。')
 save(pd.DataFrame([dict(gene=g,rule='first16_current_genes_alphabetical_not_by_significance') for g in selected]),out/'display_selection.tsv')
 fig,ax=plt.subplots(figsize=(9,4));ax.hist(qualified.bootstrap_top_frequency.dropna(),bins=np.linspace(0,1,11),color='#598C96',edgecolor='white');ax.axvline(.8,color='#BD654D',linestyle='--');ax.set_xlabel('1000次联合供者bootstrap中原最高类别保持率');ax.set_ylabel('基因数');stabilityfig=figsave(fig,'10_source_stability',qualified,'80%为描述性标记，不是80%患者一致或显著性检验；只有7个来源标签，排名会受类别覆盖影响。')
 save(pd.DataFrame(figures),out/'figure_manifest.tsv')
 # Fourteen readable A4 landscape pages, with full tables and large source figures separate.
 pdfmetrics.registerFont(TTFont('CJK','C:/Windows/Fonts/simhei.ttf'))
 pdf=out/'ccRCC_分析报告.pdf';cv=canvas.Canvas(str(pdf),pagesize=(842,595));cv.setTitle('ccRCC：CAMP发现至单细胞表达来源')
 page=0
 def pagewrite(title,paragraphs,image=None):
  nonlocal page
  page+=1;cv.setFillColorRGB(.12,.23,.31);cv.setFont('CJK',23);cv.drawString(40,552,title);y=516
  cv.setFillColorRGB(.18,.22,.25);cv.setFont('CJK',12)
  for paragraph in paragraphs:
   for line in textwrap.wrap(paragraph,width=61,replace_whitespace=False,break_long_words=True):cv.drawString(40,y,line);y-=19
   y-=9
  if image:
   im=ImageReader(str(image));iw,ih=im.getSize();w=745;h=min(350,y-65);scale=min(w/iw,h/ih);cv.drawImage(im,(842-iw*scale)/2,55+(h-ih*scale)/2,width=iw*scale,height=ih*scale)
  cv.setFont('CJK',8);cv.setFillColorRGB(.45,.48,.5);cv.drawString(40,24,'ccRCC | 当前版本与历史版本分开 | 仅公开汇总 | '+commit[:12]);cv.drawRightString(800,24,str(page));cv.showPage()
 pagewrite('ccRCC：按BRCA主线完成首轮分析',['范围：样本身份与病理 → 全量代谢物 → P<0.05探索池 → 直接生化映射 → 患者关联与RNA背景 → 全候选单细胞来源。','当前主分析：ccRCC3为17对；ccRCC4病理修正版为12对。两个队列分别分析，没有合并标本或把多区域当独立患者。',f'当前两队列共{len(rel)}条队列内关系，{int(gene.current_pool.sum())}个去重当前基因；保留历史后{len(gene)}个基因。','主要结论：代谢物与RNA差异较多，但目前没有患者主关联通过各队列全关系族的q<0.05。结果属于探索性候选资源，不能直接宣称机制或靶点验证。'])
 pagewrite('1  实际样本设计',['问题：统计n应当按什么计算？','实际结果：明确作者SUBJECT_ID连接，多个区域在患者×组织内先等权取均值。'],samplefig)
 pagewrite('2  病理与治疗资料改变了分析范围',['ccRCC4首版只依赖CAMP病理表，得到14位肿瘤患者、13对；后续原研究MetaData_M5揭示另外2位患者的病理冲突。','当前版本整患者排除髓质癌、嫌色细胞癌和未分类癌，共排除3位患者、5个标本，保留71标本、12位患者、12对。','首版发现、映射与患者结果保留历史，不用于当前主结论。','ccRCC3全部67个肿瘤的原研究病理均为clear cell，114个RNA连接全部交叉一致。治疗NO组中另有1位Cabozantinib暴露患者，因此敏感性改用更细作者字段组合。','这是明确的数据版本修正；没有覆盖原CAMP612项/310个q<0.05冻结统计。'])
 pagewrite('3  全量配对代谢物发现',['实际结果：ccRCC3为711项，P<0.05有469项、q<0.05有454项；ccRCC4为904项，分别432项和337项。','限制：作者处理尺度上的患者均值差不是浓度倍数，检验族分别BH。'],metfig)
 pagewrite('4  缺失／作者可用值敏感性',['要求患者被纳入的所有区域在相应特征上都有作者可用值，保持主分析的区域组成。','实际结果：ccRCC3可评估540项，335项q<0.05；ccRCC4可评估558项，200项q<0.05。'],sensfig)
 pagewrite('5  直接生化映射覆盖与缺口',['当前ccRCC3建立221关系/144基因，ccRCC4建立130关系/95基因；分别68和44个探索代谢特征有直接关系。','限制：仅复用有明确化学键与名称的审定人类生化库。未映射清单完整保留，不是“没有相关基因”。'],mapfig)
 pagewrite('6  全关系患者内部关联',['ccRCC3主关联19/221条P<0.05；ccRCC4为18/130条P<0.05；两队列q<0.05均为0。','解释：小样本下名义线索仍可保留，但不能当作FDR支持或否定其他未过阈值关系。'],assocfig)
 pagewrite('7  治疗背景敏感性',['治疗分层置换：ccRCC3有17条P<0.05，ccRCC4有14条P<0.05；q<0.05仍均为0。','限制：用于考察混合治疗背景，不估计治疗因果效果；作者治疗历史与采样时用药不完全等价。'],therapyfig)
 pagewrite('8  配对RNA背景',['当前基因族：ccRCC3为144基因，其中110个q<0.05；ccRCC4为95基因，其中50个q<0.05。','解释：RNA背景不作为进入单细胞的显著性门槛；RNA丰度不等于酶活。'],rnafig)
 pagewrite('9  全候选单细胞表达来源',['GSE159115：20,748个作者肿瘤组织细胞、7个患者来源标签；157个当前/历史基因全部精确覆盖。','142个基因满足至少两可评估类别和检出阈值，允许描述最高类别；只作表达背景。'],scsumfig)
 pagewrite('10  表达图：明确展示选择与缺项',['示例按当前基因字母序取前16个，不根据最小P/q或最强表达挑选；完整157基因热图/点图另附。','灰色=不可评估；每供者每类≥20细胞、每类≥3供者；无新增聚类。'],scfig)
 pagewrite('11  来源排名稳定性',['供者联合重抽样保留同一供者的类别关联；80%阈值仅为描述性标记。','限制：仅一项单细胞研究，不能宣称双研究一致；最高表达类别不一定是唯一作用或代谢物生成细胞。'],stabilityfig)
 pagewrite('12  证据整合与当前决定',[f'完整关系表保留{len(rel)}条队列内关系；完整基因表保留{len(gene)}个当前/历史基因，当前去重{int(gene.current_pool.sum())}个。',f'两队列共有{len(common)}条相同化学键＋名称＋基因关系，其中{int(common.both_nominal_p005.sum())}条两边均名义P<0.05，其中3条同向、2条反向；不能把两边显著当作方向一致或独立复现。','证据分层：代谢物发现、患者关联、RNA背景、细胞来源分别展示；不合成为不透明总分，不把RNA或细胞来源补成关联的FDR支持。','当前适合讨论探索候选与数据缺口，不适合直接宣布验证靶点。基础功能、CPTAC、外部患者复现若无新分析，明确NOT_RUN。'])
 pagewrite('13  交付、限制与复现',['交付：14页PDF、可筛选Excel、全量TSV、每张图PNG＋PDF＋图源TSV、中文图注、代码、参数与SHA256。','未完成／不可声称：穷尽生化映射、第二独立单细胞研究、作者UMAP坐标、功能或因果验证。未测与未显著严格分开。','输入提交：'+commit,'复现：python code/ccrcc_integrate_report_v1.py --run 新运行ID --input-commit 上述提交。服务器逐样本数据与凭据不进入交付包。','所有具体分析脚本、运行目录、参数、输入哈希、实际n和版本决定均保留。当前结果分支未自动合并main。'])
 cv.save();assert page==14
 readme=f'''# ccRCC 首轮数据分析交付\n\n问题：按BRCA主线推进ccRCC，统计单位、病理与治疗背景按本癌种真实设计适配。\n\n输入范围：ccRCC3 17对；ccRCC4原研究病理交叉核实后12对。基线提交 `{commit}`。版本修正详见../../VERSION_DECISIONS_CN.md。\n\n实际结果：\n\n|项目|ccRCC3|ccRCC4 修正版|\n|---|---:|---:|\n|配对患者|17|12|\n|全量保留代谢特征|711|904|\n|代谢物 P<0.05 / q<0.05|469 / 454|432 / 337|\n|直接关系 / 当前基因|221 / 144|130 / 95|\n|主关联 P<0.05 / q<0.05|19 / 0|18 / 0|\n|当前RNA q<0.05|110|50|\n\n当前去重{int(gene.current_pool.sum())}基因、含历史{len(gene)}基因；GSE159115全部覆盖，142个可描述来源排名。\n\n新手解释：代谢物差异和RNA差异不能自动证明代谢物—基因联系。当前患者主关联均未通过各队列BH q<0.05，保留的是探索线索。单细胞只说明表达位置，不说明酶活、通量、因果或药物靶点有效。\n\n限制：多区域小患者数、混合治疗背景、同队列筛选、有界映射尚不穷尽；两个CAMP队列不声称独立验证。第二独立单细胞资料未完成，现成UMAP缺失，功能/机制未新增。\n\n当前决定：主线可计算部分已交付，缺项逐条保留。采用病理修正版ccRCC4和细化治疗分层ccRCC3，首版数值不覆盖且不混入当前主表。\n\n下一步：若继续深入，应先补未映射分子的身份与直接证据、独立患者样本和第二来源研究，不从本轮名义P结果直接宣布靶点。\n\n复现：`python code/ccrcc_integrate_report_v1.py --run NEW_RUN --input-commit {commit}`。数值统计在server165执行；本脚本仅读取允许公开的汇总。\n'''
 (out/'README_CN.md').write_text(readme,encoding='utf-8');(out/'INTERPRETATION_CN.md').write_text(readme,encoding='utf-8')
 save(pd.DataFrame([dict(path_or_url=p.relative_to(R).as_posix(),sha256=sha(p),input_commit=commit) for p in dict.fromkeys(inputs)]),out/'source_manifest.tsv')
 (out/'analysis_spec.json').write_text(json.dumps(dict(version='ccrcc_integrate_report_v1',input_commit=commit,cohort_versions=COHORTS,SC=str(SC.relative_to(R)),new_statistics=False,relation_key='cohort x feature_id x gene;no best-p collapsing',gene_scope='current plus history',plot_examples='alphabetical first16 current genes',software=dict(pandas=pd.__version__,numpy=np.__version__,matplotlib=matplotlib.__version__)),indent=2),encoding='utf-8')
 val=dict(status='DONE',relations=len(rel),current_genes=int(gene.current_pool.sum()),history_union_genes=len(gene),SC_genes=len(rank),source_rank_evaluable=int(rank.status.eq('DONE').sum()),unique_relation_keys=True,unique_gene_keys=bool(gene.gene.is_unique),PDF_pages=page,current_primary_association_q005=int(rel.association_q_value.lt(.05).sum()),cross_cohort_common_relations=len(common),both_nominal_relations=int(common.both_nominal_p005.sum()),independent_replication_claim=False,raw_or_patient_data_exported=False,remaining_modules=['bounded mapping gaps','second independent SC study','existing UMAP coordinates','functional mechanism validation'])
 (out/'validation.json').write_text(json.dumps(val,indent=2),encoding='utf-8');save(pd.DataFrame([dict(file=p.relative_to(out).as_posix(),sha256=sha(p)) for p in out.rglob('*') if p.is_file() and p.name!='checksums.tsv']),out/'checksums.tsv');print(json.dumps(val))
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--run',required=True);p.add_argument('--input-commit',required=True);a=p.parse_args();main(a.run,a.input_commit)
