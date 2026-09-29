"""Validate and package the actual WPS-rendered editable deck."""
from pathlib import Path
import json,hashlib,re,shutil,html,subprocess
from PIL import Image,ImageOps,ImageDraw
from pptx import Presentation
import pymupdf as fitz
ROOT=Path(__file__).resolve().parents[1];RUN='20260929T_illustrated_main_v2';O=ROOT/'results/BRCA/07_INTEGRATION'/RUN
V=Path('D:/CodexData/visualizations/2026/09/29/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/CAMP_main_illustrated_v2')
PPT='CAMP_泛癌与LYPLA1_插图精修可编辑版.pptx';PDF='CAMP_插图精修版_预览.pdf'
old=Presentation(ROOT/'results/BRCA/07_INTEGRATION/20260929T_editable_main_v1/CAMP_泛癌发现与LYPLA1_可编辑汇报.pptx')
r=Presentation(O/PPT)
for a,b in zip(list(old.slides)[5:],list(r.slides)[5:]):
 assert [x.text for x in a.shapes if x.has_text_frame]==[x.text for x in b.shapes if x.has_text_frame]
 assert [(x.left,x.top,x.width,x.height) for x in a.shapes]==[(x.left,x.top,x.width,x.height) for x in b.shapes]
import zipfile
with zipfile.ZipFile(O/PPT) as z: assert len(z.namelist())==len(set(z.namelist()))
pdf=fitz.open(O/PDF);imgs=sorted((O/'rendered').glob('*.PNG'),key=lambda p:int(re.search(r'\d+',p.stem)[0]));assert len(r.slides)==len(pdf)==len(imgs)==18
titles=['封面','CAMP 论文与资源','CAMP 的优势与匹配','六癌种七队列设计','从代谢物到候选基因','各队列发现概况','代谢物同向与患者一致性','基因同向与患者一致性','代谢物如何指向 LYPLA1','LYPLA1 四癌组织与患者','BRCA 单细胞','COAD 单细胞','PDAC 单核','PRAD 单细胞','BRCA 空间定位','PRAD 空间定位','LYPLA1 近期文献','LYPLA1 多层证据汇总']
stats=[];alltext=''
for i,s in enumerate(r.slides,1):
 text=[sh for sh in s.shapes if sh.has_text_frame and sh.text.strip()];pictures=[sh for sh in s.shapes if sh.shape_type==13]
 assert len(text)>=4
 assert not any(sh.width>r.slide_width*.95 and sh.height>r.slide_height*.95 for sh in pictures)
 alltext+='\n'.join(sh.text for sh in text)+'\n'
 stats.append(dict(slide=i,title=titles[i-1],objects=len(s.shapes),editable_text_boxes=len(text),independent_images=len(pictures),editable_graphics=len(s.shapes)-len(text)-len(pictures)))
assert all(k in alltext for k in ['7 / 8','9 / 9','10 / 10','PDAC','LYPLA1'])
validation=dict(status='PASS',slides=18,wps_open_export='PASS',pdf_pages=len(pdf),native_text_boxes=sum(x['editable_text_boxes'] for x in stats),editable_graphics=sum(x['editable_graphics'] for x in stats),independent_images=sum(x['independent_images'] for x in stats),full_slide_flattening=False,result_slides_6_to_18_text_and_geometry_unchanged=True,new_statistical_tests=False,visual_review='All 18 WPS-rendered thumbnails reviewed; opening pages 1 through 5 additionally inspected at full resolution',limitations='Dense scientific tables retain small annotations; native vector shapes are editable but not linked Excel charts. UMAP/tissue/paper images remain raster, separately replaceable.')
(O/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2),encoding='utf-8')
(O/'editable_object_audit.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
(O/'slide_outline.tsv').write_text('slide\ttitle\n'+''.join(f'{i+1}\t{t}\n' for i,t in enumerate(titles)),encoding='utf-8')
spec=dict(version=RUN,scope='Before the three representative cases: CAMP discovery through LYPLA1 summary',source_commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),new_tests=False,layout='16:9, SYSU green, Microsoft YaHei East Asian fonts, standardized header',data='Frozen existing aggregate charts; no patient/cell matrices transferred',editable='Native text, paths, curves and tables-as-elements; raster for real tissue, UMAP clouds, paper screenshot and gradient color bars',private_sources='Remain server165; existing public figure geometry only',software=dict(pymupdf=fitz.VersionBind,python_pptx=__import__('pptx').__version__,renderer='Installed WPS KWPP.Application'))
(O/'analysis_spec.json').write_text(json.dumps(spec,ensure_ascii=False,indent=2),encoding='utf-8')
scripts=list((ROOT/'code').glob('ppt_illustrated_main*'))
(O/'code_manifest.tsv').write_text('path\tsha256\n'+''.join(f'{p.relative_to(ROOT)}\t{hashlib.sha256(p.read_bytes()).hexdigest()}\n' for p in scripts),encoding='utf-8')
readme='''# CAMP 泛癌发现与 LYPLA1：可编辑组会汇报

## 本轮问题
将三个代表案例之前的成果收拢成正式 PPTX，统一版式并提高可编辑性。

## 输入与范围
采用后续已确认版本：开场五页重建为原生文字、连接线与独立科研插图；队列概况、同向代谢物和基因（含患者一致性）、代谢物–LYPLA1桥接、组织配对、四癌全细胞/单核、选定BRCA/PRAD空间、近期文献、简洁证据汇总。插图来源与哈希见illustration_manifest.json；科学结果来源沿用前版。

## 实际结果
共18页，见slide_outline.tsv。没有纳入ASNS、SLC6A6、UCKL1三个案例。原始科学结果和阈值保持不变。原生对象计数及WPS打开/导出结果见validation.json。

## 新手解释
插图是独立可裁剪、可替换的位图，内部像素不作为矢量形状编辑。标题、正文、数字、图例、箭头、点阵和患者连线可在WPS/PowerPoint选择、改字或改样式；UMAP点云、病理照片、论文截图为独立可替换图片，没有整页截图拼接。矢量图形是原生可编辑形状，不是连接Excel数据的原生统计图表；改数据应从原分析代码重新绘制。

## 限制/反证
跨癌总览含六癌，LYPLA1专题限定BRCA/COAD/PDAC/PRAD。四癌组织平均上调，PDAC未显著；不能称四癌均显著。患者比例不等于每位患者显著。空间定位不等于癌细胞特异性或因果证据。两个PRAD切片来自同一患者。文献仅复用已核对页面，不代表本次再次系统检索。部分密集统计页保留小字；可在PPT中按汇报需求拆页。WPS已实际打开导出；未在Microsoft PowerPoint单独渲染。

## 当前决定
交付可编辑PPTX、WPS导出PDF、18页预览与来源说明。统一中文字体显式设置东亚字体，避免WPS自动替换为宋体。

## 下一步
组会审阅内容与讲述节奏；确认后将三个代表案例放入后续章节。

## 复现命令
python code/ppt_illustrated_main_v2.py
powershell -File code/ppt_illustrated_main_export_v2.ps1
python code/ppt_illustrated_main_package_v2.py
'''
(O/'README_CN.md').write_text(readme,encoding='utf-8')
o=Image.new('RGB',(1800,6*355),'#e4eae6');d=ImageDraw.Draw(o)
for i,p in enumerate(imgs):o.paste(ImageOps.contain(Image.open(p),(594,334)),((i%3)*600,(i//3)*355+20));d.text(((i%3)*600+7,(i//3)*355+3),str(i+1),fill='#00553b')
o.save(O/'contact_sheet.jpg',quality=92)
body=''.join(f'<article id="s{i}"><h2>{i:02d} · {html.escape(t)}</h2><img loading="lazy" src="rendered/{p.name}" alt="{html.escape(t)}"></article>' for i,(p,t) in enumerate(zip(imgs,titles),1))
page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>CAMP · 可编辑组会汇报</title><style>body{margin:0;background:#eaf0ec;color:#183d34;font:16px "Microsoft YaHei",sans-serif}header{padding:18px 4vw;background:white;position:sticky;top:0;z-index:2;display:flex;gap:24px;align-items:center}h1{font-size:20px;margin:0 auto 0 0}a{color:#00553b}main{max-width:1560px;margin:24px auto;padding:0 20px}article{background:white;border:1px solid #d6e2da;border-radius:12px;margin-bottom:26px;overflow:hidden}h2{font-size:16px;padding:0 20px}img{width:100%;display:block}</style><header><h1>CAMP · 泛癌发现与 LYPLA1 · 18 页</h1>'''+f'<a href="{PPT}">下载可编辑 PPTX</a><a href="{PDF}">PDF 预览</a></header><main>'+body+'</main></html>'
(O/'index.html').write_text(page,encoding='utf-8')
V.mkdir(parents=True,exist_ok=True)
for p in O.iterdir():
 if p.is_file():shutil.copy2(p,V/p.name)
shutil.copytree(O/'rendered',V/'rendered',dirs_exist_ok=True)
stage=ROOT/'coordination/stages/BRCA.tsv'
if RUN not in stage.read_text(encoding='utf-8'):
 with stage.open('a',encoding='utf-8') as f:f.write('\t'.join(['BRCA','07_INTEGRATION',RUN,'illustrated_main_v2','DONE','18-slide editable CAMP main story before representative cases',str(O.relative_to(ROOT)).replace('\\','/'),'code/ppt_illustrated_main_v2.py','presentation/camp-fourcancer-lypla1-20260928','Presentation only; frozen statistics; WPS actual export verified','Review editable group meeting deck'])+'\n')
print(json.dumps(validation,ensure_ascii=False));print(V/PPT)
