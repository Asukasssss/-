from pathlib import Path
import json,re,shutil,hashlib,html
from pptx import Presentation
from PIL import Image,ImageOps,ImageDraw
import pymupdf as fitz
ROOT=Path(__file__).resolve().parents[1];RUN='20260929T_complete_cases_v3';O=ROOT/'results/BRCA/07_INTEGRATION'/RUN
V=Path('D:/CodexData/visualizations/2026/09/29/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/CAMP_complete_cases_v3')
PPT='CAMP_完整组会汇报_含三个案例_25页.pptx';PDF='CAMP_完整组会汇报_25页预览.pdf'
r=Presentation(O/PPT);old=Presentation(ROOT/'results/BRCA/07_INTEGRATION/20260929T_illustrated_main_v2/CAMP_泛癌与LYPLA1_插图精修可编辑版.pptx')
assert len(r.slides)==len(fitz.open(O/PDF))==25
for a,b in zip(old.slides,r.slides):assert a._element.xml==b._element.xml
audit=[]
for n,s in enumerate(r.slides,1):
 texts=[x for x in s.shapes if x.has_text_frame and x.text.strip()];pics=[x for x in s.shapes if x.shape_type==13]
 assert not any(x.width>r.slide_width*.95 and x.height>r.slide_height*.95 for x in pics)
 audit.append(dict(slide=n,text_boxes=len(texts),images=len(pics),objects=len(s.shapes)))
imgs=sorted((O/'rendered').glob('*.PNG'),key=lambda p:int(re.search(r'\d+',p.stem)[0]));assert len(imgs)==25
sheet=Image.new('RGB',(1600,4*465),'#e7eeea');d=ImageDraw.Draw(sheet)
for i,p in enumerate(imgs[18:]):
 x=i%2*800;y=i//2*465;sheet.paste(ImageOps.contain(Image.open(p),(794,447)),(x,y+18));d.text((x+8,y+2),str(i+19),fill='#00553b')
sheet.save(O/'added_pages_contact.jpg',quality=92)
(O/'validation.json').write_text(json.dumps(dict(status='PASS',slides=25,pdf_pages=25,wps_export=True,first_18_slides_xml_unchanged=True,new_tests=False,full_slide_flattening=False,audit=audit),indent=2),encoding='utf-8')
(O/'README_CN.md').write_text('''# CAMP 完整组会汇报：25页

## 本轮问题
将用户指定的代表案例v3接入插图精修版。
## 输入与范围
原18页插图精修版，加CAMP_representative_cases_v3中按网页顺序的7页。19–20为ASNS，21–23为SLC6A6，24–25为UCKL1。无案例空间页。
## 实际结果
25页PPTX及WPS实际导出PDF。前18页XML逐页完全一致；新增页面统一页眉、中文字体与页码。来源及哈希见source_manifest.json。
## 新手解释
文字、反应图与患者连线是可编辑对象。UMAP点云与图像为独立可替换图片，不是整页截图。
## 限制/反证
不重算统计、不改变证据措辞。生化注释不等于本队列机制证明；相关性不代表因果。
## 当前决定
采用用户明确指定v3七页，接在LYPLA1证据汇总之后。
## 下一步
按组会时间审阅整套讲述节奏。
## 复现命令
python code/ppt_complete_cases_v3.py
使用WPS打开生成PPTX，导出25页PNG到rendered及PDF。
python code/ppt_complete_cases_package_v3.py
''',encoding='utf-8')
(O/'code_manifest.json').write_text(json.dumps({p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT/'code').glob('ppt_complete_cases*v3.py')},indent=2),encoding='utf-8')
page='<meta charset="utf-8"><title>CAMP 完整组会汇报</title><style>body{background:#eaf0ec;font:16px Microsoft YaHei;color:#00553b;margin:0}header{position:sticky;top:0;background:white;padding:18px 4%;display:flex;gap:35px}a{color:#00553b}main{max-width:1500px;margin:auto}img{width:100%;display:block;margin:18px 0 30px}</style>'
page+=f'<header><b>CAMP 完整组会汇报 · 25页</b><a href="{PPT}">可编辑PPT</a><a href="{PDF}">PDF预览</a></header><main>'
page+=''.join(f'<img loading="lazy" src="rendered/{p.name}" alt="第{i}页">' for i,p in enumerate(imgs,1))+'</main>'
(O/'index.html').write_text(page,encoding='utf-8')
for p in O.iterdir():
 if p.is_file():shutil.copy2(p,V/p.name)
shutil.copytree(O/'rendered',V/'rendered',dirs_exist_ok=True)
stage=ROOT/'coordination/stages/BRCA.tsv'
if RUN not in stage.read_text(encoding='utf-8'):
 with stage.open('a',encoding='utf-8') as f:f.write('\t'.join(['BRCA','07_INTEGRATION',RUN,'complete_cases_v3','DONE','Append user-selected seven case pages; 25-slide deck',str(O.relative_to(ROOT)).replace('\\','/'),'code/ppt_complete_cases_v3.py','presentation/camp-fourcancer-lypla1-20260928','Frozen science; editable layout; WPS export verified','Review complete presentation'])+'\n')
print('PASS: 25 pages; first 18 unchanged')
