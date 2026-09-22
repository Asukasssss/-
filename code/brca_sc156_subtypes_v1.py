"""Descriptive subtype comparison of frozen, public donor-weighted Wu summaries."""
from pathlib import Path
import hashlib,json,subprocess
import numpy as np,pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill

ROOT=Path(__file__).resolve().parents[1]
RUN='20260922T064625Z_sc156_subtypes_v1'
OUT=ROOT/'results/BRCA/06_EXTERNAL'/RUN
OUT.mkdir(parents=True,exist_ok=True)
V='sc156_subtypes_v1'
S=['ER+','HER2+','TNBC']
T=['Malignant_epithelial','Nonmalignant_epithelial','Fibroblasts','Perivascular','Endothelial','Myeloid','T_NK_cells','B_cells','Plasma_cells']
CN=dict(zip(T,['恶性上皮','非恶性上皮','成纤维细胞','血管周细胞','内皮细胞','髓系细胞','T/NK细胞','B细胞','浆细胞']))
LABEL=['Malignant epi.','Nonmalignant epi.','Fibroblasts','Perivascular','Endothelial','Myeloid','T / NK','B cells','Plasma cells']
files=[ROOT/'results/BRCA/06_EXTERNAL'/r/'Wu2021_celltype_profiles.tsv' for r in ['20260919T140105Z_scRNA117_v1','20260922T023000Z_sc39_v1']]
masterpath=ROOT/'results/BRCA/06_EXTERNAL/20260922T023000Z_sc39_v1/genes156_sc39_appended.tsv'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def save(d,name):d.to_csv(OUT/name,sep='\t',index=False,na_rep='NA',lineterminator='\n')
d=pd.concat([pd.read_csv(p,sep='\t').assign(input_file=p.relative_to(ROOT).as_posix()) for p in files],ignore_index=True)
m=pd.read_csv(masterpath,sep='\t',dtype=str,keep_default_na=False)
genes=sorted(m.gene);assert len(genes)==156 and m.gene.is_unique
assert set(d.gene)==set(genes) and not d.duplicated(['gene','partition','celltype']).any()
x=d[d.partition.isin(['subtype:'+s for s in S])].copy()
x['input_run_id']=x.run_id;x['input_analysis_version']=x.analysis_version
x['run_id']=RUN;x['analysis_version']=V;x['analysis_type']='descriptive_subtype_cell_source'
x['subtype']=x.partition.str.replace('subtype:','',regex=False)
assert len(x)==156*3*9
save(x,'subtype_celltype_profiles.tsv')
coverage=x[['subtype','celltype','n','n_cells_total','n_cells_eligible','status','reason']].drop_duplicates()
assert len(coverage)==27
save(coverage,'subtype_celltype_coverage.tsv')
def top(z):
 z=z[z.status.eq('DONE')].sort_values(['effect','celltype'],ascending=[False,True])
 if len(z)<2:return dict(status='NOT_EVALUABLE',reason='fewer_than_two_covered_celltypes',top='NA',second='NA',effect=np.nan,gap=np.nan,n=0,detection=np.nan)
 a,b=z.iloc[0],z.iloc[1];reason='low_detection' if z.mean_detection_fraction.max()<.01 else ('tied_top' if np.isclose(a.effect,b.effect,rtol=0,atol=1e-12) else '')
 return dict(status='NOT_EVALUABLE' if reason else 'DONE',reason=reason,top=a.celltype if not reason else 'NA',second=b.celltype if not reason else 'NA',effect=a.effect,gap=a.effect-b.effect,n=int(a.n),detection=a.mean_detection_fraction)
rows=[];long=[];ranges=[];recon=[]
for g in genes:
 z=x[x.gene.eq(g)];groups={s:z[z.subtype.eq(s)] for s in S}
 shared=set.intersection(*[set(a.loc[a.status.eq('DONE'),'celltype']) for a in groups.values()])
 row={'gene':g,'n_shared_celltypes':len(shared),'interpretation':'descriptive_only;not_subtype_specificity_or_causality'}
 tops=[];stops=[]
 for s in S:
  a=top(groups[s]);b=top(groups[s][groups[s].celltype.isin(shared)])
  for k,val in a.items():row[s+'_'+k]=val
  row[s+'_top_cn']=CN.get(a['top'],'不可判定');row[s+'_shared_top']=b['top'];row[s+'_shared_status']=b['status']
  long.append(dict(gene=g,subtype=s,scope='all_evaluable_celltypes',**a));long.append(dict(gene=g,subtype=s,scope='shared_celltypes',**b))
  tops.append(a['top']);stops.append(b['top'])
 row['same_top_all3']=len(set(tops))==1 and 'NA' not in tops
 row['same_top_shared3']=len(set(stops))==1 and 'NA' not in stops
 row['status']='DONE' if 'NA' not in tops else 'NOT_EVALUABLE'
 row['pattern']='SAME_TOP' if row['same_top_shared3'] else ('DIFFERENT_TOP_DESCRIPTIVE' if 'NA' not in stops else 'LOW_DETECTION_OR_COVERAGE')
 # Equal-subtype and omit-one-subtype summaries: diagnostic on common categories only.
 if len(shared)>=2 and 'NA' not in stops:
  p=z[z.celltype.isin(shared)].pivot(index='celltype',columns='subtype',values='effect').reindex(columns=S)
  eq=p.mean(axis=1);best=eq.idxmax();row['equal_subtype_top']=best
  omitted=[]
  for s in S:
   t=p.drop(columns=s).mean(axis=1).idxmax();row['omit_'+s+'_top']=t;omitted.append(t)
  row['omit_one_equal_subtype_top_changes']=any(t!=best for t in omitted)
 else:
  row['equal_subtype_top']='NA';row['omit_one_equal_subtype_top_changes']='NA'
  for s in S:row['omit_'+s+'_top']='NA'
 rows.append(row)
 for c in T:
  a=z[z.celltype.eq(c)].set_index('subtype').reindex(S)
  b=dict(gene=g,celltype=c,status='DONE' if a.status.eq('DONE').all() else 'NOT_EVALUABLE',reason='' if a.status.eq('DONE').all() else 'at_least_one_subtype_insufficient_coverage')
  for s in S:b[s+'_mean']=a.loc[s,'effect'];b[s+'_n']=a.loc[s,'n'];b[s+'_detection']=a.loc[s,'mean_detection_fraction']
  if b['status']=='DONE':
   b['highest_mean_subtype']=a.effect.idxmax();b['lowest_mean_subtype']=a.effect.idxmin();b['mean_range']=a.effect.max()-a.effect.min()
   original=d[(d.gene==g)&(d.partition=='ALL')&(d.celltype==c)].iloc[0]
   reconstructed=np.average(a.effect,weights=a.n);assert int(a.n.sum())==int(original.n)
   recon.append(abs(reconstructed-original.effect))
  else:b.update(highest_mean_subtype='NA',lowest_mean_subtype='NA',mean_range=np.nan)
  b['effect_scale']='difference_of_mean_log1p_counts_per10k;not_fold_change';ranges.append(b)
summary=pd.DataFrame(rows);save(summary,'genes156_subtype_comparison.tsv');save(pd.DataFrame(long),'subtype_top_and_runner.tsv');save(pd.DataFrame(ranges),'within_celltype_subtype_comparison.tsv')
assert max(recon)<1e-10
joined=m.merge(summary.rename(columns={c:'Wu_subtype_'+c for c in summary if c!='gene'}),on='gene',how='left',validate='one_to_one')
assert joined[m.columns].equals(m);save(joined,'genes156_subtype_appended.tsv')
spec=dict(version=V,run_id=RUN,input_commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),scope='all156=150_current+6_historical;no_P_gate',subtypes={'ER+':11,'HER2+':5,'TNBC':10},unit='published_source_donor_label',normalization='frozen mean log1p(counts per10k), equal donor weight',min_cells_per_donor_type=20,min_source_labels_per_subtype_type=3,low_detection_top_flag=.01,new_hypothesis_tests=0,new_P_q=0,seed='NA_deterministic_reuse',uncertainty='No new bootstrap or intervals from aggregate means; existing quartiles and coverage retained',treatment='mixed as in original subtype partitions;not treatment-adjusted',omit_one='equal weight for subtype means on common categories;descriptive ranking sensitivity, not donor bootstrap',pal='not included: subtype crosswalk not connected',software={'python':__import__('sys').version,'pandas':pd.__version__,'numpy':np.__version__,'matplotlib':matplotlib.__version__})
(OUT/'analysis_spec.json').write_text(json.dumps(spec,indent=2),encoding='utf-8')
save(pd.DataFrame([dict(path=p.relative_to(ROOT).as_posix(),sha256=sha(p)) for p in files+[masterpath,Path(__file__)]]),'source_manifest.tsv')
# Four pages, every gene once, same gene scale across all 27 subtype-celltype combinations.
cmap=plt.get_cmap('viridis').copy();cmap.set_bad('#dddddd')
with PdfPages(OUT/'all156_subtype_expression_patterns.pdf') as pdf:
 for page in range(4):
  gg=genes[page*39:(page+1)*39];a=np.full((39,3,9),np.nan)
  for i,g in enumerate(gg):
   for j,s in enumerate(S):
    z=x[(x.gene==g)&(x.subtype==s)].set_index('celltype').reindex(T)
    a[i,j,:]=z.effect.where(z.status.eq('DONE')).to_numpy()
  mx=np.nanmax(a,axis=(1,2));rel=np.divide(a,mx[:,None,None],out=np.full_like(a,np.nan),where=mx[:,None,None]>0)
  fig,axs=plt.subplots(1,3,figsize=(18,15),sharey=True)
  for j,(ax,s) in enumerate(zip(axs,S)):
   im=ax.imshow(np.ma.masked_invalid(rel[:,j,:]),vmin=0,vmax=1,cmap=cmap,aspect='auto');ax.set_title(s+' (source labels: '+str(spec['subtypes'][s])+')');ax.set_yticks(range(39),gg);ax.set_xticks(range(9),LABEL,rotation=55,ha='right')
  fig.subplots_adjust(left=.09,right=.9,bottom=.15,top=.93,wspace=.07);cb=fig.add_axes([.92,.55,.012,.25]);fig.colorbar(im,cax=cb,label='Relative to each gene maximum across subtypes/types')
  fig.suptitle('Wu BRCA: all156 subtype expression patterns | page '+str(page+1),fontsize=16)
  fig.text(.09,.025,'Each gene normalized by its own maximum across all 3 subtype panels; not comparable between genes.\nGray: insufficient coverage or all-zero scale. No clustering, subtype hypothesis tests or treatment adjustment.',fontsize=10)
  fig.savefig(OUT/('subtype_expression_page'+str(page+1)+'.png'),dpi=160);pdf.savefig(fig);plt.close(fig)
notes=pd.DataFrame([dict(item='范围',meaning='当前150+历史6；全156保留；不是只分析显著基因'),dict(item='解读',meaning='最高类别/均值差为描述，不能称亚型特异或主要由某亚型导致'),dict(item='缺失',meaning='每供者类别至少20细胞；每亚型类别至少3标签；灰色不代表低表达'),dict(item='低检出',meaning='最大检测比例低于1%时不强行指定最高类别；并非阴性'),dict(item='图',meaning='每基因在三个分型全部类别中共同缩放；跨基因颜色不可比'),dict(item='限制',meaning='供者标签沿用作者；治疗混合；仅Wu；无新P/q或重采样')])
book=OUT/'BRCA_全156基因_单细胞分型比较.xlsx'
with pd.ExcelWriter(book,engine='openpyxl') as w:
 for name,z in [('说明',notes),('156基因分型对照',summary),('细胞内亚型表达',pd.DataFrame(ranges)),('分型细胞类别明细',x),('覆盖',coverage),('最高与次高',pd.DataFrame(long)),('156完整历史保留',joined)]:z.to_excel(w,sheet_name=name,index=False)
wb=load_workbook(book)
for ws in wb:
 ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D')
 for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(40,max(15,len(str(col[0].value))))
wb.save(book)
val=dict(status='DONE',genes=156,subtype_celltype_rows=len(x),covered_rows=int(x.status.eq('DONE').sum()),insufficient_rows=int(x.status.ne('DONE').sum()),same_top_all3=int(summary.same_top_all3.sum()),same_top_shared3=int(summary.same_top_shared3.sum()),pattern_counts=summary.pattern.value_counts().to_dict(),max_ALL_reconstruction_error=max(recon),ALL_checks=len(recon),all_historical_columns_preserved=True,new_P_q=0,plot_pages=4)
(OUT/'validation.json').write_text(json.dumps(val,indent=2),encoding='utf-8')
examples=[]
for g in ['LYPLA1','LYPLA2','ABHD12','ENPP2','LPCAT1','ASNS','GLS','GPCPD1','GPI','NNMT','PYCR1']:
 a=summary.set_index('gene').loc[g];examples.append('|'+g+'|'+'|'.join(a[s+'_top_cn'] for s in S)+'|')
readme=f'''# BRCA全156基因：Wu单细胞分型表达背景

## 本轮问题
ER+、HER2+、TNBC中候选基因表达背景是否相似？仅描述排名和同类细胞表达，不证明亚型特异机制或驱动。

## 输入与范围
复用Wu旧117与新增39的既有供者等权汇总，当前150+历史6全部保留。ER+11、HER2+5、TNBC10来源供者标签；实际每类可评估数见覆盖表。按20细胞/供者类别、3标签/亚型类别限制。没有重读患者矩阵或重新聚类。Pal未接入可靠亚型连接，暂不纳入。

## 实际结果
共{len(x)}个基因×分型×类别组合，{val['covered_rows']}可评估，{val['insufficient_rows']}覆盖不足。各分型自身可评估类别下，{val['same_top_all3']}基因三组最高类别一致；限制共同类别后为{val['same_top_shared3']}。两口径不能混用，低检出不指定第一名。

|基因|ER+最高类别|HER2+最高类别|TNBC最高类别|
|---|---|---|---|
'''+ '\n'.join(examples)+'''

## 新手解释
先看156基因分型对照，再看同一细胞类别下三个分型的均值、检测比例和实际供者数。最高类别变化不等于该基因只在一个分型起作用；最高类别不变也不等于表达量不变。热图按每基因在全部分型和类别中的最大值缩放，颜色用于看同一基因，不能跨基因比绝对表达。原数值完整保留。

## 限制与反证
具体读表示例：LPCAT1在HER2+的B细胞均值0.129842、髓系0.129261，排名差距约0.000581，不能解释为明确的B细胞特异作用。ASNS在ER+中浆细胞最高，但HER2+浆细胞仅2个合格来源标签而未展示；因此三组自身类别下的最高排名不是完全对等比较。ASNS恶性上皮均值ER+0.0443、HER2+0.1277、TNBC0.1538，仅为描述，恶性上皮实际供者标签分别9、3、8，尚无正式亚型差异结论。
HER2+样本小，有类别缺失；不同类别覆盖不同供者。描述性均值差没有显著性或精确性认证，没有新P/q。排名间小差距应看top-second gap，不制造硬性生物学分组。低检出标记不是删除。混合治疗背景未调整，不将本结果直接称亚型效应。数据为Wu一个研究；Pal总体结果不是亚型复现。供者标签沿用原定义。
equal_subtype_top及omit字段只在共同类别对三个亚型均值等权、再依次省略一个亚型，检查排名是否改变；不是患者留出或bootstrap，不证明某亚型驱动总体信号。非线性的log表达均值差不能称浓度/表达倍数。

## 当前决定
全部156基因附加Wu_subtype前缀列，原列逐值保留；未因排名改变或缺测淘汰任何基因。当前150与历史6的范围列继续保留。

## 下一步
用细胞背景辅助形成有限候选问题。需要正式亚型差异时另立供者级模型并处理治疗和多重检验；本轮不自动启动。

## 复现
在仓库运行 `python code/brca_sc156_subtypes_v1.py`。来源文件和代码SHA256见source_manifest；检查原ALL均值是否可由各亚型按合格供者数加权重建，原列是否逐值保留。公开输出均为群组汇总。
'''
(OUT/'README_CN.md').write_text(readme,encoding='utf-8');(OUT/'.gitattributes').write_text('* -text\n')
save(pd.DataFrame([dict(file=p.name,sha256=sha(p)) for p in sorted(OUT.iterdir()) if p.is_file() and p.name!='checksums.tsv']),'checksums.tsv')
print(json.dumps(val));print('\n'.join(examples))
