"""Advisor discussion views: all117 retained; no new composite ranking or inference."""
from pathlib import Path
import csv,json,hashlib
import pandas as pd,numpy as np
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
ROOT=Path(__file__).resolve().parents[1];RUN='20260921T093756Z_candidate_discussion_v1';OUT=ROOT/'results/BRCA/07_INTEGRATION'/RUN
OLD=ROOT/'results/BRCA/06_EXTERNAL/20260921T072503Z_oslo_tang_external_v1'
def read(p):
 with p.open(encoding='utf-8-sig',newline='') as f:
  r=csv.DictReader(f,delimiter='\t');return r.fieldnames,list(r)
def write(p,c,r):
 with p.open('w',encoding='utf-8',newline='') as f:
  w=csv.DictWriter(f,c,delimiter='\t',lineterminator='\n');w.writeheader();w.writerows(r)
def fmt(x):return '未测/不可评估' if pd.isna(x) else f'{x:.4g}'
UPDATES={
'MDH1':('新增Tang支持；跨队列一致性有限','苹果酸关联Tang较强，FUSCC接近零；GSE42568肿瘤RNA较低；既有功能记录尚未确认MDH1特异干预','可供选题讨论；先核实已有MDH1特异证据，不借用MDH2；本轮不追加机制'),
'MDH2':('与MDH1分开保留','MDH1新关联不能转移为MDH2证据；GSE42568 MDH2组间差异未获支持；保留MDH2自己的既有功能记录','继续按MDH2自身适用证据讨论，不因MDH1的新结果自动升降级'),
'PNP':('新增Tang正向线索；两种代谢物分列','次黄嘌呤q0.0489、鸟嘌呤q0.07335共享同20例；不能算两次独立验证；既有功能支持存在联合治疗背景限制','列为值得关注；本轮停止于候选讨论，不启动嘌呤机制课题'),
'GPI':('CAMP支持保留；Tang未支持正关联','FUSCC为缺项；Tang的G6P和F6P已评估且未支持正关联，不能再统称外部缺测；宽区间不证明反向作用','更新候选判断；不删除，不为追求显著性持续换队列或模型'),
'ASNS':('保留既有支持；新增Tang不精确结果','Tang谷氨酰胺负向但不显著；FUSCC与小鼠Asns干预各有适用范围，仍未证明人类谷氨酰胺机制','作为已有数据专题保留；不因Tang未显著取消，也不扩写已闭合机制'),
'GLS':('保留既有支持；新增Tang不精确结果','Tang谷氨酰胺负向但不显著；不改变FUSCC条件关联；特定亚型/营养背景限制仍在','保持原限定专题；不把谷氨酰胺剥夺替代GLS自身干预'),
'SLC6A8':('保留既有支持；新增Tang不精确结果','肌酸在Tang正向但不显著；功能背景和转运体特异性需要区分','保留候选，按既有适用模型讨论；本轮不补新机制'),
'GPCPD1':('保留功能与CAMP背景；外部证据有限','Tang GPC负向但区间跨零；既有FUSCC未复现结果不覆盖；不能宣称跨队列稳定成立','作为参照候选保留；不再把已有功能当首次发现')}
def main():
 cols,rows=read(OLD/'all117_comparison_external_appended.tsv');assert len(rows)==117 and len(cols)==224
 relation=pd.read_csv(OLD/'all174_cross_cohort_context.tsv',sep='\t');rna=pd.read_csv(OLD/'GSE42568_all117_RNA.tsv',sep='\t').set_index('gene');loo=pd.read_csv(OUT/'three_relationship_LOO_summary.tsv',sep='\t')
 added=['discussion_v1_position','discussion_v1_conflicts','discussion_v1_next_scope','discussion_v1_source'];brief=[]
 for r in rows:
  g=r['gene'];u=UPDATES.get(g,('保留原候选池；本轮未专项重评','新队列未显著或缺测不作淘汰；患者关联、RNA背景、功能和细胞来源分层保留','沿用既有适用证据安排，不新增统一显著性门槛'))
  r.update(dict(zip(added,[*u,OUT.relative_to(ROOT).as_posix()])))
  rel=relation[relation.gene.eq(g)];rn=rna.loc[g]
  def association(ef,q):return '\n'.join(f'{t.metabolite_name}: ρ={fmt(getattr(t,ef))}; q={fmt(getattr(t,q))}' for t in rel.itertuples())
  lr=loo[loo.gene.eq(g)];lstr='\n'.join(f'{t.metabolite_name}: 留出范围{t.loo_rho_min:.3f}～{t.loo_rho_max:.3f}；方向保持{t.sign_retained}/20' for t in lr.itertuples()) or '本轮未安排；不是阴性'
  brief.append({'基因':g,'本轮定位':u[0],'CAMP直接代谢物及方向':r['linked_metabolites_and_direction'],'CAMP原未调关联（逐关系）':association('CAMP_unadjusted_rho','CAMP_unadjusted_q'),'FUSCC原关联（逐关系）':association('FUSCC_primary174_rho','FUSCC_primary174_q'),'Tang原关联（逐关系）':association('effect','q_value'),'Tang可评估性':r['Tang_external_status'],'GSE42568肿瘤减正常log2差':rn.effect if pd.notna(rn.effect) else 'NA','GSE42568 RNA q':rn.q_value if pd.notna(rn.q_value) else 'NA','既有功能类别（未新增文献审核）':r['review_v4_functional_class'],'既有功能摘要':r['review_v4_functional_summary'],'既有功能限制':r['review_v4_counterevidence_and_limits'],'既有细胞背景':f"Wu: {r['rob_Wu2021_top_lineage']}; Pal: {r['rob_Pal2021_reprocessed_top_lineage']}; 稳定标记: {r['rob_two_tumor_source_stable_descriptive']}",'本轮逐病例留出':lstr,'本轮矛盾及缺项':u[1],'本轮下一步范围':u[2],'历史工作优先级（未重排）':r['review_v4_work_priority'],'历史最小工作（非新任务）':r['review_v4_next_minimal_action'],'功能来源':r['review_v4_source_urls'] or 'NA'})
 write(OUT/'all117_comparison_discussion_appended.tsv',cols+added,rows);_,orig=read(OLD/'all117_comparison_external_appended.tsv');assert [{c:r[c] for c in cols} for r in rows]==orig
 brief=sorted(brief,key=lambda r:r['基因']);write(OUT/'advisor_all117.tsv',list(brief[0]),brief);focus=[next(r for r in brief if r['基因']==g) for g in UPDATES];write(OUT/'advisor_focus8.tsv',list(brief[0]),focus)
 # Full relationship table is copied without changes, not only selected q<.05 rows.
 (OUT/'all174_cross_cohort_context.tsv').write_bytes((OLD/'all174_cross_cohort_context.tsv').read_bytes())
 scopes=[{'层次':'CAMP','检验范围':'沿用历史各自家族','缺项处理':'保持原文件，不重算','解释':'展示原未调关系；不是调整模型替代品'}, {'层次':'FUSCC','检验范围':'174计划项','缺项处理':'历史内部P=1占位，公开缺项P/q为NA','解释':'原q原样保留'}, {'层次':'Tang','检验范围':'174计划关系中163可评估项','缺项处理':'未评估项无P/q，不进163项BH','解释':'与FUSCC不同；不当统一评分'}, {'层次':'GSE42568','检验范围':'117计划基因中112可评估项','缺项处理':'无唯一探针者NA','解释':'RNA背景，不是代谢关联'}, {'层次':'本轮留出','检验范围':'固定3条关系，各20次留出','缺项处理':'无新增P/q','解释':'范围不是置信区间；20次不是独立验证'}]
 write(OUT/'FDR_scope_notes.tsv',list(scopes[0]),scopes)
 wb=Workbook();wb.remove(wb.active)
 sheets=[('阅读说明',[{'项目':'用途','说明':'给导师讨论的完整候选比较，不是新靶点排名；重点8个由本轮问题指定，其余109个保留。'}, {'项目':'统计','说明':'P/q全部沿用；唯一新数值为三条关系逐病例留出ρ及范围。'}, {'项目':'功能','说明':'功能文献、细胞背景来自已有表，本轮未逐篇重新审核。MDH1与MDH2分开。'}, {'项目':'存放','说明':'逐病例散点和留出明细仅存server165；本Excel为汇总。'}]),('重点8个',focus),('全117候选',brief),('全174关系',relation.replace({np.nan:None}).to_dict('records')),('3条留出汇总',loo.replace({np.nan:None}).to_dict('records')),('RNA117背景',rna.reset_index().replace({np.nan:None}).to_dict('records')),('校正范围',scopes)]
 for name,data in sheets:
  ws=wb.create_sheet(name);headers=list(data[0]);ws.append(headers)
  for rr in data:
   ws.append([rr.get(k) for k in headers])
  ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='244A64');cell.alignment=Alignment(wrap_text=True,vertical='center')
  ws.row_dimensions[1].height=38
  for row in ws.iter_rows(min_row=2):
   for cell in row:
    cell.alignment=Alignment(wrap_text=True,vertical='top')
    if isinstance(cell.value,float):cell.number_format='0.00E+00' if 0<abs(cell.value)<.001 else '0.0000'
   ws.row_dimensions[row[0].row].height=100 if name in ['重点8个','全117候选'] else 32
  for i,h in enumerate(headers,1):ws.column_dimensions[get_column_letter(i)].width=16 if i==1 else 38 if any(w in h for w in ['关联','摘要','限制','背景','矛盾','来源','说明','范围']) else 23
 export=ROOT/'local_deliverables'/RUN;export.mkdir(parents=True,exist_ok=True);book=export/'BRCA_导师讨论_全117候选.xlsx';wb.save(book);check=load_workbook(book,read_only=True,data_only=True);assert check['全117候选'].max_row==118 and check['全174关系'].max_row==175 and check['重点8个'].max_row==9;check.close()
 val=dict(status='DONE',all117_retained=True,historical_columns_unchanged=len(cols),new_columns=4,focus_genes=8,focus_not_ranking=True,all174_relation_rows=174,workbook_sheets=7,old_P_q_unchanged=True,new_inferential_tests=0);(OUT/'integration_validation.json').write_bytes((json.dumps(val,indent=2)+'\n').encode())
 sf,sr=read(ROOT/'coordination/stages/BRCA.tsv');sr=[r for r in sr if r['run_id']!=RUN]
 for stage,scope in [('04_ROBUSTNESS','3 fixed Tang relations;20 leaveouts each;no new P/q;private scatter only'),('07_INTEGRATION','All117 preserved;advisor comparison;8 discussion examples not ranking;GPI nonconfirmation and MDH1 RNA decrease explicit')]:
  sr.append(dict(zip(sf,['BRCA',stage,RUN,'candidate_discussion_v1','DONE',scope,OUT.relative_to(ROOT).as_posix(),'code/brca_candidate_discussion_v1.py','analysis/brca-functional-review-20260919','Same20cases;post-selection;no claim of cross-cohort replication or full functional re-review','Use discussion table;Oslo crosswalk stays bounded pending;no mechanism expansion'])))
 write(ROOT/'coordination/stages/BRCA.tsv',sf,sorted(sr,key=lambda r:(r['stage_id'],r['run_id'])))
 scripts=['brca_tang_three_loo_v1.py','brca_candidate_discussion_v1.py'];write(OUT/'code_manifest.tsv',['path','sha256'],[dict(path='code/'+n,sha256=hashlib.sha256((ROOT/'code'/n).read_bytes()).hexdigest()) for n in scripts])
 write(OUT/'integration_source_manifest.tsv',['path','sha256'],[dict(path=str(f.relative_to(ROOT)),sha256=hashlib.sha256(f.read_bytes()).hexdigest()) for f in [OLD/'all117_comparison_external_appended.tsv',OLD/'all174_cross_cohort_context.tsv',OLD/'GSE42568_all117_RNA.tsv',ROOT/'code/brca_all174_external_v1.py']])
 for f in OUT.iterdir():
  if f.suffix in {'.json','.tsv','.md','.txt'}:f.write_bytes(f.read_bytes().replace(b'\r\n',b'\n'))
 ch=OUT/'checksums.sha256';ch.write_bytes(''.join(f'{hashlib.sha256(f.read_bytes()).hexdigest()}  {f.name}\n' for f in sorted(OUT.iterdir()) if f.is_file() and f!=ch).encode());print(json.dumps(val))
if __name__=='__main__':main()
