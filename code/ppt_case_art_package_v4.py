from pathlib import Path
import json,re,shutil,hashlib
from pptx import Presentation
from PIL import Image,ImageOps
import pymupdf as fitz
ROOT=Path(__file__).resolve().parents[1];RUN='20260929T_case_art_v4';O=ROOT/'results/BRCA/07_INTEGRATION'/RUN
V=Path('D:/CodexData/visualizations/2026/09/29/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/CAMP_case_art_v4')
PPT='CAMP_完整组会汇报_生化插图精修_25页.pptx';PDF='CAMP_生化插图精修_预览.pdf'
r=Presentation(O/PPT);old=Presentation(ROOT/'results/BRCA/07_INTEGRATION/20260929T_complete_cases_v3/CAMP_完整组会汇报_含三个案例_25页.pptx')
assert len(r.slides)==len(fitz.open(O/PDF))==25
for n,(a,b) in enumerate(zip(old.slides,r.slides),1):
 if n not in [19,21,24]:assert a._element.xml==b._element.xml
 else:
  def lower(s):return [(x.text,x.left,x.top,x.width,x.height) for x in s.shapes if x.has_text_frame and x.top>3.75*914400]
  assert lower(a)==lower(b)
imgs=sorted((O/'rendered').glob('*.PNG'),key=lambda p:int(re.search(r'\d+',p.stem)[0]));assert len(imgs)==25
sheet=Image.new('RGB',(1600,2700),'white')
for i,n in enumerate([19,21,24]):sheet.paste(Image.open(imgs[n-1]),(0,i*900))
sheet.save(O/'three_refined_pages.jpg',quality=92)
(O/'validation.json').write_text(json.dumps(dict(status='PASS',slides=25,wps_export=True,edited_slides=[19,21,24],other_22_slides_xml_unchanged=True,result_text_and_positions_unchanged=True,illustration_alpha=Image.open(O/'assets/scientific_cutouts.png').getextrema()[-1],new_tests=False,visual_review='Three updated slides reviewed from actual WPS exports'),indent=2),encoding='utf-8')
(O/'README_CN.md').write_text('''# CAMP 生化作用插图精修版

## 本轮问题
将三个案例中单调的反应文字区改为插图与可编辑标签组合。
## 输入与范围
基于25页完整PPT，更新19、21、24页的ASNS、SLC6A6、UCKL1生化示意。
## 实际结果
内置image_gen生成真实透明背景素材，保留alpha，以PPT原生图片裁剪置入。其余22页XML不变，三张修改页下部结果文字与位置不变。25页已由WPS实际导出。
## 新手解释
插图独立移动、裁剪和替换；标签和反应箭头是原生可编辑对象。生成图不替代任何数据图。完整生成提示词见imagegen_prompt.txt。
## 限制/反证
分子和蛋白是概念插画，不代表精确化学结构或实际蛋白结构。保留原来源和机制证据边界，未增加实验结论。
## 当前决定
交付完整PPT、PDF、预览与插图源文件。
## 下一步
组会放映审阅。
## 复现命令
python code/ppt_case_art_v4.py
WPS导出PNG至rendered，并导出PDF。
python code/ppt_case_art_package_v4.py
''',encoding='utf-8')
(O/'code_manifest.json').write_text(json.dumps({p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT/'code').glob('ppt_case_art*v4.py')},indent=2),encoding='utf-8')
page='<meta charset="utf-8"><title>CAMP 生化插图精修</title><style>body{background:#eaf0ec;font:16px Microsoft YaHei;color:#00553b;margin:0}header{position:sticky;top:0;background:white;padding:18px 4%;display:flex;gap:28px}a{color:#00553b}main{max-width:1500px;margin:auto}img{width:100%;display:block;margin:18px 0 30px}</style>'
page+=f'<header><b>CAMP 生化插图精修 · 25页</b><a href="{PPT}">可编辑PPT</a><a href="{PDF}">PDF预览</a><a href="#s19">ASNS</a><a href="#s21">SLC6A6</a><a href="#s24">UCKL1</a></header><main>'
page+=''.join(f'<img id="s{i}" loading="lazy" src="rendered/{p.name}" alt="第{i}页">' for i,p in enumerate(imgs,1))+'</main>'
(O/'index.html').write_text(page,encoding='utf-8')
for p in O.iterdir():
 if p.is_file():shutil.copy2(p,V/p.name)
for folder in ['assets','rendered']:shutil.copytree(O/folder,V/folder,dirs_exist_ok=True)
stage=ROOT/'coordination/stages/BRCA.tsv'
if RUN not in stage.read_text(encoding='utf-8'):
 with stage.open('a',encoding='utf-8') as f:f.write('\t'.join(['BRCA','07_INTEGRATION',RUN,'case_art_v4','DONE','Illustrated role panels for three cases',str(O.relative_to(ROOT)).replace('\\','/'),'code/ppt_case_art_v4.py','presentation/camp-fourcancer-lypla1-20260928','Imagegen conceptual cutouts; frozen science; WPS verified','Review illustrated deck'])+'\n')
print('PASS: 25 pages; three role panels illustrated; results unchanged')
