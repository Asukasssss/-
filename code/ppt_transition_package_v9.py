from pathlib import Path
import shutil,json,re
from pptx import Presentation
import pymupdf as fitz
ROOT=Path(__file__).resolve().parents[1];RUN='20260929T_transition_outlook_v9';O=ROOT/'results/BRCA/07_INTEGRATION'/RUN
V=Path('D:/CodexData/visualizations/2026/09/29/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/CAMP_transition_outlook_v9')
PPT='CAMP_过渡与展望完整版_27页.pptx';PDF='CAMP_27页预览.pdf'
r=Presentation(O/PPT);assert len(r.slides)==len(fitz.open(O/PDF))==27
old=Presentation(V.parent/'CAMP_reaction_inplace_v8/CAMP_反应示意原位替换_25页.pptx')
for i,a in enumerate(old.slides):
 b=r.slides[i if i<18 else i+1]
 def texts(s):return [x.text for x in s.shapes if x.has_text_frame and not(x.text.strip().isdigit() and x.left>850*12700 and x.top>490*12700)]
 assert texts(a)==texts(b)
(O/'validation.json').write_text(json.dumps(dict(status='PASS',slides=27,added=[19,27],old_slide_text_unchanged_except_numbering=True,wps_export=True,visual_review='Added pages inspected; transition banner wrapping corrected',new_statistics=False),indent=2),encoding='utf-8')
(O/'README_CN.md').write_text('''# 过渡与展望

本轮问题：用户要求在LYPLA1汇总与单癌案例之间补过渡，末尾增加展望。

输入与范围：25页原位反应替换版，插入第19页，末尾增加第27页。

实际结果：27页可编辑PPT；两页沿用中大绿与三列统一版式，引用现有独立插图。已由WPS导出检查。

新手解释：过渡从跨癌候选引向癌种背景，明确三案例仍需分别解释；展望将未来计划分成优先关系、独立证据、功能联系。

限制/反证：未称癌种特异机制成立；未来计划不代表已完成结果。公开扰动证据优先，实验验证视条件开展。

当前决定：采用27页完整版，所有原页文字除页码外保持不变。

下一步：组会审阅与讲述时间调整。

复现：python code/ppt_transition_outlook_v9.py；WPS导出PNG与PDF；python code/ppt_transition_package_v9.py。
''',encoding='utf-8')
imgs=sorted((O/'rendered').glob('*.PNG'),key=lambda p:int(re.search(r'\d+',p.stem)[0]))
page='<meta charset="utf-8"><title>CAMP 27页完整版</title><style>body{background:#eaf0ec;color:#00553b;font:16px Microsoft YaHei;margin:0}header{background:white;padding:20px;position:sticky;top:0}a{margin-right:25px;color:#00553b}main{max-width:1500px;margin:auto}img{width:100%;margin:20px 0}</style>'
page+=f'<header><a href="{PPT}">27页可编辑PPT</a><a href="{PDF}">PDF</a><a href="#s19">过渡页</a><a href="#s27">展望页</a></header><main>'
page+=''.join(f'<img id="s{i}" loading="lazy" src="rendered/{p.name}">' for i,p in enumerate(imgs,1))+'</main>'
(O/'index.html').write_text(page,encoding='utf-8')
for p in O.iterdir():
 if p.is_file() and not p.name.startswith('~'):shutil.copy2(p,V/p.name)
shutil.copytree(O/'rendered',V/'rendered',dirs_exist_ok=True)
stage=ROOT/'coordination/stages/BRCA.tsv'
if RUN not in stage.read_text(encoding='utf-8'):
 with stage.open('a',encoding='utf-8') as f:f.write('\t'.join(['BRCA','07_INTEGRATION',RUN,'transition_outlook_v9','DONE','27-page deck with transition and outlook',str(O.relative_to(ROOT)).replace('\\','/'),'code/ppt_transition_outlook_v9.py','presentation/camp-fourcancer-lypla1-20260928','Presentation only; WPS verified','Review complete deck'])+'\n')
print('PASS: 27 slides, prior result text preserved')
