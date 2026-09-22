"""Frozen aggregate review, curated overlapping modules; no new inference tests."""
from pathlib import Path
import subprocess,io,json,hashlib
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
R=Path(__file__).resolve().parents[1];SHA='8952c54b61709c875ec83881b66ee43bd1de59f3'
B='results/BRCA/07_INTEGRATION/20260922T140000Z_report_v1/appendix/'
O=R/'results/BRCA/07_INTEGRATION/20260922T160000Z_module_review_v1'
assert not O.exists();O.mkdir(parents=True)
sources=[];data={}
for name in ['metabolites','association','rna','source_status','subtypes','relations','genes']:
 path=B+name+'.tsv';b=subprocess.check_output(['git','show',SHA+':'+path],cwd=R);data[name]=pd.read_csv(io.BytesIO(b),sep='\t',dtype=str,keep_default_na=False);(O/(name+'_frozen.tsv')).write_bytes(b);sources.append(dict(path=path,input_commit=SHA,sha256=hashlib.sha256(b).hexdigest()))
# Named modules are review lenses, not tested pathways or exhaustive assignments.
M=[
('M01','氨基酸氮分配与谷氨酸分支',['glutamine','glutamate','aspartate','asparagine','proline','arginine','ornithine','citrulline'],'PRIORITY_DISCUSSION','谷氨酰胺低、谷氨酸/天冬氨酸/脯氨酸高；比较多条直接关系，不锁定ASNS。','天冬酰胺未显著；ASNS以外多为P支持；细胞来源不全相同。','核对同病例变化组合；比较ASNS/BCAT2/GOT1/2/PYCR1现有来源。'),
('M02','乙醇胺至磷酸乙醇胺短链',['ethanolamine','phosphoethanolamine'],'PRIORITY_DISCUSSION','两代谢物升高且敏感性保留；ETNK1与PCYT2分别有主/敏感性FDR关联。','ETNK1来源不稳/跨研究不同，PCYT2恶性上皮稳定；PCYT2配对RNA不显著，不能推断同一细胞合成增加。','核对两代谢物同病例共升；报告缺失CDP-乙醇胺/终产物证据。'),
('M03','胆碱与溶血磷脂重塑背景',['choline','phosphocholine',"cytidine 5'-diphosphocholine",'glycerophosphorylcholine (GPC)','1-palmitoyl-GPC (16:0)','1-stearoyl-GPC (18:0)','1-oleoylglycerophosphocholine','2-palmitoyl-GPC (16:0)','2-oleoylglycerophosphocholine','2-arachidonoylglycerophosphocholine'],'PRIORITY_DISCUSSION','胆碱/GPC降低而多种酰基GPC升高；LYPLA1、GPCPD1、ENPP2等来源不同。','不是一条线性链；1-stearoyl-GPC仅16可用对且敏感性q=0.068；外部精确脂质匹配缺项。','将精确脂质身份、可用对数与细胞来源并列；不推断细胞间转运。'),
('M04','嘌呤降解与回收背景',['adenosine','inosine','hypoxanthine','guanine','guanosine','xanthine','xanthosine','urate'],'PRIORITY_DISCUSSION','肌苷低、次黄嘌呤/鸟嘌呤高、黄嘌呤低；PNP/鸟嘌呤有内部FDR关联，PNP偏内皮。','尿酸未变；回收分支存在；Tang正向结果来自小样本且新版q不再过线。','核对共同变化与PNP/HPRT1/转运者不同来源，不把整条链称加速。'),
('M05','丝氨酸甘氨酸与同工酶相反RNA',['serine','glycine'],'PRIORITY_DISCUSSION','两代谢物稳健升高；SHMT1为39/45降低，SHMT2为40/45升高。','两同工酶与这些代谢物关联均无FDR支持；定位不能证明区室通量，SHMT2分型来源并非完全一致。','作为区室/细胞背景假说；并列两个同工酶，不声称同一患者已发生功能替代。'),
('M06','烟酰胺与甲基化产物',['nicotinamide','1-methylnicotinamide','S-adenosylhomocysteine (SAH)','methionine'],'PRIORITY_DISCUSSION','NNMT与烟酰胺负关联、与1-MNA正关联；成纤维来源稳定；两代谢物配对升高。','NNMT RNA却30/45降低；组成调整后关联减弱；SAM未纳本表，不能推断甲基供体耗竭。','明确组间变化与肿瘤内部相关不同；以细胞背景解释矛盾，不掩盖反证。'),
('M07','糖酵解节点与磷酸戊糖背景',['glucose','glucose 6-phosphate','fructose-6-phosphate','3-phosphoglycerate','pyruvate','lactate','6-phosphogluconate','ribose'],'CONTEXT_RETAIN','G6P/F6P/乳酸升高；GPI与前两者有关联并有稳定恶性上皮背景。','葡萄糖/3PG/丙酮酸不显著；6PG仅7可用对不可算；Tang不支持GPI正关联；不是全通路激活。','保留内部模块，不为GPI反复换模型。'),
('M08','多元醇与果糖',['glucose','sorbitol','fructose'],'MISSINGNESS_LIMITED','主分析山梨醇/果糖升高，SORD关联和RNA/来源初看较好。','山梨醇仅8、果糖14可用对，疾病差异不支持；SORD相关敏感性明显减弱，山梨醇rho近零。','降低优先级；先处理可用性限制，不能用稳定单细胞来源挽救代谢链。'),
('M09','氨基糖与糖基化前体背景',['N-acetylglucosamine','N-acetylglucosamine 6-phosphate','mannose','mannose 6-phosphate'],'MISSINGNESS_LIMITED','主分析多个糖前体升高，NAGK/PGM3/PMM2有P线索；NAGK髓系稳定。','GlcNAc及GlcNAc6P可用值疾病结果均未支持；PMM2稳定性不足，不能称糖基化增强。','保留补充；优先说明缺失而非立即做机制。'),
('M10','半胱氨酸与抗氧化背景',['cysteine','cystine','glutamate','glycine','glutathione, reduced (GSH)','glutathione, oxidized (GSSG)','5-oxoproline','taurine'],'CONTEXT_RETAIN','半胱氨酸等升高；CDO1 RNA45/45降低且成纤维排名稳定。','GSH可用值仅11对且不显著，GSSG不显著；CDO1检出不高；不能证明铁死亡抑制或抗氧化通量增强。','保留代谢物/RNA反向现象；不直接算处理尺度GSH/GSSG生化比值。'),
('M11','色氨酸犬尿氨酸与髓系背景',['tryptophan','kynurenine'],'CONTEXT_RETAIN','犬尿氨酸高；KYNU内部关联稳定，KYNU/KMO偏髓系。','色氨酸未显著；KYNU外部关联点估计反向且未显著；关键中间体未测，不能称免疫抑制机制。','作为组织组成问题保留，尚不定义连续机制链。'),
('M12','肌酸肌酐',['creatine','creatinine'],'CONTEXT_RETAIN','两者配对降低且敏感性保留；连接SLC6A8等历史外部证据。','两者非同一运输底物；不能由丰度推断摄取活性或肿瘤能量短缺。','保留代谢模块，与氮代谢大通路区别。'),
('M13','TCA节点不均一变化',['citrate','alpha-ketoglutarate','succinate','fumarate','malate'],'CONTEXT_RETAIN','琥珀酸低/苹果酸高；FH/MDH2有内部P线索，MDH1有Tang小队列信号。','中间多个节点不显著；不能称整个TCA加快/受抑，也不能借MDH2证据解释MDH1。','保留异质变化与缺口，不把非连续节点画成闭合链。'),
('M14','尿嘧啶降解',['uridine','uracil','5,6-dihydrouracil','beta-alanine'],'MISSINGNESS_LIMITED','尿嘧啶/二氢尿嘧啶/β丙氨酸主分析升高。','后两项可用对少；β丙氨酸敏感性q=0.103；中间节点不全，无新通量证据。','待补支持，不与覆盖充分的短链同级。'),
('M15','赖氨酸与AASS背景',['lysine','2-aminoadipate','N6-acetyllysine'],'CONTEXT_RETAIN','赖氨酸与AASS RNA降低、内部正关联。','2-氨基己二酸主分析不显著；AASS单细胞低检出，不能当完整降解链。','保留患者线索，暂停仅围绕AASS扩展。')]
members=[];decisions=[]
for mid,title,names,tier,obs,limit,nxt in M:
 decisions.append(dict(module_id=mid,module=title,decision=tier,observation=obs,limitations=limit,next_minimum=nxt))
 for name in names:members.append(dict(module_id=mid,module=title,metabolite_name=name))
mem=pd.DataFrame(members);dec=pd.DataFrame(decisions)
assert set(mem.metabolite_name)<=set(data['metabolites'].metabolite_name)
joined={}
for name in ['metabolites','association','relations']:
 joined[name]=mem.merge(data[name],on='metabolite_name',how='inner',validate='many_to_many');joined[name].to_csv(O/('module_'+name+'.tsv'),sep='\t',index=False)
gmap=joined['relations'][['module_id','module','gene']].drop_duplicates()
for name in ['rna','source_status','subtypes']:
 joined[name]=gmap.merge(data[name],on='gene',how='left');joined[name].to_csv(O/('module_'+name+'.tsv'),sep='\t',index=False)
summary=[]
for mid,title,*_ in M:
 a=joined['association'];a=a[(a.module_id==mid)&(a.analysis_type=='processed_Spearman262')];m=joined['metabolites'];m=m[(m.module_id==mid)&(m.analysis_type=='paired_processed')]
 summary.append(dict(module_id=mid,module=title,features=len(m),relations=len(a),genes=a.gene.nunique(),metabolite_P_lt005=int((pd.to_numeric(m.p_value,errors='coerce')<.05).sum()),metabolite_q_lt005=int((pd.to_numeric(m.q_value,errors='coerce')<.05).sum()),association_P_lt005=int((pd.to_numeric(a.p_value,errors='coerce')<.05).sum()),association_q_lt005=int((pd.to_numeric(a.q_value,errors='coerce')<.05).sum())))
summary=pd.DataFrame(summary).merge(dec,on=['module_id','module']);summary.to_csv(O/'module_decisions.tsv',sep='\t',index=False)
for typ,key,mapper in [('metabolites','metabolite_name',mem),('relations','relation_key',joined['relations']),('genes','gene',gmap)]:
 d=data[typ].copy()
 if key not in d and typ=='relations':key='relation_id'
 lookup=mapper.groupby(key).module_id.agg(lambda x:';'.join(sorted(set(x)))).to_dict();d['review_modules']=d[key].map(lookup).fillna('OUTSIDE_CURATED_MODULES_RETAINED');d.to_csv(O/(typ+'_full_coverage.tsv'),sep='\t',index=False)
pd.DataFrame(sources).to_csv(O/'source_manifest.tsv',sep='\t',index=False)
spec=dict(input_commit=SHA,scientific_baseline='8663004',method='curated overlapping descriptive modules; all original statistics retained',new_P_q=0,new_patient_models=0,new_cell_analysis=0,cooccurrence_counts='NOT_RUN',literature_exhaustive_review='NOT_RUN',module_enrichment='NOT_RUN',candidate_priority='qualitative;no composite score')
(O/'analysis_spec.json').write_text(json.dumps(spec,indent=2),encoding='utf-8')
counts=dict(metabolites=int(data['metabolites'].metabolite_name.nunique()),relations=len(data['relations']),genes=len(data['genes']),modules=len(M),assigned_unique_features=mem.metabolite_name.nunique())
assert counts['metabolites']==318 and counts['relations']==262 and counts['genes']==156
for name,d in data.items():assert pd.read_csv(O/(name+'_frozen.tsv'),sep='\t',dtype=str,keep_default_na=False).equals(d)
(O/'validation.json').write_text(json.dumps(dict(counts=counts,all_frozen_tables_preserved=True,no_new_statistics=True,module_counts_overlap=True),indent=2),encoding='utf-8')
wb=Workbook();wb.remove(wb.active)
tables={'先读我':pd.DataFrame({'说明':['全量318特征、262关系、156基因保留；15个人工模块可重叠，不是富集分析','现有统计家族的P/q不变；模块数目不能相加作独立证据','未归入专题不代表无价值；没有重跑患者/单细胞/文献全审','同患者多代谢物共变与基因成对RNA方向尚未新计算']}),'模块判断':summary,'模块成员':mem,**{'模块_'+k:v for k,v in joined.items()},**{'全量_'+k:v for k,v in data.items()}}
for name,d in tables.items():
 ws=wb.create_sheet(name);ws.append(list(d.columns))
 for row in d.itertuples(index=False,name=None):ws.append([None if pd.isna(x) else ("'"+x if isinstance(x,str) and x.startswith(('=','+','@')) else x) for x in row])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='087F8C')
 for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=24
wb.save(O/'BRCA_跨层模块复核.xlsx')
(O/'.gitattributes').write_bytes(b'* -text\n')
report=['# BRCA跨层结果复核：从单基因转向问题模块','\n## 问题、输入与实际范围','固定输入8952c54中的公开汇总，统计基线8663004。遍历全部318代谢物、262直接关系、156历史并集基因，整理15个可重叠人工讨论模块。逐模块核对主/敏感性和来源，不代表独立复算患者数据、全量文献审计或生化关系穷尽。全量保留，无新增P/q。', '\n## 实际结果与判断']
for row in summary.to_dict('records'):
 report += [f'\n### {row["module_id"]} {row["module"]} — {row["decision"]}',row['observation'],'限制：'+row['limitations'],'下一项最小问题：'+row['next_minimum']]
report+=['\n## 新手解释与停止点','模块只是把相关问题放一起，不是新显著通路。优先讨论氨基酸分配、乙醇胺短链、胆碱/溶血磷脂、嘌呤、SHMT同工酶及烟酰胺的不同证据模式；不是选唯一第一基因。不同细胞来源只能提示组成背景，不证明细胞间代谢交流。','所有“多少对升/降”都是逐项计数，未计算同患者共同变化交集；不得据此声称同一批病例已发生串联转换。新工作应先问同患者是否共同变化，而不是遍历通路挑故事。','\n## 复现','python code/brca_module_review_v1.py；读取固定Git对象，输出目录存在则拒绝覆盖。源表TSV与模块关联表、Excel均实际生成。','\n## 生化背景来源','https://www.kegg.jp/entry/hsa00250','https://www.kegg.jp/entry/hsa00564','https://www.kegg.jp/entry/hsa00230','https://www.kegg.jp/entry/hsa00670','https://www.kegg.jp/entry/hsa00760','https://www.ncbi.nlm.nih.gov/gene/6652','https://www.kegg.jp/entry/K00884']
(O/'README_CN.md').write_text('\n\n'.join(report),encoding='utf-8')
pd.DataFrame([dict(path=p.name,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in sorted(O.iterdir())]).to_csv(O/'checksums.tsv',sep='\t',index=False)
print(json.dumps(counts));print(O)
