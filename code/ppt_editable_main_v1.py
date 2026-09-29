"""Native, editable CAMP main story. PDF drawing/text geometry is preserved, not flattened."""
from pathlib import Path
import io,json,math,hashlib,shutil
import pymupdf as fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
ROOT=Path(__file__).resolve().parents[1]
BASE=Path('D:/CodexData/visualizations/2026/09/25/01a0d8cd-4545-75b0-ba43-922588deea1a')
RUN='20260929T_editable_main_v1';O=ROOT/'results/BRCA/07_INTEGRATION'/RUN;O.mkdir(parents=True,exist_ok=True)
V=Path('D:/CodexData/visualizations/2026/09/29/01a0ec94-4db2-7ec1-8e34-4f52c52a8ce1/CAMP_main_editable_v1');V.mkdir(parents=True,exist_ok=True)
R=Presentation();R.slide_width=Inches(13.333333);R.slide_height=Inches(7.5)
G='00553B';INK='183D34';M='6A7973';PALE='F0F5F1';LINE='DDE7E0';UP='C77462';DOWN='468B9E';FONT='Microsoft YaHei'
LOGO=BASE/'CAMP_SYSU_style/image2.png';SOURCES=[];AUDIT=[]
def rgb(c):
 if isinstance(c,str):return RGBColor.from_string(c)
 return RGBColor(*[max(0,min(255,round(float(x)*255))) for x in c])
def alpha(parent,val):
 a=OxmlElement('a:alpha');a.set('val',str(round(val*100000)));parent.append(a)
def box(s,x,y,w,h,color=PALE,kind=MSO_SHAPE.RECTANGLE):
 sh=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background();sh._element.spPr.append(OxmlElement('a:effectLst'));return sh
def txt(s,x,y,w,h,t,size=18,c=INK,b=False,align=PP_ALIGN.LEFT):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for j,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.alignment=align;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=b;p.font.color.rgb=rgb(c);p.space_after=Pt(7)
 return sh
def line(s,x,y,x2,y2,c=LINE,w=1):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x2),Inches(y2));sh.line.color.rgb=rgb(c);sh.line.width=Pt(w);return sh
def arrow(s,x,y,w=.38):return box(s,x,y,w,.21,G,MSO_SHAPE.RIGHT_ARROW)
def notes(s,text):s.notes_slide.notes_text_frame.text=text
def base(title,sub='',num=None):
 s=R.slides.add_slide(R.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('FFFFFF')
 s.shapes.add_picture(str(LOGO),Inches(.48),Inches(.15),width=Inches(2.0),height=Inches(.83))
 txt(s,10,.43,2.65,.3,'CAMP  /  组会汇报',11,M,align=PP_ALIGN.RIGHT);line(s,.63,1.03,12.69,1.03);line(s,.63,1.03,1.4,1.03,G,2)
 txt(s,.64,1.27,12.05,.6,title,27,G,True)
 if sub:txt(s,.66,1.96,12,.48,sub,13,M)
 txt(s,12.08,7.04,.6,.25,f'{num or len(R.slides):02d}',11,G,True,PP_ALIGN.RIGHT)
 return s
def molecule(s,cx,cy,scale=.6):
 pts=[(cx+math.cos(i*math.pi/3)*scale,cy+math.sin(i*math.pi/3)*scale) for i in range(6)]
 for i,(x,y) in enumerate(pts):line(s,x,y,*pts[(i+1)%6],G,2)
 for i,(x,y) in enumerate(pts):box(s,x-.07,y-.07,.14,.14,G if i%2 else '99B77E',MSO_SHAPE.OVAL)
def dna(s,x,y):
 for i in range(9):
  a=math.sin(i*.7)*.25;yy=y+i*.09;line(s,x-a,yy,x+a,yy,'9DBEA9',1.5)
  for xx in [x-a,x+a]:box(s,xx-.025,yy-.025,.05,.05,G,MSO_SHAPE.OVAL)
def footer(s,t):line(s,.64,6.88,12.69,6.88);txt(s,.65,7.02,11.1,.25,t,9.5,M)

# 01 Title, all text and network editable.
s=base('',num=1)
txt(s,.73,1.66,8.6,.38,'CAMP · 泛癌代谢研究',15,M)
txt(s,.73,2.38,9.3,1.32,'从代谢异常出发\n寻找潜在干预候选',34,G,True)
txt(s,.76,4.04,8.6,.65,'跨癌种发现与多层证据整合｜以 LYPLA1 为重点',20,INK)
line(s,.76,5.02,8.3,5.02,G,1.5)
txt(s,.76,5.31,8,.5,'患者匹配  /  代谢物—基因关系  /  细胞与空间背景',15,M)
molecule(s,10.65,3.17,.7);dna(s,10.65,4.45)
box(s,.76,6.39,2.12,.39,PALE);txt(s,.87,6.47,2,.3,'组会汇报 · 2026.09',12,G)
notes(s,'本报告先说明CAMP资源、患者/样本匹配和映射方法，再展示泛癌共同方向，随后聚焦LYPLA1。到证据汇总结束，三个代表案例不纳入本文件。')

# 02 Real paper screenshot, editable surrounding content.
s=base('CAMP 是什么？','Cancer Atlas of Metabolic Profiles · 原始论文与资源定位')
box(s,.66,2.70,5.35,3.70,'F7F9F7');paper=BASE/'CAMP_IMAGE_fivepages_v3/paper_header_authors_real.png';im=Image.open(paper);w=4.98;h=w*im.height/im.width
s.shapes.add_picture(str(paper),Inches(.84),Inches(2.96),width=Inches(w),height=Inches(h))
txt(s,.88,5.60,4.98,.66,'Nature Metabolism · 2023\nBenedetti et al.',15,G,True)
txt(s,6.52,2.77,5.8,1.18,'整合肿瘤及参照组织的代谢组与转录组，\n研究基因表达与代谢物丰度如何共同变化。',21,INK)
for x,n,lab in [(6.55,'11','类癌症'),(8.55,'15','个数据集'),(10.65,'988','份组织标本')]:
 txt(s,x,4.35,1.8,.62,n,34,G,True);txt(s,x,5.07,1.8,.35,lab,15,M)
txt(s,6.55,5.82,5.75,.58,'核心：用可匹配的两种组学，连接代谢水平与基因表达。',17,G,True)
footer(s,'来源：Nature Metabolism 5, 1029–1044 (2023) · doi:10.1038/s42255-023-00817-8')
notes(s,'左侧为原始PDF截图，保留文章标题与作者，未用生成式工具伪造。CAMP共988标本（764肿瘤、224邻近正常），11癌种15数据集。组学匹配与同患者肿瘤—正常配对是两种不同匹配，逐队列核对，不能把全部988标本都称为配对患者。来源原始论文PDF，已在项目中核对。')

# 03 Two matching concepts emphasized.
s=base('为什么选择 CAMP？','患者配对控制个体背景，组学匹配连接代谢异常与候选基因')
xs=[.67,4.91,9.15];ws=[3.62,3.62,3.51]
for x,w,n,title in zip(xs,ws,['01','02','03'],['同患者组织配对','对应标本组学匹配','跨癌种比较']):
 box(s,x,2.71,w,3.59,'F4F8F4');txt(s,x+.22,2.9,.45,.4,n,19,'8BA696',True);txt(s,x+.79,2.91,w-.9,.5,title,19,G,True)
for xx,lab,c in [(1.07,'肿瘤',UP),(2.80,'参照',DOWN)]:
 box(s,xx,3.88,.86,.86,c,MSO_SHAPE.OVAL);txt(s,xx-.03,4.96,1,.4,lab,16,G,True,PP_ALIGN.CENTER)
line(s,1.51,3.7,3.23,3.7,G,1.6);txt(s,1.19,3.38,2.36,.3,'同一患者',13,M,align=PP_ALIGN.CENTER)
txt(s,.96,5.64,3.04,.5,'减少患者间基础差异',16,INK,True)
box(s,5.95,3.62,1.45,.52,'D9E8DE');txt(s,6.0,3.73,1.36,.3,'对应组织标本',13,G,True,PP_ALIGN.CENTER)
line(s,6.68,4.18,5.86,4.61,G,1.5);line(s,6.68,4.18,7.56,4.61,G,1.5)
for xx,lab in [(5.25,'代谢组'),(7.05,'RNA')]:box(s,xx,4.68,1.26,.49,'FFFFFF');txt(s,xx,4.78,1.26,.3,lab,15,G,True,PP_ALIGN.CENTER)
txt(s,5.22,5.64,3.06,.5,'在匹配标本中分析关联',16,INK,True)
for k,lab in enumerate(['BRCA','COAD','GBM','PDAC','PRAD','ccRCC']):
 xx=9.44+(k%2)*1.51;yy=3.72+(k//2)*.6;box(s,xx,yy,1.27,.43,'FFFFFF');txt(s,xx,yy+.075,1.27,.28,lab,13,G,True,PP_ALIGN.CENTER)
txt(s,9.44,5.64,3.06,.5,'识别共同方向与背景差异',15,INK,True)
footer(s,'仅在元数据支持的队列中使用患者配对；GBM 的跨来源参照单独解释。')
notes(s,'本页分别解释患者内肿瘤—参照配对、代谢组—转录组标本匹配，二者不可混同。相关分析仅使用匹配的肿瘤标本，不用正常与肿瘤混合引入组间差异。')

# 04 Cohort design, native grid.
s=base('六癌种、七队列：按真实设计分别分析','本项目使用 CAMP 的一个子集；配对、非配对与治疗背景保持区分')
coh=[('BRCA','乳腺癌','患者内配对'),('COAD','结肠癌','患者内配对'),('PDAC','胰腺导管癌','患者内配对'),('PRAD','前列腺癌','患者内配对'),('ccRCC','肾透明细胞癌','两个队列，分别分析'),('GBM','胶质母细胞瘤','非配对 · 跨来源参照')]
for i,(a,b,c) in enumerate(coh):
 x=.69+(i%3)*4.22;y=2.73+(i//3)*1.62
 box(s,x,y,3.66,1.32,PALE);box(s,x,y,.04,1.32,G if a!='GBM' else 'B69B60')
 txt(s,x+.23,y+.16,3,.4,a,22,G,True);txt(s,x+.24,y+.64,3.2,.3,b,15,INK);txt(s,x+.24,y+1.02,3.2,.27,c,12,M)
txt(s,.75,6.31,11.9,.39,'队列覆盖与参照不同，因此分别估计方向，再比较可重复的证据。',18,G,True)
footer(s,'ccRCC 两队列在癌种计数时仅计一次；细胞验证中的治疗背景另行标注。')
notes(s,'七队列BRCA、COAD、GBM、PDAC、PRAD、ccRCC3、ccRCC4。LYPLA1专题只纳入BRCA、COAD、PDAC、PRAD。禁止由患者编号或排序推断配对。')

# 05 Correct parallel-evidence workflow.
s=base('从代谢物到候选基因：生化映射与多层证据','先建立候选关系，再并行检查统计支持和细胞背景')
stages=[('差异代谢物','肿瘤—参照比较\n名义 P < 0.05'),('核对已有注释','名称与化学标识对齐\n异构体 / 混合峰分别处理'),('生化关系映射','反应底物 / 产物、直接转运\n必要复合体成员另注'),('代谢物—基因候选池','按代谢特征 × 基因去重\n条件性 / 未解决关系另列')]
for i,(a,b) in enumerate(stages):
 x=.68+i*3.18;box(s,x,2.79,2.75,1.61,PALE);txt(s,x+.17,3.0,2.43,.43,a,17,G,True);txt(s,x+.17,3.63,2.4,.64,b,12.5,M)
 if i<3:arrow(s,x+2.82,3.43,.29)
line(s,11.58,4.41,11.58,4.75,G,1.5);line(s,2.51,4.75,11.58,4.75,G,1.5)
for x,a,b in [(.69,'RNA 是否改变？','配对表达与患者方向一致性'),(4.96,'代谢与 RNA 是否关联？','匹配肿瘤标本中的实测相关'),(9.22,'候选在哪里表达？','单细胞定位与空间组织背景')]:
 line(s,x+1.67,4.75,x+1.67,5.1,G,1.5);box(s,x,5.1,3.43,1.10,'F7F9F7');txt(s,x+.15,5.26,3.1,.38,a,16,G,True);txt(s,x+.15,5.84,3.1,.3,b,12.5,M)
footer(s,'生化映射提出候选；患者关联与细胞证据并行检查，不以相关显著作为唯一入场条件。')
notes(s,'生化关系来源：已有审定关系、UniProt/Rhea及文献。同通路不等于直接关系；候选可以在关联未显著时继续检查功能与细胞背景。本图为工作流，不宣称因果或已测代谢通量。')

def el(tag,**attrs):
 a=OxmlElement(tag)
 for k,v in attrs.items():a.set(k,str(v))
 return a
def drawing(s,d,scale):
 rect=d['rect'];x0,y0,x1,y1=rect;w=max(.01,x1-x0);h=max(.01,y1-y0)
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x0*scale),Pt(y0*scale),Pt(w*scale),Pt(h*scale));sp=sh._element.spPr
 old=sp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
 if old is not None:sp.remove(old)
 g=el('a:custGeom')
 for tag in ['a:avLst','a:gdLst','a:ahLst','a:cxnLst']:g.append(el(tag))
 g.append(el('a:rect',l='0',t='0',r='r',b='b'));pl=el('a:pathLst');p=el('a:path',w=round(w*1000),h=round(h*1000));pl.append(p);g.append(pl);sp.insert(1,g)
 last=None
 def point(tag,pt):
  a=el(tag);a.append(el('a:pt',x=round((pt[0]-x0)*1000),y=round((pt[1]-y0)*1000)));p.append(a)
 for item in d['items']:
  kind=item[0]
  if kind=='re':
   r=item[1];coords=[r.tl,r.tr,r.br,r.bl] if item[2]==1 else [r.tl,r.bl,r.br,r.tr];point('a:moveTo',coords[0])
   for z in coords[1:]:point('a:lnTo',z)
   p.append(el('a:close'));last=None
  elif kind=='qu':
   q=item[1];point('a:moveTo',q.ul)
   for z in [q.ur,q.lr,q.ll]:point('a:lnTo',z)
   p.append(el('a:close'));last=None
  elif kind in ['l','c']:
   if last is None or abs(last[0]-item[1][0])+abs(last[1]-item[1][1])>.001:point('a:moveTo',item[1])
   if kind=='l':point('a:lnTo',item[2]);last=item[2]
   else:
    a=el('a:cubicBezTo')
    for z in item[2:]:a.append(el('a:pt',x=round((z[0]-x0)*1000),y=round((z[1]-y0)*1000)))
    p.append(a);last=item[-1]
 if d.get('closePath'):p.append(el('a:close'))
 if d.get('fill') is not None:
  sh.fill.solid();sh.fill.fore_color.rgb=rgb(d['fill']);alpha(sh.fill.fore_color._xFill[0],d.get('fill_opacity',1))
 else:sh.fill.background()
 if d.get('color') is not None:
  sh.line.color.rgb=rgb(d['color']);sh.line.width=Pt(max(.12,(d.get('width') or .5)*scale));alpha(sh.line.color._xFill[0],d.get('stroke_opacity',1))
 else:sh.line.fill.background()
 sp.append(el('a:effectLst'));sh.name='Editable vector geometry'
 return sh

def convert(path,page_num):
 doc=fitz.open(path);page=doc[0];scale=960/page.rect.width;s=R.slides.add_slide(R.slide_layouts[6]);blocks=page.get_text('dict')['blocks'];log=page.get_bboxlog();events=[]
 for d in page.get_drawings():events.append((d['seqno'],'path',d))
 images=[b for b in blocks if b['type']==1]
 for b in images:
  matches=[(sum(abs(x-y) for x,y in zip(b['bbox'],r)),i) for i,(k,r) in enumerate(log) if k=='fill-image'];seq=min(matches)[1] if matches else 1
  events.append((seq,'image',b))
 for i,(k,bb) in enumerate(log):
  if k=='fill-shade':events.append((i,'shade',bb))
 for seq,kind,d in sorted(events,key=lambda z:z[0]):
  if kind=='path':
   sh=drawing(s,d,scale)
   if page_num in [15,16] and d['rect'].height<.1 and d['rect'].width>page.rect.width*.8 and .4<d['rect'].y0/page.rect.height<.8:
    sh.width=Pt((page.rect.width*.905-d['rect'].x0)*scale)
  else:
   if kind=='image':
    bb=d['bbox'];data=d['image'];im=Image.open(io.BytesIO(data))
    if d.get('mask'):
     im=im.convert('RGBA');im.putalpha(Image.open(io.BytesIO(d['mask'])).convert('L'));out=io.BytesIO();im.save(out,format='PNG');data=out.getvalue()
   else:bb=d;data=page.get_pixmap(clip=fitz.Rect(bb),matrix=fitz.Matrix(3,3),alpha=True).tobytes('png')
   x,y,x2,y2=bb
   if x2>x and y2>y:s.shapes.add_picture(io.BytesIO(data),Pt(x*scale),Pt(y*scale),width=Pt((x2-x)*scale),height=Pt((y2-y)*scale)).name='Source image: point cloud / tissue / logo / color scale'
 native_text=0
 for b in blocks:
  if b['type']!=0:continue
  for ln in b['lines']:
   direction=ln['dir'];vertical=abs(direction[1])>.5
   for sp in ln['spans']:
    content=sp['text'];x,y,x2,y2=sp['bbox'];size=sp['size']*scale
    if page_num==10 and '为什么聚焦' in content:content='LYPLA1：四癌组织表达与患者方向一致性';x2=max(x2,x+790)
    if x>page.rect.width*.89 and y>page.rect.height*.92 and (content.strip().isdigit() or content.strip() in ['证据汇总','研究进展']):content=f'{page_num:02d}'
    if not content.strip():continue
    if vertical:
     width=(y2-y)*scale+size*.4;height=size*1.45;cx=(x+x2)/2*scale;cy=(y+y2)/2*scale;left=cx-width/2;top=cy-height/2
    else:left=x*scale;top=(sp['origin'][1]-sp['size']*1.075)*scale;width=(x2-x)*scale+size*.45;height=size*1.52
    sh=s.shapes.add_textbox(Pt(left),Pt(top),Pt(max(width,1)),Pt(height));tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    p=tf.paragraphs[0];p.text=content;p.font.name=FONT;p.font.size=Pt(size);p.font.bold='Bold' in sp['font'];p.font.color.rgb=RGBColor.from_string(f"{sp['color']:06X}");p.space_before=p.space_after=Pt(0);p.line_spacing=1.0
    if vertical:sh.rotation=270 if direction[1]<0 else 90
    sh.name='Editable text: '+content[:40];native_text+=1
 notes(s,f'来源：{path}\n原有冻结结果，不重算统计、不改变阈值。标题、文字、线条、点阵及患者连线为原生可编辑PPT对象。UMAP点云、H&E组织影像及色带为独立图像。原图表中的科学限制保留。')
 SOURCES.append(dict(slide=page_num,source=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
 AUDIT.append(dict(slide=page_num,native_text=native_text,shapes=len(s.shapes),raster_images=sum(sh.shape_type==13 for sh in s.shapes)))
 return s

selected=[('CAMP_pages_06_07_v2','06_*.pdf'),('CAMP_page07_patient_consistency_v4','07_*.pdf'),('CAMP_page08_patient_consistency_v3','08_*.pdf'),('CAMP_LYPLA1_metabolite_bridge_v2','09_*.pdf'),('CAMP_page09_LYPLA1_paired_v1','09_*.pdf')]
selected += [('CAMP_four_cancer_singlecell_v3',f'{n}_*.pdf') for n in range(11,15)]
selected += [('CAMP_selected_spatial_v1',f'{n}_*.pdf') for n in [15,16]]
selected += [('CAMP_LYPLA1_recent_literature_v1','LYPLA1_近期*.pdf'),('CAMP_LYPLA1_evidence_summary_v2','LYPLA1_多层*.pdf')]
for n,(folder,pat) in enumerate(selected,6):
 p=next((BASE/folder).glob(pat));convert(p,n);print('PAGE',n,p.name,flush=True)
assert len(R.slides)==18
# Consistent header on every page, independent of the source figure's margins.
for sn,slide in enumerate(R.slides,1):
 for sh in list(slide.shapes):
  if sh.top+sh.height<Inches(1.065):slide.shapes._spTree.remove(sh._element)
 slide.shapes.add_picture(str(LOGO),Inches(.65),Inches(.13),width=Inches(1.93),height=Inches(.75))
 txt(slide,10.02,.42,2.62,.28,'CAMP  /  组会汇报',11,M,align=PP_ALIGN.RIGHT)
 line(slide,.65,.94,12.67,.94);line(slide,.65,.94,1.40,.94,G,2)
 if sn in range(2,6):
  for sh in slide.shapes:
   if sh.has_text_frame and sh.text and 1.24<sh.top/Inches(1)<1.4:sh.top=Inches(1.20)
 for sh in slide.shapes:
  for e in sh._element.xpath('.//a:effectRef'):e.set('idx','0')
# Explicit East Asian fonts are essential: WPS otherwise substitutes a serif font.
for slide in R.slides:
 for sh in slide.shapes:
  for rp in sh._element.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr'):
   rp.set('lang','zh-CN')
   for tag,face in [('a:latin','Microsoft YaHei'),('a:ea','微软雅黑'),('a:cs','Microsoft YaHei')]:
    old=rp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+tag.split(':')[1])
    if old is not None:rp.remove(old)
    rp.append(el(tag,typeface=face))
name='CAMP_泛癌发现与LYPLA1_可编辑汇报.pptx';R.save(O/name)
for i,s in enumerate(R.slides,1):
 if i<=5:AUDIT.insert(i-1,dict(slide=i,native_text=sum(sh.has_text_frame and bool(sh.text) for sh in s.shapes),shapes=len(s.shapes),raster_images=sum(sh.shape_type==13 for sh in s.shapes)))
(O/'editable_object_audit.json').write_text(json.dumps(AUDIT,ensure_ascii=False,indent=2),encoding='utf-8')
(O/'source_manifest.json').write_text(json.dumps(SOURCES,ensure_ascii=False,indent=2),encoding='utf-8')
extra=[LOGO,paper,BASE/'CAMP_IMAGE_fivepages_v3/CAMP_original.pdf',BASE/'CAMP_IMAGE_mapping_redesign_v3/prompts.txt',Path('C:/Users/Administrator/Desktop/Pre_IBD.pptx')]
(O/'intro_source_manifest.json').write_text(json.dumps([dict(source=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest(),use='Resource definition, existing workflow, or user style reference') for p in extra],ensure_ascii=False,indent=2),encoding='utf-8')
for f in O.glob('*'):
 if f.is_file():shutil.copy2(f,V/f.name)
print('CREATED',O/name,flush=True)
