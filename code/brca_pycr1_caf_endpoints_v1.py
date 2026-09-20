"""Re-extract published CAF endpoint summaries, without new hypothesis tests."""
import argparse,hashlib,json,sys
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl

PREFIX='cancer cohort stage_id run_id analysis_version analysis_type metabolite_key metabolite_name gene unit n n_reference effect_type effect ci_lower ci_upper p_value q_value test_family family_n_evaluable status reason source_id'.split()
def main(root):
    src=root/'source';out=root/'public';out.mkdir(exist_ok=True)
    w=openpyxl.load_workbook(src/'42255_2022_582_MOESM14_ESM.xlsx',data_only=True)
    rows=[];checks=[]
    def add(panel,model,endpoint,condition,values,reference,reason):
        v=np.asarray(values,float);c=np.asarray(reference,float)
        r={k:'NA' for k in PREFIX}
        r.update(cancer='BRCA',cohort='Kay2022_'+model,stage_id='05_FUNCTION',run_id=root.name,analysis_version='caf_endpoint_v1',analysis_type='published_source_descriptive_reanalysis',gene='PYCR1',unit='author_experiment_replicate',n=len(v),n_reference=len(c),effect_type='ratio_of_group_means_on_author_normalized_scale',effect=float(v.mean()/c.mean()),status='DONE',reason=reason,source_id='10.1038/s42255-022-00582-0;'+panel)
        r.update(endpoint=endpoint,condition=condition,mean=float(v.mean()),sd=float(v.std(ddof=1)),reference_mean=float(c.mean()),reference_sd=float(c.std(ddof=1)),new_independent_validation=False)
        rows.append(r)
    # Fig3i has 3 explicitly labelled experiment blocks; verify author's protein/loading ratios.
    ws=w['Figure 3i'];x=[]
    for col in range(2,11):
        value=ws.cell(3,col).value/ws.cell(4,col).value
        assert np.isclose(value,ws.cell(6,col).value,rtol=1e-12)
        x.append(value)
    x=np.asarray(x).reshape(3,3)
    assert [ws.cell(1,c).value for c in [2,5,8]]==['Replicate 1','Replicate 2','Replicate 3']
    add('Fig3i','pCAF2','ECM_COL6A1_Ponceau','shPYCR1',x[:,1],x[:,0],'Three source replicate blocks; no new P/q; same published study')
    add('Fig3i','pCAF2','ECM_COL6A1_Ponceau','shPYCR1_plus_proline',x[:,2],x[:,0],'Proline supplementation condition; ratio to control, not percent mediation')
    checks.append({'check':'Fig3i_COL6A1_over_Ponceau','status':'PASS','n_cells_verified':9})
    # RNA uses author's pre-averaged technical replicate fields, not each qPCR well as independent n.
    for sheet,model,ctl,kd in [('Figure 3n','cCAF','sictl','sipycr1'),('Figure 3o','pCAF2','shctl','shpycr1')]:
        grouped={ctl:[],kd:[]}
        for row in w[sheet].values:
            if len(row)>5 and isinstance(row[5],(float,int)) and str(row[2]).upper()=='COL1A1':
                grouped[str(row[1]).lower()].append(row[5])
        assert len(grouped[ctl])==len(grouped[kd])==3
        add(sheet.replace('Figure ','Fig'),model,'COL1A1_RNA_author_reference_normalized',kd,grouped[kd],grouped[ctl],'Three author replicate averages; technical wells not counted; nonsignificance is not equivalence')
    pd.DataFrame(rows).to_csv(out/'endpoint_summary.tsv',sep='\t',index=False)
    s4=pd.read_excel(src/'42255_2022_582_MOESM6_ESM.xlsx')
    intensity=[c for c in s4 if str(c).startswith('Intensity sh')]
    assert len(intensity)==6 and not any('pycr1' in str(c).lower() for c in s4.columns)
    audit={'status':'NEEDS_REVIEW','source':'Supplementary Data 4','site_rows':len(s4),'unique_peptide_sequences':s4.Sequence.nunique(),'intensity_columns':intensity,'PYCR1_knockdown_column_found':False,'website_description_disagrees_with_table':True,'interpretation':'Labelled proline-containing COL1A1 peptide detection in control-labelled cultures; no available knockdown contrast; site rows not independent replicates'}
    (out/'supplement4_audit.json').write_text(json.dumps(audit,indent=2)+'\n')
    p=pd.read_excel(src/'42255_2022_582_MOESM14_ESM.xlsx',sheet_name='Figure 3p')
    a,b,c=[p.iloc[:,i].to_numpy(float) for i in [1,2,3]]
    assert p.iloc[:,0].is_unique
    pd.DataFrame([{'source':'Fig3p','n_ECM_gene_rows':len(p),'n_KD_higher_than_control':int((b>a).sum()),'n_proline_lower_than_KD':int((c<b).sum()),'median_control':float(np.median(a)),'median_KD':float(np.median(b)),'median_KD_plus_proline':float(np.median(c)),'status':'DONE','reason':'Descriptive across author-selected genes; no replicate-resolved columns; genes are not independent experiment replicates; no P/q'}]).to_csv(out/'translation_summary.tsv',sep='\t',index=False)
    issues=[{'source':'Fig3g','status':'NEEDS_REVIEW','issue':'Three intervention rows labelled siPYCR1,siPYCR2,siPYCR3; cannot silently relabel as PYCR1 replicate1/2/3'}, {'source':'Fig3h','status':'NEEDS_REVIEW','issue':'Four labelled replicate blocks in workbook vs three biological replicates in figure caption; no selection or formal test'}, {'source':'Fig3p','status':'PARTIAL','issue':'22 gene summaries without three separate biological replicate columns; no replicate-level inference'}, {'source':'GSE220931_ENA','status':'NEEDS_REVIEW','issue':'Nine ENA run/sample titles corroborate GEO titles but do not resolve conflicting genotype fields, processed matrix units or biological independence'}]
    pd.DataFrame(issues).to_csv(out/'unresolved_items.tsv',sep='\t',index=False)
    ena=pd.read_csv(src/'GSE220931_ENA_metadata.tsv',sep='\t');assert len(ena)==9 and ena.run_accession.is_unique
    validation={'status':'DONE','scope':'bounded descriptive endpoint re-extraction; not full source-data validation','checks':checks,'RNA_biological_averages_per_group':3,'S4_rows':len(s4),'S4_unique_sequences':int(s4.Sequence.nunique()),'ENA_runs':len(ena),'new_tests':0,'causal_or_independent_replication_claim':False,'unresolved_items_retained':len(issues),'runtime':{'python':sys.version,'numpy':np.__version__,'pandas':pd.__version__,'openpyxl':openpyxl.__version__}}
    (out/'validation.json').write_text(json.dumps(validation,indent=2)+'\n')
    manifest=[]
    for f in sorted(src.iterdir()):
        if f.is_file():manifest.append({'path':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
    pd.DataFrame(manifest).to_csv(out/'source_manifest.tsv',sep='\t',index=False)
    print(pd.DataFrame(rows)[['endpoint','condition','effect']].to_string(index=False));print(json.dumps(audit));print(json.dumps(validation))
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--root',type=Path,required=True);main(p.parse_args().root)
